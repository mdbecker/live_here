#!/usr/bin/env python3
"""
Update the county-rankings workbook with NOAA/NCEI annual/seasonal
multivariate climate normals, re-rank affected climate factors, rerun the
100k randomized runoff MCMC, and write an updated workbook plus an audit CSV.

Expected inputs
---------------
1. Existing county-rankings workbook with sheets:
   - Rank Matrix
   - Value Matrix
   - EPA Walkability Update, used for county FIPS mapping
2. NOAA/NCEI annual/seasonal multivariate normals tar.gz containing one CSV per
   station. The useful fields for this pass are:
   - ANN-SNOW-NORMAL: annual normal snowfall in inches
   - ANN-PRCP-NORMAL: annual normal precipitation in inches
   - ANN-PRCP-AVGNDS-GE001HI: annual days with precipitation >= 0.01 inches
   - seasonal PRCP normals such as JJA-PRCP-NORMAL
3. EPA WalkabilityIndex ZIP from the previous workflow, used only to compute
   population-weighted county centroids from block-group geometry/population.

Method
------
- Annual snowfall is updated directly from NOAA ANN-SNOW-NORMAL. Station values
  are inverse-distance-squared blended to county population-weighted centroids.
  Inches are converted to feet.
- Drought days are not directly observed in this NOAA normals file. Actual D1+
  drought frequency should ultimately come from the U.S. Drought Monitor.
  However, the multivariate normals do provide a better real-data dryness proxy
  than the prior hand estimate. The script builds a NOAA drought-risk proxy from:
      * dry days = 365.25 - ANN-PRCP-AVGNDS-GE001HI
      * annual precipitation scarcity
      * summer precipitation scarcity
  The proxy is then quantile-mapped onto the previous workbook's D1+ drought-day
  scale so the units remain comparable while the ordering is driven by NOAA
  precipitation normals.
- Missing-data policy:
    * Precipitation/drought: nearest-station interpolation first; calibrated old
      estimate fallback if a county lacks a usable station blend.
    * Snowfall: nearest snow-normal station interpolation first; calibrated old
      estimate fallback for missing/unreliable snow coverage. Missing snow is
      not treated as zero unless supported by nearby zero-snow stations or prior
      modeled zero snow.
- Updates:
    * Value Matrix -> Est. annual snowfall, ft
    * Value Matrix -> Est. D1+ drought days/year
    * Rank Matrix -> Annual snowfall rank
    * Rank Matrix -> Drought days rank
  Both ranks are lower-is-better.
- Reruns randomized runoff with the same rank-matrix rules: random permutation
  cycles over rank factors; the worst rank on the selected factor is eliminated;
  tied worst rows are broken randomly.

Install dependencies
--------------------
pip install pandas numpy scipy pyproj pyogrio openpyxl numba

Example
-------
python update_noaa_multivariate_and_rerun_mcmc.py \
  --input-workbook county_rankings_epa_walkability_transit_aqi_noaa_temperature_100k.xlsx \
  --noaa-tar us-climate-normals_2006-2020_v1.0.1_annualseasonal_multivariate_by-station_c20230404.tar.gz \
  --epa-walkability-zip EPA_WalkabilityIndex.zip \
  --output-workbook county_rankings_epa_walkability_transit_aqi_noaa_temperature_multivariate_100k.xlsx \
  --county-rollup-csv noaa_multivariate_county_rollup.csv
"""

from __future__ import annotations

import argparse
import csv
import io
import math
import re
import statistics
import tarfile
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

import numpy as np
import pandas as pd
from numba import njit
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
try:
    from openpyxl.workbook.properties import CalcProperties
except Exception:  # pragma: no cover
    CalcProperties = None
from openpyxl.utils import get_column_letter
from pyproj import Transformer
from scipy.spatial import cKDTree
import pyogrio

DEFAULT_ITERATIONS = 100_000
DEFAULT_SEED = 2026060901

RANK_MATRIX_SHEET = "Rank Matrix"
VALUE_MATRIX_SHEET = "Value Matrix"
TOP30_SHEET = "Top 30"
RESULT_COMPARISON_SHEET = "Result Comparison"
RUNOFF_SHEET = "Runoff Simulation"
METHODOLOGY_SHEET = "Sources & Methodology"
LOG_SHEET = "Data Acquisition Log"
QA_SHEET = "QA Checks"
WALKABILITY_AUDIT_SHEET = "EPA Walkability Update"
IMPACT_SHEET = "NOAA Multivariate Impact"
UPDATE_SHEET = "NOAA Multivariate Update"

DROUGHT_VALUE_COL = "Est. D1+ drought days/year"
DROUGHT_RANK_COL = "Drought days rank"
SNOW_VALUE_COL = "Est. annual snowfall, ft"
SNOW_RANK_COL = "Annual snowfall rank"

HEADER_FILL = "1F4E79"
BORDER_COLOR = "B7B7B7"

STATE_DIVISION = {
    "CT": "New England", "ME": "New England", "MA": "New England", "NH": "New England", "RI": "New England", "VT": "New England",
    "NJ": "Middle Atlantic", "NY": "Middle Atlantic", "PA": "Middle Atlantic",
    "IL": "East North Central", "IN": "East North Central", "MI": "East North Central", "OH": "East North Central", "WI": "East North Central",
    "IA": "West North Central", "KS": "West North Central", "MN": "West North Central", "MO": "West North Central", "NE": "West North Central", "ND": "West North Central", "SD": "West North Central",
    "DE": "South Atlantic", "DC": "South Atlantic", "FL": "South Atlantic", "GA": "South Atlantic", "MD": "South Atlantic", "NC": "South Atlantic", "SC": "South Atlantic", "VA": "South Atlantic", "WV": "South Atlantic",
    "AL": "East South Central", "KY": "East South Central", "MS": "East South Central", "TN": "East South Central",
    "AR": "West South Central", "LA": "West South Central", "OK": "West South Central", "TX": "West South Central",
    "AZ": "Mountain", "CO": "Mountain", "ID": "Mountain", "MT": "Mountain", "NV": "Mountain", "NM": "Mountain", "UT": "Mountain", "WY": "Mountain",
    "AK": "Pacific", "CA": "Pacific", "HI": "Pacific", "OR": "Pacific", "WA": "Pacific",
}


def normalize_county_name(name: str) -> str:
    if name is None:
        return ""
    s = str(name).strip()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.upper().replace("ST.", "ST").replace("SAINT ", "ST ")
    s = s.replace("'", "")
    s = re.sub(r"\s*\(.*?\)", "", s)
    suffixes = [
        " COUNTY AND BOROUGH", " CITY AND BOROUGH", " CENSUS AREA",
        " MUNICIPALITY", " CONSOLIDATED GOVERNMENT", " UNIFIED GOVERNMENT",
        " BOROUGH", " PARISH", " COUNTY", " CITY",
    ]
    changed = True
    while changed:
        changed = False
        for suffix in suffixes:
            if s.endswith(suffix):
                s = s[: -len(suffix)]
                changed = True
                break
    s = re.sub(r"[^A-Z0-9]+", " ", s).strip()
    return re.sub(r"\s+", " ", s)


def safe_float(x):
    if x is None or x == "":
        return None
    try:
        val = float(str(x).strip())
    except Exception:
        return None
    if not math.isfinite(val) or val in {-9999.0, -999.0, -8888.0}:
        return None
    return val


def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = list(rows[0])
    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        padded = list(row) + [None] * (len(headers) - len(row))
        data.append(dict(zip(headers, padded[: len(headers)])))
    return headers, data


def write_matrix(ws, matrix: List[List[object]]):
    ws.delete_rows(1, ws.max_row)
    for r_idx, row in enumerate(matrix, 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, value)


def recreate_sheet(wb, name: str, after_name: str | None = None):
    if name in wb.sheetnames:
        del wb[name]
    if after_name and after_name in wb.sheetnames:
        idx = wb.sheetnames.index(after_name)
        return wb.create_sheet(name, idx + 1)
    return wb.create_sheet(name)


def style_sheet(ws, max_width: int = 42):
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color=BORDER_COLOR)
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 8), max_width)


def get_fips_for_workbook(wb) -> Dict[Tuple[str, str], str]:
    if WALKABILITY_AUDIT_SHEET not in wb.sheetnames:
        raise ValueError(f"Workbook must contain {WALKABILITY_AUDIT_SHEET!r} for County FIPS mapping")
    headers, rows = read_sheet(wb[WALKABILITY_AUDIT_SHEET])
    if not {"County", "State", "County FIPS"}.issubset(set(headers)):
        raise ValueError(f"{WALKABILITY_AUDIT_SHEET!r} must include County, State, and County FIPS columns")
    return {
        (r["County"], r["State"]): str(r["County FIPS"]).zfill(5)
        for r in rows
        if r.get("County") is not None and r.get("State") is not None and r.get("County FIPS") is not None
    }


def rank_average(values: Iterable[float], higher_is_better: bool) -> List[float]:
    pairs = [(float(v), i) for i, v in enumerate(values)]
    pairs.sort(key=lambda t: t[0], reverse=higher_is_better)
    ranks = [0.0] * len(pairs)
    i = 0
    while i < len(pairs):
        j = i + 1
        while j < len(pairs) and pairs[j][0] == pairs[i][0]:
            j += 1
        avg_rank = (i + 1 + j) / 2.0
        for k in range(i, j):
            ranks[pairs[k][1]] = avg_rank
        i = j
    return ranks


@njit
def run_randomized_runoff(ranks: np.ndarray, sorted_order: np.ndarray, iterations: int, seed: int):
    np.random.seed(seed)
    n, f = ranks.shape
    wins = np.zeros(n, np.int64)
    elim_sum = np.zeros(n, np.float64)
    factor_order = np.empty(f, np.int64)
    active = np.empty(n, np.bool_)
    ptr = np.empty(f, np.int64)
    ties = np.empty(n, np.int64)

    for _ in range(iterations):
        for i in range(n):
            active[i] = True
        for j in range(f):
            ptr[j] = 0
        active_count = n
        pos = f
        round_no = 0

        while active_count > 1:
            if pos >= f:
                for j in range(f):
                    factor_order[j] = j
                for j in range(f - 1, 0, -1):
                    k = np.random.randint(j + 1)
                    tmp = factor_order[j]
                    factor_order[j] = factor_order[k]
                    factor_order[k] = tmp
                pos = 0

            col = factor_order[pos]
            pos += 1

            p = ptr[col]
            while p < n and not active[sorted_order[col, p]]:
                p += 1
            ptr[col] = p

            first_idx = sorted_order[col, p]
            worst_rank = ranks[first_idx, col]
            tie_count = 0
            q = p
            while q < n:
                row_idx = sorted_order[col, q]
                if ranks[row_idx, col] != worst_rank:
                    break
                if active[row_idx]:
                    ties[tie_count] = row_idx
                    tie_count += 1
                q += 1

            loser = ties[np.random.randint(tie_count)]
            active[loser] = False
            round_no += 1
            elim_sum[loser] += round_no
            active_count -= 1

        for i in range(n):
            if active[i]:
                winner = i
                break
        wins[winner] += 1
        elim_sum[winner] += n

    return wins, elim_sum / iterations


def clean_cell(value):
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return None if math.isnan(float(value)) else float(value)
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def mean_abs(rows, col):
    vals = [abs(float(r[col])) for r in rows if r.get(col) is not None]
    return statistics.mean(vals) if vals else None


def median_abs(rows, col):
    vals = [abs(float(r[col])) for r in rows if r.get(col) is not None]
    return statistics.median(vals) if vals else None


def median_or_none(vals):
    vals = [float(v) for v in vals if v is not None and math.isfinite(float(v))]
    return statistics.median(vals) if vals else None


def get_years(row, variable):
    return safe_float(row.get(f"years_{variable}"))


def parse_noaa_multivariate_station_metrics(noaa_tar: str):
    """Parse one-row annual/seasonal station CSVs into station-level metrics."""
    station_rows = []
    files_read = 0
    with tarfile.open(noaa_tar, "r:gz") as tar:
        for member in tar.getmembers():
            if not member.isfile() or not member.name.lower().endswith(".csv"):
                continue
            f = tar.extractfile(member)
            if f is None:
                continue
            try:
                text = io.TextIOWrapper(f, encoding="utf-8", errors="replace", newline="")
                reader = csv.DictReader(text)
                row = next(reader, None)
            except Exception:
                continue
            if not row:
                continue
            files_read += 1
            sid = row.get("STATION") or Path(member.name).stem
            lat = safe_float(row.get("LATITUDE"))
            lon = safe_float(row.get("LONGITUDE"))
            elev_m = safe_float(row.get("ELEVATION"))
            if lat is None or lon is None:
                continue
            rec = {
                "GHCN_ID": sid,
                "lat": lat,
                "lon": lon,
                "elev_m": elev_m,
                "station_name": row.get("NAME"),
            }
            # Precipitation metrics.
            ann_prcp = safe_float(row.get("ANN-PRCP-NORMAL"))
            wet_days = safe_float(row.get("ANN-PRCP-AVGNDS-GE001HI"))
            prcp_years = get_years(row, "ANN-PRCP-NORMAL")
            wet_years = get_years(row, "ANN-PRCP-AVGNDS-GE001HI")
            if ann_prcp is not None and ann_prcp >= 0 and (prcp_years is None or prcp_years >= 10):
                rec["ann_prcp_in"] = ann_prcp
            if wet_days is not None and 0 <= wet_days <= 366 and (wet_years is None or wet_years >= 10):
                rec["wet_days_ge001"] = wet_days
            for season in ["DJF", "MAM", "JJA", "SON"]:
                val = safe_float(row.get(f"{season}-PRCP-NORMAL"))
                years = get_years(row, f"{season}-PRCP-NORMAL")
                if val is not None and val >= 0 and (years is None or years >= 10):
                    rec[f"{season.lower()}_prcp_in"] = val
            # Snow metrics.
            snow_in = safe_float(row.get("ANN-SNOW-NORMAL"))
            snow_years = get_years(row, "ANN-SNOW-NORMAL")
            if snow_in is not None and snow_in >= 0 and (snow_years is None or snow_years >= 10):
                rec["ann_snow_in"] = snow_in
            snow_days = safe_float(row.get("ANN-SNOW-AVGNDS-GE010TI"))
            snow_days_years = get_years(row, "ANN-SNOW-AVGNDS-GE010TI")
            if snow_days is not None and 0 <= snow_days <= 366 and (snow_days_years is None or snow_days_years >= 10):
                rec["snow_days_ge1in"] = snow_days
            # Extra diagnostics, not used for ranking in this pass.
            for temp_col in ["ANN-TAVG-NORMAL", "ANN-TMAX-NORMAL", "ANN-TMIN-NORMAL", "ANN-HTDD-NORMAL", "ANN-CLDD-NORMAL"]:
                val = safe_float(row.get(temp_col))
                years = get_years(row, temp_col)
                if val is not None and (years is None or years >= 10):
                    rec[temp_col.lower().replace("-", "_")] = val
            station_rows.append(rec)

    df = pd.DataFrame(station_rows)
    if df.empty:
        raise ValueError("No station rows parsed from NOAA multivariate tar")

    # Build a station-level drought proxy. Values are dimensionless at this stage;
    # county-level values are quantile-mapped to the old D1+ drought-days scale.
    prcp_df = df.dropna(subset=["ann_prcp_in", "wet_days_ge001"]).copy()
    prcp_df["dry_days_ge001"] = (365.25 - prcp_df["wet_days_ge001"]).clip(0, 365.25)
    median_prcp = float(prcp_df["ann_prcp_in"].clip(lower=1.0).median())
    # Summer dryness is important because water stress during growing/warm season
    # is more relevant to drought stress than the same dry days in winter.
    if "jja_prcp_in" in prcp_df.columns:
        prcp_df["jja_share"] = (prcp_df["jja_prcp_in"] / prcp_df["ann_prcp_in"].replace(0, np.nan)).clip(0.02, 1.0)
        median_jja_share = float(prcp_df["jja_share"].median())
    else:
        prcp_df["jja_share"] = np.nan
        median_jja_share = 0.25
    prcp_df["drought_proxy_raw"] = (
        prcp_df["dry_days_ge001"]
        * np.sqrt(median_prcp / prcp_df["ann_prcp_in"].clip(lower=1.0))
        * np.sqrt(median_jja_share / prcp_df["jja_share"].fillna(median_jja_share).clip(lower=0.05))
    )
    for col in ["dry_days_ge001", "drought_proxy_raw"]:
        df.loc[prcp_df.index, col] = prcp_df[col]

    stats = {
        "station_csv_files_read": files_read,
        "station_rows_with_coordinates": int(len(df)),
        "stations_with_precip_normal": int(df["ann_prcp_in"].notna().sum()) if "ann_prcp_in" in df else 0,
        "stations_with_drought_proxy": int(df["drought_proxy_raw"].notna().sum()) if "drought_proxy_raw" in df else 0,
        "stations_with_snow_normal": int(df["ann_snow_in"].notna().sum()) if "ann_snow_in" in df else 0,
        "median_station_prcp_in": median_prcp if len(prcp_df) else None,
        "median_station_jja_share": median_jja_share if len(prcp_df) else None,
    }
    return df, stats



def _unit_sphere_xyz(lat_deg, lon_deg):
    lat = np.radians(np.asarray(lat_deg, dtype=float))
    lon = np.radians(np.asarray(lon_deg, dtype=float))
    clat = np.cos(lat)
    return np.column_stack([clat * np.cos(lon), clat * np.sin(lon), np.sin(lat)])


def _haversine_km_from_chord(chord):
    chord = np.clip(np.asarray(chord, dtype=float), 0.0, 2.0)
    return 2.0 * 6371.0088 * np.arcsin(chord / 2.0)


def _parse_semicolon_weights(ids_text, weights_text):
    ids = [x.strip() for x in str(ids_text).split(';') if x.strip()]
    weights = []
    for x in str(weights_text).split(';'):
        x = x.strip()
        if not x:
            continue
        try:
            weights.append(float(x))
        except Exception:
            pass
    n = min(len(ids), len(weights))
    ids, weights = ids[:n], weights[:n]
    if not ids:
        return [], []
    total = sum(w for w in weights if math.isfinite(w) and w > 0)
    if total <= 0:
        weights = [1.0 / len(ids)] * len(ids)
    else:
        weights = [max(0.0, w) / total for w in weights]
    return ids, weights


def build_county_rollup_from_temperature_station_weights(stations: pd.DataFrame, station_weight_csv: str):
    """Reuse the county-to-station weights produced by the prior NOAA temperature step.

    This is much faster and more reproducible than rereading EPA block-group geometry.
    Those weights were already calculated from population-weighted county centroids.
    Precipitation is calculated from the same weighted station set. Snowfall is calculated
    from the nearest NOAA snow-normal stations to a pseudo-centroid inferred from that
    weighted station set, because many temperature/precipitation stations lack snow normals.
    """
    weights_df = pd.read_csv(station_weight_csv, dtype={"County FIPS": str})
    required = {"County", "State", "County FIPS", "Station IDs used", "Station weights"}
    missing = required - set(weights_df.columns)
    if missing:
        raise ValueError(f"Station weight CSV missing required columns: {sorted(missing)}")
    station_by_id = stations.set_index("GHCN_ID", drop=False)
    rows = []
    pseudo_lat = []
    pseudo_lon = []
    # Precip/drought from prior station blend.
    precip_cols = ["ann_prcp_in", "wet_days_ge001", "dry_days_ge001", "drought_proxy_raw"]
    if "jja_prcp_in" in stations.columns:
        precip_cols.append("jja_prcp_in")
    for _, src in weights_df.iterrows():
        ids, weights = _parse_semicolon_weights(src.get("Station IDs used"), src.get("Station weights"))
        rec = {
            "County": src.get("County"),
            "State": src.get("State"),
            "County FIPS": str(src.get("County FIPS")).zfill(5),
        }
        coord_rows = []
        coord_weights = []
        precip_rows = []
        precip_weights = []
        for sid, w in zip(ids, weights):
            if sid not in station_by_id.index:
                continue
            st = station_by_id.loc[sid]
            if isinstance(st, pd.DataFrame):
                st = st.iloc[0]
            if pd.notna(st.get("lat")) and pd.notna(st.get("lon")):
                coord_rows.append(st)
                coord_weights.append(w)
            if all(pd.notna(st.get(col)) for col in precip_cols):
                precip_rows.append(st)
                precip_weights.append(w)
        if coord_rows:
            cw = np.asarray(coord_weights, dtype=float)
            cw = cw / cw.sum() if cw.sum() > 0 else np.ones(len(cw)) / len(cw)
            lat = float(np.sum([float(st["lat"]) * w for st, w in zip(coord_rows, cw)]))
            lon = float(np.sum([float(st["lon"]) * w for st, w in zip(coord_rows, cw)]))
        else:
            lat = np.nan
            lon = np.nan
        pseudo_lat.append(lat)
        pseudo_lon.append(lon)
        rec["Pseudo centroid lat from temp stations"] = lat
        rec["Pseudo centroid lon from temp stations"] = lon
        if precip_rows:
            pw = np.asarray(precip_weights, dtype=float)
            pw = pw / pw.sum() if pw.sum() > 0 else np.ones(len(pw)) / len(pw)
            for col in precip_cols:
                rec[col] = float(np.sum([float(st[col]) * w for st, w in zip(precip_rows, pw)]))
            rec["precip status"] = "NOAA precip reused temperature-station blend"
            rec["precip method"] = "reused prior NOAA temperature step county station weights; renormalized to stations with precipitation metrics"
            rec["precip stations used"] = int(len(precip_rows))
            rec["precip nearest station km"] = src.get("Nearest station km")
            rec["precip max station km used"] = src.get("Max station km used")
            rec["precip station IDs used"] = ";".join(str(st["GHCN_ID"]) for st in precip_rows)
            rec["precip station names used"] = ";".join(str(st.get("station_name")) for st in precip_rows)
            rec["precip station weights"] = ";".join(f"{w:.4f}" for w in pw)
        else:
            rec["precip status"] = "calibrated fallback; no reusable precipitation station metrics"
            rec["precip method"] = "no prior weighted stations had precipitation metrics"
        rows.append(rec)

    rollup = pd.DataFrame(rows)

    # Snow from nearest snow-normal stations to pseudo centroids.
    snow_stations = stations.dropna(subset=["ann_snow_in", "lat", "lon"]).copy()
    if snow_stations.empty:
        return rollup, {"snow_station_count": 0, "precip_station_weight_rows": len(weights_df)}
    snow_xyz = _unit_sphere_xyz(snow_stations["lat"].to_numpy(), snow_stations["lon"].to_numpy())
    tree = cKDTree(snow_xyz)
    county_xyz = _unit_sphere_xyz(rollup["Pseudo centroid lat from temp stations"].to_numpy(), rollup["Pseudo centroid lon from temp stations"].to_numpy())
    query_k = min(8, len(snow_stations))
    chord, idxs = tree.query(county_xyz, k=query_k)
    if query_k == 1:
        chord = chord[:, None]
        idxs = idxs[:, None]
    dists_km_all = _haversine_km_from_chord(chord)
    for i in range(len(rollup)):
        ds = dists_km_all[i]
        ids_idx = idxs[i]
        nearby = np.where(ds <= 175.0)[0]
        if len(nearby) >= 2:
            pos = nearby[:5]
            status = "NOAA snow nearby station blend"
            method = "nearest NOAA snow-normal stations within 175 km of pseudo-centroid from prior temperature interpolation"
        else:
            pos = np.arange(min(5, len(ids_idx)))
            status = "NOAA snow nearest-station fallback"
            method = "fewer than two snow-normal stations within 175 km; used nearest snow-normal stations to pseudo-centroid"
        dsel = np.maximum(ds[pos], 1.0)
        w = 1.0 / (dsel ** 2)
        w = w / w.sum()
        sub = snow_stations.iloc[ids_idx[pos]].copy()
        rollup.loc[i, "ann_snow_in"] = float(np.sum(sub["ann_snow_in"].to_numpy(dtype=float) * w))
        rollup.loc[i, "snow status"] = status
        rollup.loc[i, "snow method"] = method
        rollup.loc[i, "snow stations used"] = int(len(pos))
        rollup.loc[i, "snow nearest station km"] = float(ds[0])
        rollup.loc[i, "snow max station km used"] = float(ds[pos].max())
        rollup.loc[i, "snow station IDs used"] = ";".join(sub["GHCN_ID"].astype(str).tolist())
        rollup.loc[i, "snow station names used"] = ";".join(sub["station_name"].astype(str).tolist())
        rollup.loc[i, "snow station weights"] = ";".join(f"{x:.4f}" for x in w)
    return rollup, {"snow_station_count": len(snow_stations), "precip_station_weight_rows": len(weights_df)}


def compute_county_centroids(
    workbook_fips: Iterable[str],
    epa_walkability_zip: str | None = None,
    county_centroids_csv: str | None = None,
):
    if county_centroids_csv and Path(county_centroids_csv).exists():
        centroids = pd.read_csv(county_centroids_csv, dtype={"fips5": str})
        required = {"fips5", "county_x", "county_y", "crs_wkt"}
        missing = required - set(centroids.columns)
        if missing:
            raise ValueError(f"County centroids CSV missing required columns: {sorted(missing)}")
        return centroids

    if not epa_walkability_zip or not Path(epa_walkability_zip).exists():
        raise ValueError(
            "Need either --county-centroids-csv or --epa-walkability-zip to compute county centroids. "
            "Use EPA_WalkabilityIndex.zip from the prior workflow."
        )

    vsi_path = f"/vsizip/{Path(epa_walkability_zip).resolve()}/Natl_WI.gdb"
    gdf = pyogrio.read_dataframe(
        vsi_path,
        layer="NationalWalkabilityIndex",
        columns=["STATEFP", "COUNTYFP", "TotPop"],
        read_geometry=True,
    )
    crs_wkt = str(gdf.crs)
    target_fips = set(str(f).zfill(5) for f in workbook_fips if f is not None)
    gdf["fips5"] = gdf["STATEFP"].astype(str).str.zfill(2) + gdf["COUNTYFP"].astype(str).str.zfill(3)
    gdf = gdf[gdf["fips5"].isin(target_fips)].copy()
    gdf["TotPop"] = pd.to_numeric(gdf["TotPop"], errors="coerce").fillna(0)
    cent = gdf.geometry.centroid
    tmp = pd.DataFrame({
        "fips5": gdf["fips5"].values,
        "pop": gdf["TotPop"].values,
        "x": cent.x,
        "y": cent.y,
    })
    rows = []
    for fips, grp in tmp.groupby("fips5"):
        weights = grp["pop"].clip(lower=0).to_numpy()
        if weights.sum() > 0:
            x = float(np.average(grp["x"].to_numpy(), weights=weights))
            y = float(np.average(grp["y"].to_numpy(), weights=weights))
        else:
            x = float(grp["x"].mean())
            y = float(grp["y"].mean())
        rows.append({
            "fips5": str(fips).zfill(5),
            "county_x": x,
            "county_y": y,
            "bg_count": int(len(grp)),
            "population_used": float(weights.sum()),
            "crs_wkt": crs_wkt,
        })
    centroids = pd.DataFrame(rows)
    if county_centroids_csv:
        Path(county_centroids_csv).parent.mkdir(parents=True, exist_ok=True)
        centroids.to_csv(county_centroids_csv, index=False)
    return centroids


def project_stations(stations: pd.DataFrame, crs_wkt: str) -> pd.DataFrame:
    transformer = Transformer.from_crs("EPSG:4269", crs_wkt, always_xy=True)
    x, y = transformer.transform(stations["lon"].to_numpy(), stations["lat"].to_numpy())
    out = stations.copy()
    out["x"] = x
    out["y"] = y
    out = out[np.isfinite(out["x"]) & np.isfinite(out["y"])].copy()
    return out


def interpolate_metric_to_counties(
    stations: pd.DataFrame,
    centroids: pd.DataFrame,
    metric_cols: List[str],
    metric_label: str,
    k: int = 5,
    radius_km: float = 150.0,
    min_nearby: int = 2,
):
    """Inverse-distance-squared station blend for one or more station columns."""
    crs_wkt = centroids["crs_wkt"].iloc[0]
    station_cols = ["GHCN_ID", "station_name", "lat", "lon", "elev_m"] + metric_cols
    use = stations.dropna(subset=metric_cols)[station_cols].copy()
    use = project_stations(use, crs_wkt)
    if use.empty:
        raise ValueError(f"No finite NOAA station coordinates for {metric_label}")
    tree = cKDTree(use[["x", "y"]].to_numpy())
    query_k = min(max(k * 2, k), len(use))
    dists_m, idxs = tree.query(centroids[["county_x", "county_y"]].to_numpy(), k=query_k)
    if query_k == 1:
        dists_m = dists_m[:, None]
        idxs = idxs[:, None]

    rows = []
    for ci, county in centroids.reset_index(drop=True).iterrows():
        ds_km = dists_m[ci] / 1000.0
        ids = idxs[ci]
        nearby_positions = np.where(ds_km <= radius_km)[0]
        if len(nearby_positions) >= min_nearby:
            use_pos = nearby_positions[:k]
            status = f"NOAA {metric_label} nearby station blend"
            method = f"inverse-distance-squared blend of nearest {metric_label} stations within {radius_km:g} km"
        else:
            use_pos = np.arange(min(k, len(ids)))
            status = f"NOAA {metric_label} nearest-station fallback"
            method = f"fewer than {min_nearby} {metric_label} stations within {radius_km:g} km; used nearest stations"
        dsel = np.maximum(ds_km[use_pos], 1.0)
        weights = 1.0 / (dsel ** 2)
        weights = weights / weights.sum()
        sub = use.iloc[ids[use_pos]].copy()
        rec = {
            "County FIPS": str(county["fips5"]).zfill(5),
            f"{metric_label} status": status,
            f"{metric_label} method": method,
            f"{metric_label} stations used": int(len(use_pos)),
            f"{metric_label} nearest station km": float(ds_km[0]),
            f"{metric_label} max station km used": float(ds_km[use_pos].max()),
            f"{metric_label} station IDs used": ";".join(sub["GHCN_ID"].astype(str).tolist()),
            f"{metric_label} station names used": ";".join(sub["station_name"].astype(str).tolist()),
            f"{metric_label} station weights": ";".join(f"{w:.4f}" for w in weights),
        }
        for col in metric_cols:
            rec[col] = float(np.sum(sub[col].to_numpy(dtype=float) * weights))
        rows.append(rec)
    return pd.DataFrame(rows), {f"{metric_label}_station_count": len(use), f"{metric_label}_radius_km": radius_km, f"{metric_label}_k": k}


def quantile_map_to_old_scale(raw_values: Dict[Tuple[str, str], float], old_values: Dict[Tuple[str, str], float]) -> Dict[Tuple[str, str], float]:
    """Map raw proxy ordering onto old value distribution to preserve scale."""
    keys = [k for k, v in raw_values.items() if v is not None and math.isfinite(float(v)) and k in old_values]
    if not keys:
        return {}
    raw_sorted = sorted(keys, key=lambda k: (float(raw_values[k]), k[1], k[0]))
    old_sorted_vals = sorted(float(old_values[k]) for k in keys if old_values[k] is not None and math.isfinite(float(old_values[k])))
    n = min(len(raw_sorted), len(old_sorted_vals))
    mapped = {}
    for i, k in enumerate(raw_sorted[:n]):
        mapped[k] = float(old_sorted_vals[i])
    return mapped


def calibrated_fallbacks(value_rows, measured_by_key, old_col):
    ratios_by_state = defaultdict(list)
    ratios_by_division = defaultdict(list)
    national_ratios = []
    national_vals = []
    for r in value_rows:
        key = (r["County"], r["State"])
        old = safe_float(r[old_col])
        if key not in measured_by_key or old is None or old < 0:
            continue
        new = float(measured_by_key[key])
        if old == 0:
            # Additive rather than ratio calibration for zero old values.
            continue
        ratio = new / old
        state = str(r["State"]).upper()
        division = STATE_DIVISION.get(state, "Unknown")
        ratios_by_state[state].append(ratio)
        ratios_by_division[division].append(ratio)
        national_ratios.append(ratio)
        national_vals.append(new)
    national_ratio = median_or_none(national_ratios) or 1.0
    national_median = median_or_none(national_vals) or 0.0
    state_ratio = {s: median_or_none(v) for s, v in ratios_by_state.items()}
    division_ratio = {d: median_or_none(v) for d, v in ratios_by_division.items()}
    state_n = {s: len(v) for s, v in ratios_by_state.items()}
    division_n = {d: len(v) for d, v in ratios_by_division.items()}
    fallbacks = {}
    for r in value_rows:
        key = (r["County"], r["State"])
        if key in measured_by_key:
            continue
        old = safe_float(r[old_col])
        state = str(r["State"]).upper()
        division = STATE_DIVISION.get(state, "Unknown")
        sr = state_ratio.get(state)
        dr = division_ratio.get(division)
        sn = state_n.get(state, 0)
        dn = division_n.get(division, 0)
        if sr is not None and sn >= 5:
            ratio = sr
            method = f"state calibrated old estimate; state_n={sn}"
        elif sr is not None and dr is not None:
            w_state = sn / (sn + 5.0)
            w_div = 0.50 * (1.0 - w_state)
            w_nat = 1.0 - w_state - w_div
            ratio = w_state * sr + w_div * dr + w_nat * national_ratio
            method = f"blended state/division/national calibrated old estimate; state_n={sn}; division_n={dn}"
        elif dr is not None:
            ratio = 0.70 * dr + 0.30 * national_ratio
            method = f"division/national calibrated old estimate; division_n={dn}"
        else:
            ratio = national_ratio
            method = "national calibrated old estimate"
        if old is None:
            imputed = national_median
        elif old == 0:
            imputed = 0.0
        else:
            imputed = max(0.0, old * ratio)
        fallbacks[key] = (imputed, method)
    return fallbacks


def update_workbook(
    input_workbook: str,
    noaa_tar: str,
    output_workbook: str,
    county_rollup_csv: str | None,
    epa_walkability_zip: str | None,
    county_centroids_csv: str | None,
    station_weight_csv: str | None,
    iterations: int,
    seed: int,
):
    wb = load_workbook(input_workbook)
    if CalcProperties is not None:
        try:
            wb.calculation = CalcProperties(calcMode="auto")
        except Exception:
            pass

    rank_headers, rank_rows = read_sheet(wb[RANK_MATRIX_SHEET])
    value_headers, value_rows = read_sheet(wb[VALUE_MATRIX_SHEET])
    for col in [DROUGHT_VALUE_COL, SNOW_VALUE_COL]:
        if col not in value_headers:
            raise ValueError(f"Missing value column: {col}")
    for col in [DROUGHT_RANK_COL, SNOW_RANK_COL]:
        if col not in rank_headers:
            raise ValueError(f"Missing rank column: {col}")

    fips_lookup = get_fips_for_workbook(wb)
    for r in value_rows:
        r["_fips5"] = fips_lookup.get((r["County"], r["State"]))

    original = {}
    for r in rank_rows:
        key = (r["County"], r["State"])
        original[key] = {
            "Original runoff rank": r["Runoff rank"],
            "Original avg rank": r["Avg Rank Python Check"],
            "Original drought days rank": r[DROUGHT_RANK_COL],
            "Original snowfall rank": r[SNOW_RANK_COL],
            "Original runoff avg elimination round": r["Runoff avg elimination round"],
            "Original runoff wins": r["Runoff wins"],
            "Original runoff win rate": r["Runoff win rate"],
        }
    old_values_by_key = {
        (r["County"], r["State"]): {
            "Old modeled drought days": safe_float(r[DROUGHT_VALUE_COL]),
            "Old modeled snowfall ft": safe_float(r[SNOW_VALUE_COL]),
        }
        for r in value_rows
    }

    stations, station_stats = parse_noaa_multivariate_station_metrics(noaa_tar)
    if station_weight_csv and Path(station_weight_csv).exists():
        rollup, reuse_stats = build_county_rollup_from_temperature_station_weights(stations, station_weight_csv)
    else:
        unique_fips = sorted({r["_fips5"] for r in value_rows if r.get("_fips5")})
        centroids = compute_county_centroids(unique_fips, epa_walkability_zip, county_centroids_csv)
        prcp_metric_cols = ["ann_prcp_in", "wet_days_ge001", "dry_days_ge001", "drought_proxy_raw"]
        if "jja_prcp_in" in stations.columns:
            prcp_metric_cols.append("jja_prcp_in")
        prcp_rollup, prcp_stats = interpolate_metric_to_counties(
            stations, centroids, prcp_metric_cols, "precip", k=5, radius_km=150.0, min_nearby=2
        )
        snow_rollup, snow_stats = interpolate_metric_to_counties(
            stations, centroids, ["ann_snow_in"], "snow", k=5, radius_km=175.0, min_nearby=2
        )
        rollup = prcp_rollup.merge(snow_rollup, on="County FIPS", how="outer")
        reuse_stats = {}
    rollup_by_fips = {str(r["County FIPS"]).zfill(5): r for r in rollup.to_dict("records")}

    # Build measured values by workbook key.
    raw_drought_by_key = {}
    snow_by_key = {}
    update_info = {}
    for r in value_rows:
        key = (r["County"], r["State"])
        fips = r.get("_fips5")
        rec = rollup_by_fips.get(str(fips).zfill(5)) if fips else None
        info = dict(rec) if rec else {}
        if rec and safe_float(rec.get("drought_proxy_raw")) is not None:
            raw_drought_by_key[key] = float(rec["drought_proxy_raw"])
        if rec and safe_float(rec.get("ann_snow_in")) is not None:
            # Use NOAA snow interpolation unless the nearest snow station is extremely far.
            nearest = safe_float(rec.get("snow nearest station km"))
            if nearest is None or nearest <= 350:
                snow_by_key[key] = max(0.0, float(rec["ann_snow_in"]) / 12.0)
        update_info[key] = info

    old_drought_values = {k: v["Old modeled drought days"] for k, v in old_values_by_key.items()}
    mapped_drought_by_key = quantile_map_to_old_scale(raw_drought_by_key, old_drought_values)

    # Fallbacks retained for future cases where station coverage is incomplete.
    drought_fallbacks = calibrated_fallbacks(value_rows, mapped_drought_by_key, DROUGHT_VALUE_COL)
    snow_fallbacks = calibrated_fallbacks(value_rows, snow_by_key, SNOW_VALUE_COL)

    for r in value_rows:
        key = (r["County"], r["State"])
        info = update_info[key]
        if key in mapped_drought_by_key:
            drought_value = float(mapped_drought_by_key[key])
            drought_status = info.get("precip status") or "NOAA precip station blend"
            drought_method = "NOAA precipitation-normal drought proxy, quantile-mapped to prior D1+ drought-day scale"
        else:
            drought_value, fallback_method = drought_fallbacks.get(key, (safe_float(r[DROUGHT_VALUE_COL]) or 0.0, "unmodified old estimate fallback"))
            drought_status = "calibrated fallback; no usable NOAA precipitation station blend"
            drought_method = fallback_method
        if key in snow_by_key:
            snow_value = float(snow_by_key[key])
            snow_status = info.get("snow status") or "NOAA snow station blend"
            snow_method = info.get("snow method") or "inverse-distance-squared snow station blend"
        else:
            snow_value, fallback_method = snow_fallbacks.get(key, (safe_float(r[SNOW_VALUE_COL]) or 0.0, "unmodified old estimate fallback"))
            snow_status = "calibrated fallback; no usable NOAA snow station blend"
            snow_method = fallback_method

        r[DROUGHT_VALUE_COL] = float(max(0.0, drought_value))
        r[SNOW_VALUE_COL] = float(max(0.0, snow_value))
        info["Drought update status"] = drought_status
        info["Drought imputation method"] = drought_method
        info["New NOAA drought-proxy days/year"] = r[DROUGHT_VALUE_COL]
        info["Snow update status"] = snow_status
        info["Snow imputation method"] = snow_method
        info["New NOAA snowfall ft"] = r[SNOW_VALUE_COL]
        update_info[key] = info

    # Re-rank both factors, lower is better.
    drought_ranks = rank_average([r[DROUGHT_VALUE_COL] for r in value_rows], higher_is_better=False)
    snow_ranks = rank_average([r[SNOW_VALUE_COL] for r in value_rows], higher_is_better=False)
    ranks_by_key = {}
    for idx, r in enumerate(value_rows):
        ranks_by_key[(r["County"], r["State"])] = (drought_ranks[idx], snow_ranks[idx])
    for r in rank_rows:
        key = (r["County"], r["State"])
        r[DROUGHT_RANK_COL], r[SNOW_RANK_COL] = ranks_by_key[key]

    factor_cols = [
        c for c in rank_headers
        if isinstance(c, str) and c.endswith(" rank") and c not in {"Avg-based rank", "Runoff rank"}
    ]
    if len(factor_cols) != 21:
        raise ValueError(f"Expected 21 direct rank factor columns, found {len(factor_cols)}")

    factor_array = np.array([[float(r[c]) for c in factor_cols] for r in rank_rows], dtype=np.float64)
    avg_ranks = factor_array.mean(axis=1)
    avg_based_ranks = rank_average(avg_ranks, higher_is_better=False)
    sorted_order = np.argsort(-factor_array, axis=0).T.astype(np.int64)
    wins, avg_elim = run_randomized_runoff(factor_array, sorted_order, iterations, seed)

    sort_rows = []
    for i, r in enumerate(rank_rows):
        sort_rows.append((i, avg_elim[i], wins[i], avg_ranks[i], str(r["State"]), str(r["County"])))
    sort_rows.sort(key=lambda t: (-t[1], -t[2], t[3], t[4], t[5]))
    runoff_rank = [0] * len(rank_rows)
    for pos, (i, *_rest) in enumerate(sort_rows, 1):
        runoff_rank[i] = pos

    for i, r in enumerate(rank_rows):
        r["Avg Rank"] = round(float(avg_ranks[i]), 6)
        r["Avg Rank Python Check"] = round(float(avg_ranks[i]), 6)
        r["Avg Check Delta"] = 0.0
        r["Avg-based rank"] = float(avg_based_ranks[i])
        r["Runoff avg elimination round"] = round(float(avg_elim[i]), 6)
        r["Runoff wins"] = int(wins[i])
        r["Runoff win rate"] = float(wins[i]) / float(iterations)
        r["Runoff rank"] = int(runoff_rank[i])

    rank_by_key = {(r["County"], r["State"]): r for r in rank_rows}
    audit_rows = []
    for vr in value_rows:
        key = (vr["County"], vr["State"])
        old = original[key]
        new = rank_by_key[key]
        info = update_info[key]
        old_vals = old_values_by_key[key]
        annual_prcp = safe_float(info.get("ann_prcp_in"))
        wet_days = safe_float(info.get("wet_days_ge001"))
        dry_days = safe_float(info.get("dry_days_ge001"))
        jja_prcp = safe_float(info.get("jja_prcp_in"))
        audit_rows.append({
            "County": key[0],
            "State": key[1],
            "County FIPS": info.get("County FIPS") or vr.get("_fips5"),
            "Drought update status": info.get("Drought update status"),
            "Old modeled drought days": old_vals["Old modeled drought days"],
            "NOAA annual precip inches": annual_prcp,
            "NOAA wet days >=0.01in": wet_days,
            "NOAA dry days <0.01in": dry_days,
            "NOAA JJA precip inches": jja_prcp,
            "NOAA drought proxy raw": safe_float(info.get("drought_proxy_raw")),
            "New NOAA drought-proxy days/year": vr[DROUGHT_VALUE_COL],
            "Original drought rank": old["Original drought days rank"],
            "New drought rank": new[DROUGHT_RANK_COL],
            "Drought rank delta": new[DROUGHT_RANK_COL] - old["Original drought days rank"],
            "Drought imputation method": info.get("Drought imputation method"),
            "Precip nearest station km": safe_float(info.get("precip nearest station km")),
            "Precip max station km used": safe_float(info.get("precip max station km used")),
            "Precip station names used": info.get("precip station names used"),
            "Snow update status": info.get("Snow update status"),
            "Old modeled snowfall ft": old_vals["Old modeled snowfall ft"],
            "New NOAA snowfall ft": vr[SNOW_VALUE_COL],
            "Original snowfall rank": old["Original snowfall rank"],
            "New snowfall rank": new[SNOW_RANK_COL],
            "Snowfall rank delta": new[SNOW_RANK_COL] - old["Original snowfall rank"],
            "Snow imputation method": info.get("Snow imputation method"),
            "Snow nearest station km": safe_float(info.get("snow nearest station km")),
            "Snow max station km used": safe_float(info.get("snow max station km used")),
            "Snow station names used": info.get("snow station names used"),
            "Original avg rank": old["Original avg rank"],
            "New avg rank": new["Avg Rank Python Check"],
            "Avg rank delta": new["Avg Rank Python Check"] - old["Original avg rank"],
            "Original runoff rank": old["Original runoff rank"],
            "New runoff rank": new["Runoff rank"],
            "Runoff rank delta": new["Runoff rank"] - old["Original runoff rank"],
            "Original runoff avg elimination round": old["Original runoff avg elimination round"],
            "New runoff avg elimination round": new["Runoff avg elimination round"],
            "Runoff avg elimination delta": new["Runoff avg elimination round"] - old["Original runoff avg elimination round"],
            "Original runoff wins": old["Original runoff wins"],
            "New runoff wins": new["Runoff wins"],
            "Original runoff win rate": old["Original runoff win rate"],
            "New runoff win rate": new["Runoff win rate"],
        })

    # Sort outputs by new runoff rank.
    rank_rows_sorted = sorted(rank_rows, key=lambda r: r["Runoff rank"])
    value_rows_sorted = sorted(value_rows, key=lambda r: rank_by_key[(r["County"], r["State"])] ["Runoff rank"])
    audit_sorted = sorted(audit_rows, key=lambda r: r["New runoff rank"])

    top_old20 = {k for k, v in original.items() if int(v["Original runoff rank"]) <= 20}
    top_new20 = {(r["County"], r["State"]) for r in rank_rows if int(r["Runoff rank"]) <= 20}
    top_old30 = {k for k, v in original.items() if int(v["Original runoff rank"]) <= 30}
    top_new30 = {(r["County"], r["State"]) for r in rank_rows if int(r["Runoff rank"]) <= 30}
    drought_station_backed = sum(1 for r in audit_rows if str(r["Drought update status"]).startswith("NOAA"))
    drought_fallback_count = len(audit_rows) - drought_station_backed
    snow_station_backed = sum(1 for r in audit_rows if str(r["Snow update status"]).startswith("NOAA"))
    snow_fallback_count = len(audit_rows) - snow_station_backed

    precip_nearest_vals = [float(r["Precip nearest station km"]) for r in audit_rows if r.get("Precip nearest station km") is not None]
    snow_nearest_vals = [float(r["Snow nearest station km"]) for r in audit_rows if r.get("Snow nearest station km") is not None]

    impact_rows = [
        ["NOAA station CSV files read", station_stats["station_csv_files_read"]],
        ["NOAA station rows with coordinates", station_stats["station_rows_with_coordinates"]],
        ["NOAA stations with precipitation normal", station_stats["stations_with_precip_normal"]],
        ["NOAA stations with drought proxy", station_stats["stations_with_drought_proxy"]],
        ["NOAA stations with snow normal", station_stats["stations_with_snow_normal"]],
        ["Workbook county rows", len(rank_rows)],
        ["Drought rows station-backed", drought_station_backed],
        ["Drought rows fallback", drought_fallback_count],
        ["Snowfall rows station-backed", snow_station_backed],
        ["Snowfall rows fallback", snow_fallback_count],
        ["Precip nearest station km min", min(precip_nearest_vals) if precip_nearest_vals else None],
        ["Precip nearest station km max", max(precip_nearest_vals) if precip_nearest_vals else None],
        ["Snow nearest station km min", min(snow_nearest_vals) if snow_nearest_vals else None],
        ["Snow nearest station km max", max(snow_nearest_vals) if snow_nearest_vals else None],
        ["Top-20 runoff overlap", f"{len(top_old20 & top_new20)} / 20"],
        ["Top-30 runoff overlap", f"{len(top_old30 & top_new30)} / 30"],
        ["Mean abs drought-rank change", mean_abs(audit_rows, "Drought rank delta")],
        ["Median abs drought-rank change", median_abs(audit_rows, "Drought rank delta")],
        ["Mean abs snowfall-rank change", mean_abs(audit_rows, "Snowfall rank delta")],
        ["Median abs snowfall-rank change", median_abs(audit_rows, "Snowfall rank delta")],
        ["Mean abs Avg Rank change", mean_abs(audit_rows, "Avg rank delta")],
        ["Median abs Avg Rank change", median_abs(audit_rows, "Avg rank delta")],
        ["Mean abs runoff-rank change", mean_abs(audit_rows, "Runoff rank delta")],
        ["Median abs runoff-rank change", median_abs(audit_rows, "Runoff rank delta")],
        ["Largest abs runoff-rank move", max(abs(float(r["Runoff rank delta"])) for r in audit_rows)],
    ]

    # Write main sheets.
    def rows_to_matrix(headers, rows):
        return [headers] + [[clean_cell(r.get(h)) for h in headers] for r in rows]

    write_matrix(wb[RANK_MATRIX_SHEET], rows_to_matrix(rank_headers, rank_rows_sorted))
    write_matrix(wb[VALUE_MATRIX_SHEET], rows_to_matrix(value_headers, value_rows_sorted))
    write_matrix(wb[TOP30_SHEET], rows_to_matrix(rank_headers, rank_rows_sorted[:30]))

    comp_headers = ["New top-30 position", "County", "State", "Original runoff rank", "New runoff rank", "Runoff rank delta",
                    "Original runoff avg elimination round", "New runoff avg elimination round", "Original runoff wins", "New runoff wins", "Original avg rank", "New avg rank", "Avg rank delta"]
    comp_rows = []
    for pos, r in enumerate(rank_rows_sorted[:30], 1):
        key = (r["County"], r["State"])
        old = original[key]
        comp_rows.append([pos, r["County"], r["State"], old["Original runoff rank"], r["Runoff rank"], r["Runoff rank"] - old["Original runoff rank"],
                          old["Original runoff avg elimination round"], r["Runoff avg elimination round"],
                          old["Original runoff wins"], r["Runoff wins"], old["Original avg rank"], r["Avg Rank Python Check"], r["Avg Rank Python Check"] - old["Original avg rank"]])
    write_matrix(wb[RESULT_COMPARISON_SHEET], [comp_headers] + comp_rows)

    ws_impact = recreate_sheet(wb, IMPACT_SHEET, after_name=RESULT_COMPARISON_SHEET)
    write_matrix(ws_impact, [["Metric", "Value"]] + impact_rows)
    ws_impact.cell(1, 4, "Top runoff movers")
    mover_headers = ["County", "State", "Original runoff rank", "New runoff rank", "Runoff rank delta",
                     "Original avg rank", "New avg rank", "Avg rank delta",
                     "Old modeled drought days", "New NOAA drought-proxy days/year", "Drought rank delta",
                     "Old modeled snowfall ft", "New NOAA snowfall ft", "Snowfall rank delta",
                     "NOAA annual precip inches", "NOAA dry days <0.01in", "Snow nearest station km"]
    movers = sorted(audit_rows, key=lambda r: abs(float(r["Runoff rank delta"])), reverse=True)[:15]
    for c, h in enumerate(mover_headers, 4):
        ws_impact.cell(2, c, h)
    for r_idx, row in enumerate(movers, 3):
        for c_idx, h in enumerate(mover_headers, 4):
            ws_impact.cell(r_idx, c_idx, clean_cell(row.get(h)))

    ws_audit = recreate_sheet(wb, UPDATE_SHEET, after_name=IMPACT_SHEET)
    audit_headers = list(audit_sorted[0].keys())
    write_matrix(ws_audit, [audit_headers] + [[clean_cell(r[h]) for h in audit_headers] for r in audit_sorted])

    run_rows = [
        ["Parameter", "Value", "Notes"],
        ["Iterations", iterations, "100k randomized runoff simulations rerun after NOAA annual/seasonal multivariate update."],
        ["Random seed", seed, "Same seed retained for comparability; deterministic with this script/version."],
        ["Rows/counties", len(rank_rows), "No counties added or removed."],
        ["Rank factors used", len(factor_cols), "Avg Rank excluded from random factor selection."],
        ["Column-selection rule", "Random permutation cycle", "All rank columns are picked exactly once per cycle before reshuffling."],
        ["Tie handling", "Random among tied worst ranks", "If multiple active counties tie for worst rank in selected factor."],
        ["Data import outcome", f"Drought: {drought_station_backed} station-backed / {drought_fallback_count} fallback; Snow: {snow_station_backed} station-backed / {snow_fallback_count} fallback", "NOAA station normals interpolated to county population-weighted centroids."],
        ["Factor values changed?", "Drought days and annual snowfall", "NOAA precipitation normals inform drought proxy; NOAA snow normals directly update snowfall."],
        ["Drought method", "NOAA precipitation-normal drought proxy", "Dry days, annual precipitation scarcity, and summer precipitation scarcity; quantile-mapped to prior D1+ drought-day scale."],
        ["Snowfall method", "NOAA ANN-SNOW-NORMAL station interpolation", "Annual snow normal in inches converted to feet after county station blend."],
        ["Missing-data policy", "Spatial station interpolation first; calibrated old-estimate fallback if no usable station blend", "No fallback is expected for drought; snow fallback may occur in sparse snow-normal coverage."],
        [None, None, None],
        ["Runoff rank", "County", "State", "Runoff avg elimination round", "Runoff wins", "Runoff win rate", "Avg Rank Python Check"],
    ]
    for r in rank_rows_sorted[:30]:
        run_rows.append([r["Runoff rank"], r["County"], r["State"], r["Runoff avg elimination round"], r["Runoff wins"], r["Runoff win rate"], r["Avg Rank Python Check"]])
    write_matrix(wb[RUNOFF_SHEET], run_rows)

    if METHODOLOGY_SHEET in wb.sheetnames:
        ws = wb[METHODOLOGY_SHEET]
        for row in ws.iter_rows(min_row=2):
            value_col = str(row[1].value).strip() if row[1].value is not None else ""
            if value_col == DROUGHT_VALUE_COL:
                row[4].value = "Updated using NOAA/NCEI annual/seasonal precipitation normals as drought-risk proxy"
                row[5].value = ("NOAA multivariate normals do not directly observe D1+ drought frequency. Updated as a real-data proxy using annual dry days derived from ANN-PRCP-AVGNDS-GE001HI, annual precipitation scarcity from ANN-PRCP-NORMAL, and summer precipitation scarcity from JJA-PRCP-NORMAL. County values use inverse-distance-squared nearest-station blending, then quantile-map proxy ordering to the prior D1+ drought-day scale. Lower value ranks better. Future preferred source remains U.S. Drought Monitor county time series.")
                row[6].value = "https://www.ncei.noaa.gov/products/land-based-station/us-climate-normals ; uploaded us-climate-normals_2006-2020_v1.0.1_annualseasonal_multivariate_by-station_c20230404.tar.gz ; https://droughtmonitor.unl.edu/DmData/DataDownload.aspx"
            elif value_col == SNOW_VALUE_COL:
                row[4].value = "Imported NOAA/NCEI annual/seasonal multivariate snow normals"
                row[5].value = ("Updated from uploaded NOAA annual/seasonal multivariate normals. Annual snowfall uses ANN-SNOW-NORMAL station normals, inverse-distance-squared blended to population-weighted county centroids and converted from inches to feet. Lower value ranks better. Missing/unreliable snow station coverage falls back to calibrated old estimates.")
                row[6].value = "https://www.ncei.noaa.gov/products/land-based-station/us-climate-normals ; uploaded us-climate-normals_2006-2020_v1.0.1_annualseasonal_multivariate_by-station_c20230404.tar.gz"

    if LOG_SHEET in wb.sheetnames:
        ws = wb[LOG_SHEET]
        next_row = ws.max_row + 1
        log_values = [
            "NOAA/NCEI annual-seasonal multivariate normals 2006-2020",
            "Annual snowfall; drought-days proxy",
            "Uploaded us-climate-normals_2006-2020_v1.0.1_annualseasonal_multivariate_by-station_c20230404.tar.gz",
            "Imported station annual/seasonal normals; spatially blended snow/precip stations to county population-weighted centroids; quantile-mapped drought proxy to prior scale",
            f"Drought {drought_station_backed} station-backed / {drought_fallback_count} fallback; snow {snow_station_backed} station-backed / {snow_fallback_count} fallback; reran 100k MCMC",
            "https://www.ncei.noaa.gov/products/land-based-station/us-climate-normals",
        ]
        for c, val in enumerate(log_values, 1):
            ws.cell(next_row, c, val)

    if QA_SHEET in wb.sheetnames:
        qa_rows = [
            ["Check", "Value", "Result / Notes"],
            ["County rows", len(rank_rows), "Expected 413 based on input workbook."],
            ["Direct rank factor count", len(factor_cols), "Direct factor columns ending in ' rank', excluding Runoff rank and Avg-based rank."],
            ["NOAA station CSV files read", station_stats["station_csv_files_read"], "Annual/seasonal multivariate station files."],
            ["NOAA stations with drought proxy", station_stats["stations_with_drought_proxy"], "Stations with precipitation normal and wet-day normal."],
            ["NOAA stations with snow normal", station_stats["stations_with_snow_normal"], "Stations with ANN-SNOW-NORMAL."],
            ["Drought rows station-backed", drought_station_backed, "County rows assigned NOAA precipitation proxy values."],
            ["Drought rows calibrated fallback", drought_fallback_count, "No usable precipitation station blend."],
            ["Snow rows station-backed", snow_station_backed, "County rows assigned NOAA snow-normal values."],
            ["Snow rows calibrated fallback", snow_fallback_count, "No usable snow station blend or snow stations too far away."],
            ["Runoff wins total", int(sum(r["Runoff wins"] for r in rank_rows)), f"Should equal iterations ({iterations})."],
            ["Runoff win rate sum", float(sum(r["Runoff win rate"] for r in rank_rows)), "Should equal 1.0, subject to floating-point precision."],
            ["Changed factors", "Drought days; annual snowfall", "Other value/rank columns preserved from source workbook."],
        ]
        write_matrix(wb[QA_SHEET], qa_rows)

    for sheet_name in [RANK_MATRIX_SHEET, VALUE_MATRIX_SHEET, TOP30_SHEET, RESULT_COMPARISON_SHEET, RUNOFF_SHEET, QA_SHEET, "Sources Not Imported", IMPACT_SHEET, UPDATE_SHEET, METHODOLOGY_SHEET, LOG_SHEET]:
        if sheet_name in wb.sheetnames:
            style_sheet(wb[sheet_name])

    Path(output_workbook).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_workbook)
    if county_rollup_csv:
        pd.DataFrame(audit_sorted).to_csv(county_rollup_csv, index=False)

    return {
        "drought_station_backed": drought_station_backed,
        "drought_fallback_count": drought_fallback_count,
        "snow_station_backed": snow_station_backed,
        "snow_fallback_count": snow_fallback_count,
        "impact_rows": impact_rows,
        "top5": [(r["Runoff rank"], r["County"], r["State"], r["Avg Rank Python Check"], r["Runoff wins"]) for r in rank_rows_sorted[:5]],
        "movers": movers,
        "audit_rows": audit_rows,
    }


def parse_args():
    parser = argparse.ArgumentParser(description="Update county ranking workbook with NOAA annual/seasonal multivariate normals and rerun MCMC.")
    parser.add_argument("--input-workbook", required=True)
    parser.add_argument("--noaa-tar", required=True)
    parser.add_argument("--output-workbook", required=True)
    parser.add_argument("--county-rollup-csv", default=None)
    parser.add_argument("--epa-walkability-zip", default=None)
    parser.add_argument("--county-centroids-csv", default=None)
    parser.add_argument("--station-weight-csv", default=None, help="Prior NOAA temperature rollup CSV with Station IDs used and Station weights. Fast path; avoids rereading EPA geometry.")
    parser.add_argument("--iterations", type=int, default=DEFAULT_ITERATIONS)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    return parser.parse_args()


def main():
    args = parse_args()
    result = update_workbook(
        input_workbook=args.input_workbook,
        noaa_tar=args.noaa_tar,
        output_workbook=args.output_workbook,
        county_rollup_csv=args.county_rollup_csv,
        epa_walkability_zip=args.epa_walkability_zip,
        county_centroids_csv=args.county_centroids_csv,
        station_weight_csv=args.station_weight_csv,
        iterations=args.iterations,
        seed=args.seed,
    )
    print(f"Saved updated workbook: {args.output_workbook}")
    if args.county_rollup_csv:
        print(f"Saved county rollup CSV: {args.county_rollup_csv}")
    print(f"Drought station-backed county rows: {result['drought_station_backed']}")
    print(f"Drought fallback rows: {result['drought_fallback_count']}")
    print(f"Snow station-backed county rows: {result['snow_station_backed']}")
    print(f"Snow fallback rows: {result['snow_fallback_count']}")
    print("Top 5 after update:")
    for row in result["top5"]:
        print(row)


if __name__ == "__main__":
    main()
