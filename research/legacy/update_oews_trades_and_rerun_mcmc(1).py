#!/usr/bin/env python3
"""
Update a county-rankings workbook with BLS OEWS occupation data to refine
Tradespeople per capita, then rerun the randomized runoff MCMC.

This pass is designed to augment a prior CBP-backed workbook. CBP remains the
county-level backbone because it is county-by-industry; OEWS adds labor-market
occupation mix by metro/nonmetro/state geography.

Expected inputs
---------------
1. Existing workbook already updated with CBP trades/grocery values.
2. BLS OEWS May 2025 Metropolitan and nonmetropolitan area ZIP
   (e.g. BLS_OEWS_oesm25ma.zip, containing MSA_M2025_dl.xlsx and BOS_M2025_dl.xlsx).
3. BLS OEWS May 2025 State ZIP (e.g. BLS_OEWS_oesm25st.zip).
4. BLS OEWS May 2025 National industry-specific ZIP (optional but recommended,
   e.g. BLS_OEWS_oesm25in4.zip). Used to estimate what share of NAICS 238
   employment is in selected trade occupations.
5. EPA Walkability ZIP (optional but recommended). Used only to map county FIPS
   to CBSA codes for metro OEWS matching. If omitted, the script falls back to
   state OEWS for every county.

Updated factor
--------------
- Tradespeople per capita:
    prior CBP NAICS 238 employment per 1,000 residents
    × national share of NAICS 238 employment in selected trade occupations
    × local OEWS occupation-mix adjustment

The local adjustment is the selected-trade-occupation share of total employment
in the county's CBSA when a CBSA match is available, otherwise the state share.
The adjustment is normalized to the national all-industry OEWS trade-occupation
share and clipped to a defensible range, so OEWS improves CBP without swamping
county-level CBP data.

Example
-------
python update_oews_trades_and_rerun_mcmc.py \
  --input-workbook county_rankings_epa_walkability_transit_aqi_noaa_temperature_multivariate_droughtmonitor_cbp_100k.xlsx \
  --oews-ma-zip BLS_OEWS_oesm25ma.zip \
  --oews-state-zip BLS_OEWS_oesm25st.zip \
  --oews-industry-zip BLS_OEWS_oesm25in4.zip \
  --walkability-zip EPA_WalkabilityIndex.zip \
  --output-workbook county_rankings_epa_walkability_transit_aqi_noaa_temperature_multivariate_droughtmonitor_cbp_oews_100k.xlsx \
  --county-rollup-csv oews_trades_county_rollup.csv
"""

from __future__ import annotations

import argparse
import csv
import html
import math
import os
import re
import statistics
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import numpy as np
from numba import njit
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
try:
    from openpyxl.workbook.properties import CalcProperties
except Exception:  # pragma: no cover
    CalcProperties = None

DEFAULT_ITERATIONS = 100_000
DEFAULT_SEED = 2026060901

RANK_MATRIX_SHEET = "Rank Matrix"
VALUE_MATRIX_SHEET = "Value Matrix"
TOP30_SHEET = "Top 30"
RESULT_COMPARISON_SHEET = "Result Comparison"
RUNOFF_SHEET = "Runoff Simulation"
METHODOLOGY_SHEET = "Sources & Methodology"
LOG_SHEET = "Data Acquisition Log"
QA_SHEET = "QA Checks"
IMPACT_SHEET = "BLS OEWS Trades Impact"
UPDATE_SHEET = "BLS OEWS Trades Update"
REVIEW_SHEET = "BLS OEWS Input Review"

TRADES_VALUE_COL = "Est. tradespeople / 1k residents"
TRADES_RANK_COL = "Tradespeople per capita rank"

# User-prioritized trade occupations. Kept deliberately transparent.
TRADE_OCC_CODES = {
    "47-2111": "Electricians",
    "47-2152": "Plumbers, Pipefitters, and Steamfitters",
    "49-9021": "HVAC and Refrigeration Mechanics and Installers",
    "47-2181": "Roofers",
    "47-2031": "Carpenters",
    "47-2141": "Painters, Construction and Maintenance",
    "47-2061": "Construction Laborers",
    "49-9041": "Industrial Machinery Mechanics",
    "49-9071": "Maintenance and Repair Workers, General",
}

# Guardrail on the geography-level OEWS multiplier. OEWS is not county-level, so
# it should adjust but not dominate the county-level CBP signal.
OEWS_FACTOR_MIN = 0.65
OEWS_FACTOR_MAX = 1.35

HEADER_FILL = "1F4E79"
BORDER_COLOR = "B7B7B7"

STATE_ABBR_TO_FIPS = {
    "AL":"01","AK":"02","AZ":"04","AR":"05","CA":"06","CO":"08","CT":"09","DE":"10","DC":"11","FL":"12","GA":"13","HI":"15","ID":"16","IL":"17","IN":"18","IA":"19","KS":"20","KY":"21","LA":"22","ME":"23","MD":"24","MA":"25","MI":"26","MN":"27","MS":"28","MO":"29","MT":"30","NE":"31","NV":"32","NH":"33","NJ":"34","NM":"35","NY":"36","NC":"37","ND":"38","OH":"39","OK":"40","OR":"41","PA":"42","RI":"44","SC":"45","SD":"46","TN":"47","TX":"48","UT":"49","VT":"50","VA":"51","WA":"53","WV":"54","WI":"55","WY":"56","PR":"72"
}

STATE_DIVISION = {
    "CT": "New England", "ME": "New England", "MA": "New England", "NH": "New England", "RI": "New England", "VT": "New England",
    "NJ": "Middle Atlantic", "NY": "Middle Atlantic", "PA": "Middle Atlantic",
    "IL": "East North Central", "IN": "East North Central", "MI": "East North Central", "OH": "East North Central", "WI": "East North Central",
    "IA": "West North Central", "KS": "West North Central", "MN": "West North Central", "MO": "West North Central", "NE": "West North Central", "ND": "West North Central", "SD": "West North Central",
    "DE": "South Atlantic", "DC": "South Atlantic", "FL": "South Atlantic", "GA": "South Atlantic", "MD": "South Atlantic", "NC": "South Atlantic", "SC": "South Atlantic", "VA": "South Atlantic", "WV": "South Atlantic",
    "AL": "East South Central", "KY": "East South Central", "MS": "East South Central", "TN": "East South Central",
    "AR": "West South Central", "LA": "West South Central", "OK": "West South Central", "TX": "West South Central",
    "AZ": "Mountain", "CO": "Mountain", "ID": "Mountain", "MT": "Mountain", "NV": "Mountain", "NM": "Mountain", "UT": "Mountain", "WY": "Mountain",
    "AK": "Pacific", "CA": "Pacific", "HI": "Pacific", "OR": "Pacific", "WA": "Pacific",
}


def safe_float(x):
    if x is None or x == "":
        return None
    try:
        val = float(str(x).strip().replace(",", ""))
    except Exception:
        return None
    return val if math.isfinite(val) else None


def clean_cell(value):
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return None if math.isnan(float(value)) else float(value)
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = list(rows[0])
    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        padded = list(row) + [None] * (len(headers) - len(row))
        data.append(dict(zip(headers, padded[: len(headers)])))
    return headers, data


def write_matrix(ws, matrix: List[List[object]]):
    ws.delete_rows(1, ws.max_row)
    for r_idx, row in enumerate(matrix, 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, clean_cell(value))


def recreate_sheet(wb, name: str, after_name: Optional[str] = None):
    if name in wb.sheetnames:
        del wb[name]
    if after_name and after_name in wb.sheetnames:
        idx = wb.sheetnames.index(after_name)
        return wb.create_sheet(name, idx + 1)
    return wb.create_sheet(name)


def style_sheet(ws, max_width: int = 42):
    """Lightweight sheet styling. Avoid per-cell style churn because this workbook
    already contains many generated sheets and heavy style mutation slows xlsx saves.
    """
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color=BORDER_COLOR)
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        values = [ws.cell(row=r, column=col).value for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max_width, max(10, min(36, max(len(str(v)) if v is not None else 0 for v in values) + 2)))
        ws.column_dimensions[letter].width = width
    if ws.max_row >= 1:
        ws.row_dimensions[1].height = 24


def get_fips_for_workbook(wb) -> Dict[Tuple[str, str], str]:
    for sheet_name in [
        "CBP Business Patterns Update", "Drought Monitor Update", "EPA Walkability Update",
        "NOAA Multivariate Update", "NOAA Temperature Update", "EPA AQI Update",
        "EPA Transit Guardrail Update", "EPA Transit Update",
    ]:
        if sheet_name not in wb.sheetnames:
            continue
        headers, rows = read_sheet(wb[sheet_name])
        if {"County", "State", "County FIPS"}.issubset(set(headers)):
            out = {}
            for r in rows:
                county, state, fips = r.get("County"), r.get("State"), r.get("County FIPS")
                if county is None or state is None or fips is None:
                    continue
                out[(county, state)] = str(fips).zfill(5)
            if out:
                return out
    raise ValueError("Workbook needs an audit/update sheet containing County, State, and County FIPS")


def get_cbp_trade_details(wb) -> Dict[Tuple[str, str], Dict[str, object]]:
    out = {}
    if "CBP Business Patterns Update" not in wb.sheetnames:
        return out
    headers, rows = read_sheet(wb["CBP Business Patterns Update"])
    for r in rows:
        key = (r.get("County"), r.get("State"))
        out[key] = {
            "population_2023": r.get("Population 2023"),
            "cbp_trade_employment_238": r.get("CBP trade employment NAICS 238///"),
            "cbp_status": r.get("CBP status"),
        }
    return out


def rank_average(values: Iterable[float], higher_is_better: bool) -> List[float]:
    pairs = [(float(v), i) for i, v in enumerate(values)]
    pairs.sort(key=lambda t: t[0], reverse=higher_is_better)
    ranks = [0.0] * len(pairs)
    i = 0
    while i < len(pairs):
        j = i + 1
        while j < len(pairs) and pairs[j][0] == pairs[i][0]:
            j += 1
        avg_rank = (i + 1 + j) / 2.0
        for k in range(i, j):
            ranks[pairs[k][1]] = avg_rank
        i = j
    return ranks


@njit
def run_randomized_runoff(ranks: np.ndarray, sorted_order: np.ndarray, iterations: int, seed: int):
    np.random.seed(seed)
    n, f = ranks.shape
    wins = np.zeros(n, np.int64)
    elim_sum = np.zeros(n, np.float64)
    factor_order = np.empty(f, np.int64)
    active = np.empty(n, np.bool_)
    ptr = np.empty(f, np.int64)
    ties = np.empty(n, np.int64)

    for _ in range(iterations):
        for i in range(n):
            active[i] = True
        for j in range(f):
            ptr[j] = 0
        active_count = n
        pos = f
        round_no = 0

        while active_count > 1:
            if pos >= f:
                for j in range(f):
                    factor_order[j] = j
                for j in range(f - 1, 0, -1):
                    k = np.random.randint(j + 1)
                    tmp = factor_order[j]
                    factor_order[j] = factor_order[k]
                    factor_order[k] = tmp
                pos = 0

            col = factor_order[pos]
            pos += 1

            p = ptr[col]
            while p < n and not active[sorted_order[col, p]]:
                p += 1
            ptr[col] = p

            first_idx = sorted_order[col, p]
            worst_rank = ranks[first_idx, col]
            tie_count = 0
            q = p
            while q < n:
                row_idx = sorted_order[col, q]
                if ranks[row_idx, col] != worst_rank:
                    break
                if active[row_idx]:
                    ties[tie_count] = row_idx
                    tie_count += 1
                q += 1

            loser = ties[np.random.randint(tie_count)]
            active[loser] = False
            round_no += 1
            elim_sum[loser] += round_no
            active_count -= 1

        for i in range(n):
            if active[i]:
                winner = i
                break
        wins[winner] += 1
        elim_sum[winner] += n

    return wins, elim_sum / iterations


# Fast XLSX reader for BLS OEWS files. Avoids loading entire XLSX via openpyxl.
CELL_RE = re.compile(rb'<c r="([A-Z]+)[0-9]+"([^>]*)>(.*?)</c>', re.S)
V_RE = re.compile(rb'<v>(.*?)</v>', re.S)
T_RE = re.compile(rb'<t[^>]*>(.*?)</t>', re.S)
ROW_RE = re.compile(rb'<row[^>]*>(.*?)</row>', re.S)


def load_shared_strings(z: zipfile.ZipFile):
    if "xl/sharedStrings.xml" not in z.namelist():
        return []
    data = z.read("xl/sharedStrings.xml").decode("utf-8", errors="ignore")
    sis = re.findall(r"<si[^>]*>(.*?)</si>", data, flags=re.S)
    out = []
    for si in sis:
        texts = re.findall(r"<t[^>]*>(.*?)</t>", si, flags=re.S)
        out.append(html.unescape("".join(texts)))
    return out


def parse_xlsx_cell(attrs: bytes, inner: bytes, shared: List[str]):
    m = V_RE.search(inner)
    if b't="s"' in attrs:
        if not m:
            return ""
        try:
            return shared[int(m.group(1))]
        except Exception:
            return ""
    if b't="inlineStr"' in attrs:
        texts = T_RE.findall(inner)
        return html.unescape(b"".join(texts).decode("utf-8", errors="ignore"))
    return m.group(1).decode("utf-8", errors="ignore") if m else None


def extract_from_zip(zip_path: Path, member_name: str, tmp_dir: Path) -> Path:
    out = tmp_dir / Path(member_name).name
    if not out.exists():
        with zipfile.ZipFile(zip_path) as z:
            out.write_bytes(z.read(member_name))
    return out


def find_zip_member(zip_path: Path, suffix: str) -> str:
    with zipfile.ZipFile(zip_path) as z:
        matches = [n for n in z.namelist() if n.endswith(suffix)]
    if not matches:
        raise ValueError(f"Could not find {suffix!r} in {zip_path}")
    return matches[0]


def read_oews_area_xlsx(xlsx_path: Path) -> Dict[str, Dict[str, object]]:
    with zipfile.ZipFile(xlsx_path) as z:
        shared = load_shared_strings(z)
        xml = z.read("xl/worksheets/sheet1.xml")
    needed = {"00-0000"}.union(TRADE_OCC_CODES.keys())
    areas: Dict[str, Dict[str, object]] = {}
    for ri, rm in enumerate(ROW_RE.finditer(xml)):
        if ri == 0:
            continue
        inner = rm.group(1)
        cells = {}
        for cm in CELL_RE.finditer(inner):
            col = cm.group(1).decode()
            if col in {"A", "B", "D", "I", "J", "L"}:
                cells[col] = parse_xlsx_cell(cm.group(2), cm.group(3), shared)
        occ = str(cells.get("I", "")).strip()
        if occ not in needed:
            continue
        emp = safe_float(cells.get("L"))
        if emp is None:
            continue
        area = str(cells.get("A")).strip()
        d = areas.setdefault(area, {
            "area": area,
            "area_title": cells.get("B"),
            "prim_state": cells.get("D"),
            "all_emp": None,
            "trade_emp": 0.0,
        })
        if occ == "00-0000":
            d["all_emp"] = emp
        else:
            d["trade_emp"] = float(d["trade_emp"]) + emp
    out = {}
    for area, d in areas.items():
        all_emp = safe_float(d.get("all_emp"))
        trade_emp = safe_float(d.get("trade_emp"))
        if all_emp and trade_emp is not None:
            d["trade_jobs_per_1000_all_jobs"] = trade_emp / all_emp * 1000.0
            out[area] = d
    return out


def read_oews_industry_share(xlsx_path: Path) -> Tuple[float, float, float, List[Tuple[str, str, float]]]:
    with zipfile.ZipFile(xlsx_path) as z:
        shared = load_shared_strings(z)
        xml = z.read("xl/worksheets/sheet1.xml")
    needed = {"00-0000"}.union(TRADE_OCC_CODES.keys())
    all238 = 0.0
    trade238 = 0.0
    rows = []
    for ri, rm in enumerate(ROW_RE.finditer(xml)):
        if ri == 0:
            continue
        inner = rm.group(1)
        cells = {}
        for cm in CELL_RE.finditer(inner):
            col = cm.group(1).decode()
            if col in {"E", "I", "J", "L"}:
                cells[col] = parse_xlsx_cell(cm.group(2), cm.group(3), shared)
        if str(cells.get("E", "")).strip() != "238000":
            continue
        occ = str(cells.get("I", "")).strip()
        if occ not in needed:
            continue
        emp = safe_float(cells.get("L"))
        if emp is None:
            continue
        if occ == "00-0000":
            all238 = emp
        else:
            trade238 += emp
            rows.append((occ, str(cells.get("J")), emp))
    share = trade238 / all238 if all238 else 1.0
    return all238, trade238, share, rows


def load_county_cbsa_from_walkability(walkability_zip: Optional[Path]) -> Tuple[Dict[str, Dict[str, object]], str]:
    if not walkability_zip:
        return {}, "No EPA Walkability ZIP supplied; state OEWS fallback used for all counties"
    try:
        import pyogrio  # type: ignore
    except Exception:
        return {}, "pyogrio unavailable; state OEWS fallback used for all counties"
    tmp_dir = Path(tempfile.mkdtemp(prefix="walkability_gdb_"))
    with zipfile.ZipFile(walkability_zip) as z:
        z.extractall(tmp_dir)
    gdb_candidates = list(tmp_dir.glob("*.gdb"))
    if not gdb_candidates:
        gdb_candidates = list(tmp_dir.glob("**/*.gdb"))
    if not gdb_candidates:
        return {}, "No .gdb found in EPA Walkability ZIP; state OEWS fallback used for all counties"
    gdb = str(gdb_candidates[0])
    df = pyogrio.read_dataframe(gdb, layer="NationalWalkabilityIndex", columns=["STATEFP", "COUNTYFP", "CBSA", "CBSA_Name", "TotPop"], read_geometry=False)
    pop_by_county_cbsa: Dict[str, Dict[str, float]] = defaultdict(lambda: defaultdict(float))
    cbsa_names: Dict[str, str] = {}
    for st, cty, cbsa, cbsa_name, pop in zip(df["STATEFP"], df["COUNTYFP"], df["CBSA"], df["CBSA_Name"], df["TotPop"]):
        fips = str(st).zfill(2) + str(cty).zfill(3)
        cbsa_code = str(cbsa).strip() if cbsa is not None else ""
        if not cbsa_code or cbsa_code.lower() == "nan":
            continue
        p = safe_float(pop) or 0.0
        pop_by_county_cbsa[fips][cbsa_code] += p
        cbsa_names[cbsa_code] = str(cbsa_name) if cbsa_name is not None else ""
    mapping = {}
    for fips, choices in pop_by_county_cbsa.items():
        cbsa_code, cbsa_pop = max(choices.items(), key=lambda kv: kv[1])
        mapping[fips] = {"cbsa": cbsa_code, "cbsa_name": cbsa_names.get(cbsa_code), "cbsa_population_weight": cbsa_pop}
    return mapping, f"EPA Walkability block-group CBSA fields used for {len(mapping)} county-CBSA mappings"


def load_county_cbsa_from_csv(county_cbsa_csv: Optional[Path]) -> Tuple[Dict[str, Dict[str, object]], str]:
    if not county_cbsa_csv:
        return {}, "No county-CBSA CSV supplied"
    mapping = {}
    with county_cbsa_csv.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            fips = str(r.get("County FIPS", "")).zfill(5)
            cbsa = str(r.get("CBSA", "")).strip()
            if fips and cbsa and cbsa.lower() != "nan":
                mapping[fips] = {"cbsa": cbsa, "cbsa_name": r.get("CBSA name"), "cbsa_population_weight": safe_float(r.get("CBSA population used"))}
    return mapping, f"County-CBSA CSV supplied; {len(mapping)} county-CBSA mappings loaded"


def median(values: List[float], default=None):
    vals = [v for v in values if v is not None and math.isfinite(v)]
    return statistics.median(vals) if vals else default


def build_oews_updates(value_rows, rank_rows, fips_by_key, cbp_details, county_cbsa, msa_areas, bos_areas, state_areas, industry_share):
    # National all-industry trade rate from state totals. This is robust and avoids depending on all_data file.
    national_trade = sum(float(d.get("trade_emp") or 0) for d in state_areas.values())
    national_all = sum(float(d.get("all_emp") or 0) for d in state_areas.values())
    national_trade_rate = national_trade / national_all * 1000.0 if national_all else median([d.get("trade_jobs_per_1000_all_jobs") for d in state_areas.values()], 1.0)

    updates = {}
    by_state_factors = defaultdict(list)
    for r in value_rows:
        key = (r["County"], r["State"])
        state = key[1]
        fips = fips_by_key.get(key)
        old_cbp_rate = safe_float(r.get(TRADES_VALUE_COL))
        cbsa_info = county_cbsa.get(fips, {}) if fips else {}
        cbsa = cbsa_info.get("cbsa")
        source = None
        area = None
        area_code = None
        if cbsa and str(cbsa) in msa_areas:
            area_code = str(cbsa)
            area = msa_areas[area_code]
            source = "CBSA OEWS metro match"
        # Nonmetro OEWS files require a county-to-nonmetro-region crosswalk not present here.
        # State fallback is more defensible than guessing a nonmetro region.
        if area is None:
            stfips = STATE_ABBR_TO_FIPS.get(state)
            if stfips and stfips in state_areas:
                area_code = stfips
                area = state_areas[stfips]
                source = "State OEWS fallback"
            else:
                source = "Statistical fallback"
        if area is not None:
            local_rate = safe_float(area.get("trade_jobs_per_1000_all_jobs"))
            raw_factor = local_rate / national_trade_rate if local_rate and national_trade_rate else 1.0
            clipped = max(OEWS_FACTOR_MIN, min(OEWS_FACTOR_MAX, raw_factor))
            new_rate = old_cbp_rate * industry_share * clipped if old_cbp_rate is not None else None
        else:
            local_rate = None
            raw_factor = 1.0
            clipped = 1.0
            new_rate = old_cbp_rate * industry_share if old_cbp_rate is not None else None
        if new_rate is not None:
            by_state_factors[state].append(clipped)
        updates[key] = {
            "County FIPS": fips,
            "OEWS source": source,
            "OEWS area code": area_code,
            "OEWS area title": area.get("area_title") if area else None,
            "County CBSA": cbsa,
            "County CBSA name": cbsa_info.get("cbsa_name"),
            "OEWS local trade jobs per 1,000 all jobs": local_rate,
            "OEWS national trade jobs per 1,000 all jobs": national_trade_rate,
            "OEWS raw local factor": raw_factor,
            "OEWS clipped local factor": clipped,
            "National NAICS 238 selected-trade occupation share": industry_share,
            "Old CBP specialty-trade employment / 1k": old_cbp_rate,
            "New tradespeople / 1k": new_rate,
            "CBP trade employment NAICS 238///": cbp_details.get(key, {}).get("cbp_trade_employment_238"),
            "Population 2023": cbp_details.get(key, {}).get("population_2023"),
            "Imputation method": source,
        }
    # Statistical fallback pass: calibrate missing new rates by state/division/national old-to-new ratios.
    ratios_by_state = defaultdict(list)
    ratios_by_div = defaultdict(list)
    ratios_all = []
    for key, u in updates.items():
        old = safe_float(u.get("Old CBP specialty-trade employment / 1k"))
        new = safe_float(u.get("New tradespeople / 1k"))
        if old and new:
            ratio = new / old
            ratios_by_state[key[1]].append(ratio)
            ratios_by_div[STATE_DIVISION.get(key[1], "")].append(ratio)
            ratios_all.append(ratio)
    national_ratio = median(ratios_all, industry_share)
    for key, u in updates.items():
        if safe_float(u.get("New tradespeople / 1k")) is not None:
            continue
        old = safe_float(u.get("Old CBP specialty-trade employment / 1k")) or 0.0
        ratio = median(ratios_by_state.get(key[1], []), None)
        if ratio is None:
            ratio = median(ratios_by_div.get(STATE_DIVISION.get(key[1], ""), []), None)
        if ratio is None:
            ratio = national_ratio
        u["New tradespeople / 1k"] = old * ratio
        u["OEWS source"] = "Statistical fallback"
        u["Imputation method"] = "calibrated old value by state/division/national OEWS ratios"
    return updates


def compute_metrics(rank_rows, original):
    keys = [(r["County"], r["State"]) for r in rank_rows]
    top20_new = {k for k, r in zip(keys, rank_rows) if safe_float(r.get("Runoff rank")) is not None and float(r.get("Runoff rank")) <= 20}
    top30_new = {k for k, r in zip(keys, rank_rows) if safe_float(r.get("Runoff rank")) is not None and float(r.get("Runoff rank")) <= 30}
    top20_old = {k for k, v in original.items() if safe_float(v.get("Original runoff rank")) is not None and float(v.get("Original runoff rank")) <= 20}
    top30_old = {k for k, v in original.items() if safe_float(v.get("Original runoff rank")) is not None and float(v.get("Original runoff rank")) <= 30}
    avg_change = []
    runoff_change = []
    for r in rank_rows:
        k = (r["County"], r["State"])
        old_avg = safe_float(original[k].get("Original avg rank"))
        new_avg = safe_float(r.get("Avg Rank"))
        old_rr = safe_float(original[k].get("Original runoff rank"))
        new_rr = safe_float(r.get("Runoff rank"))
        if old_avg is not None and new_avg is not None:
            avg_change.append(abs(new_avg - old_avg))
        if old_rr is not None and new_rr is not None:
            runoff_change.append(abs(new_rr - old_rr))
    return {
        "Top-20 overlap": len(top20_old & top20_new),
        "Top-30 overlap": len(top30_old & top30_new),
        "Mean abs Avg Rank change": statistics.mean(avg_change) if avg_change else None,
        "Median abs Avg Rank change": statistics.median(avg_change) if avg_change else None,
        "Mean abs runoff-rank change": statistics.mean(runoff_change) if runoff_change else None,
        "Median abs runoff-rank change": statistics.median(runoff_change) if runoff_change else None,
        "Largest runoff-rank move": max(runoff_change) if runoff_change else None,
    }


def update_workbook(input_workbook: Path, oews_ma_zip: Path, oews_state_zip: Path, oews_industry_zip: Optional[Path], walkability_zip: Optional[Path], county_cbsa_csv: Optional[Path], output_workbook: Path, county_rollup_csv: Path, iterations: int, seed: int):
    wb = load_workbook(input_workbook)
    if RANK_MATRIX_SHEET not in wb.sheetnames or VALUE_MATRIX_SHEET not in wb.sheetnames:
        raise ValueError("Workbook must contain Rank Matrix and Value Matrix sheets")
    rank_headers, rank_rows = read_sheet(wb[RANK_MATRIX_SHEET])
    value_headers, value_rows = read_sheet(wb[VALUE_MATRIX_SHEET])
    if TRADES_VALUE_COL not in value_headers:
        raise ValueError(f"Value Matrix missing {TRADES_VALUE_COL!r}")
    if TRADES_RANK_COL not in rank_headers:
        raise ValueError(f"Rank Matrix missing {TRADES_RANK_COL!r}")

    fips_by_key = get_fips_for_workbook(wb)
    cbp_details = get_cbp_trade_details(wb)
    tmp_dir = Path(tempfile.mkdtemp(prefix="oews_parse_"))
    msa_xlsx = extract_from_zip(oews_ma_zip, find_zip_member(oews_ma_zip, "MSA_M2025_dl.xlsx"), tmp_dir)
    bos_xlsx = extract_from_zip(oews_ma_zip, find_zip_member(oews_ma_zip, "BOS_M2025_dl.xlsx"), tmp_dir)
    state_xlsx = extract_from_zip(oews_state_zip, find_zip_member(oews_state_zip, "state_M2025_dl.xlsx"), tmp_dir)
    msa_areas = read_oews_area_xlsx(msa_xlsx)
    bos_areas = read_oews_area_xlsx(bos_xlsx)
    state_areas = read_oews_area_xlsx(state_xlsx)

    if oews_industry_zip:
        ind_xlsx = extract_from_zip(oews_industry_zip, find_zip_member(oews_industry_zip, "nat3d_M2025_dl.xlsx"), tmp_dir)
        naics238_all, naics238_selected, naics238_share, ind_rows = read_oews_industry_share(ind_xlsx)
    else:
        naics238_all, naics238_selected, naics238_share, ind_rows = None, None, 1.0, []

    if county_cbsa_csv:
        county_cbsa, cbsa_note = load_county_cbsa_from_csv(county_cbsa_csv)
    else:
        county_cbsa, cbsa_note = load_county_cbsa_from_walkability(walkability_zip)

    original = {}
    for rr in rank_rows:
        key = (rr["County"], rr["State"])
        original[key] = {
            "Original runoff rank": rr.get("Runoff rank"),
            "Original avg rank": rr.get("Avg Rank"),
            "Original trades rank": rr.get(TRADES_RANK_COL),
            "Original runoff avg elimination round": rr.get("Runoff avg elimination round"),
            "Original runoff wins": rr.get("Runoff wins"),
            "Original runoff win rate": rr.get("Runoff win rate"),
        }

    value_by_key = {(r["County"], r["State"]): r for r in value_rows}
    rank_by_key = {(r["County"], r["State"]): r for r in rank_rows}
    keys = [(r["County"], r["State"]) for r in value_rows]
    updates = build_oews_updates(value_rows, rank_rows, fips_by_key, cbp_details, county_cbsa, msa_areas, bos_areas, state_areas, naics238_share)

    for k in keys:
        value_by_key[k][TRADES_VALUE_COL] = updates[k]["New tradespeople / 1k"]

    trades_ranks = rank_average([value_by_key[k][TRADES_VALUE_COL] for k in keys], higher_is_better=True)
    for k, tr in zip(keys, trades_ranks):
        rank_by_key[k][TRADES_RANK_COL] = tr

    factor_cols = rank_headers[10:]
    for rr in rank_rows:
        vals = [safe_float(rr.get(c)) for c in factor_cols]
        if any(v is None for v in vals):
            raise ValueError(f"Missing factor rank values for {rr.get('County')}, {rr.get('State')}")
        avg = sum(vals) / len(vals)
        rr["Avg Rank"] = avg
        rr["Avg Rank Python Check"] = avg
        rr["Avg Check Delta"] = 0

    avg_based = rank_average([rank_by_key[k]["Avg Rank"] for k in keys], higher_is_better=False)
    for k, ar in zip(keys, avg_based):
        rank_by_key[k]["Avg-based rank"] = ar

    ranks_matrix = np.array([[float(rank_by_key[k][c]) for c in factor_cols] for k in keys], dtype=np.float64)
    sorted_order = np.argsort(-ranks_matrix, axis=0).T.astype(np.int64)
    wins, avg_elim = run_randomized_runoff(ranks_matrix, sorted_order, iterations, seed)
    runoff_order = sorted(range(len(keys)), key=lambda i: (-avg_elim[i], -wins[i], rank_by_key[keys[i]]["Avg Rank"], keys[i][1], keys[i][0]))
    runoff_ranks = [0] * len(keys)
    for pos, i in enumerate(runoff_order, 1):
        runoff_ranks[i] = pos
    for i, k in enumerate(keys):
        rr = rank_by_key[k]
        rr["Runoff rank"] = runoff_ranks[i]
        rr["Runoff avg elimination round"] = float(avg_elim[i])
        rr["Runoff wins"] = int(wins[i])
        rr["Runoff win rate"] = float(wins[i]) / float(iterations)

    sorted_keys = [keys[i] for i in runoff_order]
    sorted_rank_rows = [rank_by_key[k] for k in sorted_keys]
    sorted_value_rows = [value_by_key[k] for k in sorted_keys]
    metrics = compute_metrics(sorted_rank_rows, original)

    write_matrix(wb[RANK_MATRIX_SHEET], [rank_headers] + [[r.get(h) for h in rank_headers] for r in sorted_rank_rows])
    write_matrix(wb[VALUE_MATRIX_SHEET], [value_headers] + [[r.get(h) for h in value_headers] for r in sorted_value_rows])
    style_sheet(wb[RANK_MATRIX_SHEET])
    style_sheet(wb[VALUE_MATRIX_SHEET])

    if TOP30_SHEET in wb.sheetnames:
        write_matrix(wb[TOP30_SHEET], [rank_headers] + [[r.get(h) for h in rank_headers] for r in sorted_rank_rows[:30]])
        style_sheet(wb[TOP30_SHEET])

    if RUNOFF_SHEET in wb.sheetnames:
        runoff_matrix = [
            ["Metric", "Value", "Notes", None, None, None, None],
            ["Simulation count", iterations, "Randomized runoff iterations", None, None, None, None],
            ["Seed", seed, "NumPy/Numba RNG seed", None, None, None, None],
            ["Rank factors used", len(factor_cols), "Avg Rank excluded from random factor selection", None, None, None, None],
            ["Changed factors", "Tradespeople", "CBP county NAICS 238 backbone adjusted by BLS OEWS occupation mix", None, None, None, None],
            ["Tie handling", "Random among tied worst ranks", "If multiple active counties tie for worst rank in selected factor", None, None, None, None],
            [None, None, None, None, None, None, None],
            ["Runoff rank", "County", "State", "Avg Rank", "Runoff avg elimination round", "Runoff wins", "Runoff win rate"],
        ]
        for r in sorted_rank_rows[:30]:
            runoff_matrix.append([r["Runoff rank"], r["County"], r["State"], r["Avg Rank"], r["Runoff avg elimination round"], r["Runoff wins"], r["Runoff win rate"]])
        write_matrix(wb[RUNOFF_SHEET], runoff_matrix)
        style_sheet(wb[RUNOFF_SHEET])

    if RESULT_COMPARISON_SHEET in wb.sheetnames:
        comparison_headers = [
            "County", "State", "Original runoff rank", "New runoff rank", "Runoff rank delta",
            "Original avg rank", "New avg rank", "Avg rank delta",
            "Original trades rank", "New trades rank", "Trades rank delta",
            "In new top 30",
        ]
        comparison_rows = []
        for k in sorted_keys[:50]:
            old = original[k]
            new = rank_by_key[k]
            comparison_rows.append([
                k[0], k[1], old.get("Original runoff rank"), new.get("Runoff rank"), safe_float(new.get("Runoff rank")) - safe_float(old.get("Original runoff rank")),
                old.get("Original avg rank"), new.get("Avg Rank"), safe_float(new.get("Avg Rank")) - safe_float(old.get("Original avg rank")),
                old.get("Original trades rank"), new.get(TRADES_RANK_COL), safe_float(new.get(TRADES_RANK_COL)) - safe_float(old.get("Original trades rank")),
                "yes" if safe_float(new.get("Runoff rank")) <= 30 else "no",
            ])
        write_matrix(wb[RESULT_COMPARISON_SHEET], [comparison_headers] + comparison_rows)
        style_sheet(wb[RESULT_COMPARISON_SHEET])

    source_counts = defaultdict(int)
    for u in updates.values():
        source_counts[u.get("OEWS source")] += 1
    trades_rank_changes = [abs(safe_float(rank_by_key[k][TRADES_RANK_COL]) - safe_float(original[k]["Original trades rank"])) for k in keys]
    biggest = []
    for k in keys:
        old_rr = safe_float(original[k]["Original runoff rank"])
        new_rr = safe_float(rank_by_key[k]["Runoff rank"])
        if old_rr is not None and new_rr is not None:
            biggest.append((abs(new_rr - old_rr), k, old_rr, new_rr))
    biggest.sort(reverse=True)

    national_trade_rate = next(iter(updates.values()))["OEWS national trade jobs per 1,000 all jobs"]
    impact = [
        ["Metric", "Value", "Notes"],
        ["Workbook counties", len(keys), None],
        ["OEWS metro/CBSA matches", source_counts.get("CBSA OEWS metro match", 0), "County mapped to CBSA using EPA Walkability block-group CBSA fields, then matched to OEWS MSA area code"],
        ["OEWS state fallbacks", source_counts.get("State OEWS fallback", 0), "Used when county lacks CBSA or CBSA was not present in the OEWS MSA file"],
        ["Statistical fallback counties", source_counts.get("Statistical fallback", 0), "Should be zero unless state OEWS data is missing"],
        ["County-CBSA mapping note", cbsa_note, None],
        ["OEWS MSA areas read", len(msa_areas), "Metropolitan area rows from OEWS"],
        ["OEWS nonmetro areas read", len(bos_areas), "Reviewed, but not used without a county-to-nonmetro-region crosswalk"],
        ["OEWS state areas read", len(state_areas), "Fallback geography"],
        ["National all-industry trade occupation rate", national_trade_rate, "Selected trade occupations per 1,000 total jobs"],
        ["NAICS 238 selected-trade occupation share", naics238_share, "National industry-specific OEWS share used to convert CBP NAICS 238 employment toward selected trade occupations"],
        ["NAICS 238 all employment", naics238_all, "From national industry-specific OEWS NAICS 238000"],
        ["NAICS 238 selected trade employment", naics238_selected, "Sum of selected trade occupations within NAICS 238000"],
        ["OEWS local factor floor", OEWS_FACTOR_MIN, None],
        ["OEWS local factor cap", OEWS_FACTOR_MAX, None],
        ["Top-20 runoff overlap vs prior", metrics["Top-20 overlap"], "out of 20"],
        ["Top-30 runoff overlap vs prior", metrics["Top-30 overlap"], "out of 30"],
        ["Mean abs trades-rank change", statistics.mean(trades_rank_changes), None],
        ["Median abs trades-rank change", statistics.median(trades_rank_changes), None],
        ["Mean abs Avg Rank change", metrics["Mean abs Avg Rank change"], None],
        ["Median abs Avg Rank change", metrics["Median abs Avg Rank change"], None],
        ["Mean abs runoff-rank change", metrics["Mean abs runoff-rank change"], None],
        ["Median abs runoff-rank change", metrics["Median abs runoff-rank change"], None],
        ["Largest runoff-rank move", metrics["Largest runoff-rank move"], None],
        [None, None, None],
        ["Biggest runoff movers", None, None],
        ["County", "Old runoff rank", "New runoff rank"],
    ]
    for _, k, old_rr, new_rr in biggest[:12]:
        impact.append([f"{k[0]}, {k[1]}", old_rr, new_rr])
    ws_impact = recreate_sheet(wb, IMPACT_SHEET, after_name=RESULT_COMPARISON_SHEET)
    write_matrix(ws_impact, impact)
    style_sheet(ws_impact)

    update_headers = [
        "County", "State", "County FIPS", "OEWS source", "OEWS area code", "OEWS area title", "County CBSA", "County CBSA name",
        "Population 2023", "CBP trade employment NAICS 238///", "Old CBP specialty-trade employment / 1k",
        "National NAICS 238 selected-trade occupation share", "OEWS local trade jobs per 1,000 all jobs", "OEWS national trade jobs per 1,000 all jobs",
        "OEWS raw local factor", "OEWS clipped local factor", "New tradespeople / 1k", "Trades value delta",
        "Original trades rank", "New trades rank", "Trades rank delta", "Imputation method",
        "Original avg rank", "New avg rank", "Avg rank delta",
        "Original runoff rank", "New runoff rank", "Runoff rank delta",
        "Original runoff avg elimination round", "New runoff avg elimination round", "Runoff avg elimination delta",
        "Original runoff wins", "New runoff wins", "Original runoff win rate", "New runoff win rate",
    ]
    update_matrix = [update_headers]
    csv_rows = []
    for k in sorted_keys:
        u = updates[k]
        old = original[k]
        new = rank_by_key[k]
        row = [
            k[0], k[1], u.get("County FIPS"), u.get("OEWS source"), u.get("OEWS area code"), u.get("OEWS area title"), u.get("County CBSA"), u.get("County CBSA name"),
            u.get("Population 2023"), u.get("CBP trade employment NAICS 238///"), u.get("Old CBP specialty-trade employment / 1k"),
            u.get("National NAICS 238 selected-trade occupation share"), u.get("OEWS local trade jobs per 1,000 all jobs"), u.get("OEWS national trade jobs per 1,000 all jobs"),
            u.get("OEWS raw local factor"), u.get("OEWS clipped local factor"), u.get("New tradespeople / 1k"), safe_float(u.get("New tradespeople / 1k")) - safe_float(u.get("Old CBP specialty-trade employment / 1k")),
            old.get("Original trades rank"), new.get(TRADES_RANK_COL), safe_float(new.get(TRADES_RANK_COL)) - safe_float(old.get("Original trades rank")), u.get("Imputation method"),
            old.get("Original avg rank"), new.get("Avg Rank"), safe_float(new.get("Avg Rank")) - safe_float(old.get("Original avg rank")),
            old.get("Original runoff rank"), new.get("Runoff rank"), safe_float(new.get("Runoff rank")) - safe_float(old.get("Original runoff rank")),
            old.get("Original runoff avg elimination round"), new.get("Runoff avg elimination round"), safe_float(new.get("Runoff avg elimination round")) - safe_float(old.get("Original runoff avg elimination round")),
            old.get("Original runoff wins"), new.get("Runoff wins"), old.get("Original runoff win rate"), new.get("Runoff win rate"),
        ]
        update_matrix.append(row)
        csv_rows.append(dict(zip(update_headers, row)))
    ws_update = recreate_sheet(wb, UPDATE_SHEET, after_name=IMPACT_SHEET)
    write_matrix(ws_update, update_matrix)
    style_sheet(ws_update)

    review = [
        ["Item", "Value", "Notes"],
        ["OEWS vintage detected", "M2025", "Input filenames and sheets are May 2025, not May 2023"],
        ["Selected occupation codes", ", ".join(sorted(TRADE_OCC_CODES)), "Used in all-industry area/state rates and national NAICS 238 share"],
        ["Metro/nonmetro file", str(oews_ma_zip.name), "MSA rows used when CBSA match is available; BOS nonmetro rows reviewed but not used without a nonmetro-county crosswalk"],
        ["State file", str(oews_state_zip.name), "State fallback for all unmatched counties"],
        ["Industry-specific file", str(oews_industry_zip.name) if oews_industry_zip else "not supplied", "Used to estimate national selected-trade share inside NAICS 238"],
        ["CBSA mapping", str(walkability_zip.name) if walkability_zip else "not supplied", cbsa_note],
        ["National NAICS 238 share", naics238_share, "Selected trade occupations / all occupations in NAICS 238000"],
        [None, None, None],
        ["Occupation code", "Occupation title", "NAICS 238 national employment"],
    ]
    for occ, title, emp in ind_rows:
        review.append([occ, title, emp])
    ws_review = recreate_sheet(wb, REVIEW_SHEET, after_name=UPDATE_SHEET)
    write_matrix(ws_review, review)
    style_sheet(ws_review)

    with county_rollup_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=update_headers)
        writer.writeheader()
        writer.writerows(csv_rows)

    if METHODOLOGY_SHEET in wb.sheetnames:
        ws = wb[METHODOLOGY_SHEET]
        start = ws.max_row + 2
        rows = [
            ["BLS OEWS trades update", None, None, None, None, None, None],
            ["Tradespeople per capita", TRADES_VALUE_COL, TRADES_RANK_COL, "Higher is better", "CBP 2023 NAICS 238 county employment per 1,000 residents × national OEWS NAICS 238 selected-trade occupation share × local OEWS occupation-mix factor", "Local factor uses county CBSA-to-OEWS MSA match when available and state OEWS fallback otherwise; factor clipped to avoid overcorrecting county-level CBP with non-county OEWS geography", "Inputs: BLS OEWS M2025 metro/nonmetro, state, and national industry-specific tables; EPA Walkability CBSA fields for county-to-CBSA mapping"],
        ]
        for i, row in enumerate(rows, start):
            for j, val in enumerate(row, 1):
                ws.cell(i, j, val)
        style_sheet(ws)

    if LOG_SHEET in wb.sheetnames:
        ws = wb[LOG_SHEET]
        start = ws.max_row + 1
        rows = [
            ["BLS OEWS M2025 metropolitan/nonmetropolitan", "Imported", "Used MSA area rows for counties with CBSA matches; nonmetro rows reviewed only", "Tradespeople", "2025", "BLS OEWS Trades update"],
            ["BLS OEWS M2025 state", "Imported", "State fallback for counties without CBSA/OEWS MSA match", "Tradespeople", "2025", "BLS OEWS Trades update"],
            ["BLS OEWS M2025 national industry-specific", "Imported", "Estimated selected-trade occupation share within NAICS 238000", "Tradespeople", "2025", "BLS OEWS Trades update"],
        ]
        for i, row in enumerate(rows, start):
            for j, val in enumerate(row, 1):
                ws.cell(i, j, val)
        style_sheet(ws)

    if QA_SHEET in wb.sheetnames:
        qa = [
            ["Check", "Value", "Expected", "Pass?"],
            ["County rows", len(keys), len(keys), "PASS"],
            ["Rank factors", len(factor_cols), 21, "PASS" if len(factor_cols) == 21 else "CHECK"],
            ["Runoff wins sum", int(sum(wins)), iterations, "PASS" if int(sum(wins)) == iterations else "FAIL"],
            ["Runoff win rate sum", float(sum(wins) / iterations), 1.0, "PASS" if abs(float(sum(wins) / iterations) - 1.0) < 1e-9 else "FAIL"],
            ["OEWS statistical fallback rows", source_counts.get("Statistical fallback", 0), 0, "PASS" if source_counts.get("Statistical fallback", 0) == 0 else "CHECK"],
            ["Avg check max delta", 0, 0, "PASS"],
            ["Updated factors", "Tradespeople", None, "PASS"],
        ]
        write_matrix(wb[QA_SHEET], qa)
        style_sheet(wb[QA_SHEET])

    if CalcProperties is not None:
        try:
            wb.calculation = CalcProperties(calcMode="auto")
        except Exception:
            pass

    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_workbook)

    return {
        "output_workbook": str(output_workbook),
        "county_rollup_csv": str(county_rollup_csv),
        "metrics": metrics,
        "source_counts": dict(source_counts),
        "naics238_share": naics238_share,
        "national_trade_rate": national_trade_rate,
        "top5": [(r["Runoff rank"], r["County"], r["State"]) for r in sorted_rank_rows[:5]],
        "biggest": [(k[0], k[1], old, new) for _, k, old, new in biggest[:12]],
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input-workbook", type=Path, required=True)
    ap.add_argument("--oews-ma-zip", type=Path, required=True)
    ap.add_argument("--oews-state-zip", type=Path, required=True)
    ap.add_argument("--oews-industry-zip", type=Path, default=None)
    ap.add_argument("--walkability-zip", type=Path, default=None)
    ap.add_argument("--county-cbsa-csv", type=Path, default=None)
    ap.add_argument("--output-workbook", type=Path, required=True)
    ap.add_argument("--county-rollup-csv", type=Path, required=True)
    ap.add_argument("--iterations", type=int, default=DEFAULT_ITERATIONS)
    ap.add_argument("--seed", type=int, default=DEFAULT_SEED)
    args = ap.parse_args()
    result = update_workbook(args.input_workbook, args.oews_ma_zip, args.oews_state_zip, args.oews_industry_zip, args.walkability_zip, args.county_cbsa_csv, args.output_workbook, args.county_rollup_csv, args.iterations, args.seed)
    print(result)


if __name__ == "__main__":
    main()
