#!/usr/bin/env python3
"""
Update the county-rankings workbook with EPA Access to Jobs and Workers via Transit
DBF data, apply a missing-data guardrail for counties outside EPA/GTFS coverage,
re-rank the public transportation factor, recompute average ranks, rerun the
100k randomized-runoff MCMC, and write an updated workbook plus an audit CSV.

Expected inputs
---------------
1. Existing workbook with sheets:
   - Rank Matrix
   - Value Matrix
   - Top 30
   - Result Comparison
   - Runoff Simulation
   - Sources & Methodology
   - Data Acquisition Log
   - QA Checks
   - preferably EPA Walkability Update, for County FIPS mapping created in the
     walkability update pass
2. EPA_SLD_Trans45_DBF.zip containing SLD_Trans45.dbf.
3. Optional county FIPS lookup CSV with columns fips,name,state, used only if
   the workbook does not already include an EPA Walkability Update sheet.

Method
------
- Reads the DBF directly from the ZIP with a small built-in DBF parser.
- Uses the DBF's 10-character shapefile field names:
    TrAccess_I  = TrAccess_Indexi
    Pct_Jobs_b  = Pct_Jobs_byTr
    Pct_Pop_by  = Pct_Pop_byTr
    Pct_Wrks_b  = Pct_Wrks_byTr
- Computes a county-level EPA transit score for counties present in the DBF:
    100 * (0.50*TrAccess_I + 0.30*Pct_Jobs_byTr
           + 0.10*Pct_Pop_byTr + 0.10*Pct_Wrks_byTr)
  using a simple mean of block-group component values within county FIPS.
- By default, counties outside the EPA DBF coverage are placed in a lower
  fallback score band below every county with measured EPA data. Their relative
  order within that lower band is preserved using the prior modeled transit
  estimate. This avoids rewarding missing GTFS/EPA coverage while not treating
  all missing counties as equal.
- Re-ranks public transportation only; all other factor values/ranks are
  preserved.
- Reruns the randomized runoff with the same rank-matrix rules:
  random permutation cycles over rank factors; worst rank on selected factor is
  eliminated; tied worst rows are broken randomly.

Install dependencies
--------------------
pip install openpyxl numpy numba

Example
-------
python update_transit_and_rerun_mcmc.py \
  --input-workbook county_rankings_epa_walkability_100k.xlsx \
  --epa-transit-zip EPA_SLD_Trans45_DBF.zip \
  --output-workbook county_rankings_epa_walkability_transit_guardrailed_100k.xlsx \
  --county-rollup-csv epa_transit_county_rollup_guardrailed.csv
"""

from __future__ import annotations

import argparse
import csv
import math
import re
import statistics
import struct
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

import numpy as np
from numba import njit
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
try:
    from openpyxl.workbook.properties import CalcProperties
except Exception:  # pragma: no cover
    CalcProperties = None


DEFAULT_ITERATIONS = 100_000
DEFAULT_SEED = 2026060901

RANK_MATRIX_SHEET = "Rank Matrix"
VALUE_MATRIX_SHEET = "Value Matrix"
TOP30_SHEET = "Top 30"
RESULT_COMPARISON_SHEET = "Result Comparison"
RUNOFF_SHEET = "Runoff Simulation"
METHODOLOGY_SHEET = "Sources & Methodology"
LOG_SHEET = "Data Acquisition Log"
QA_SHEET = "QA Checks"
WALKABILITY_AUDIT_SHEET = "EPA Walkability Update"
IMPACT_SHEET = "EPA Transit Guardrail Impact"
UPDATE_SHEET = "EPA Transit Guardrail Update"

TRANSIT_VALUE_COL = "Est. public transportation score"
TRANSIT_RANK_COL = "Public transportation rank"

HEADER_FILL = "1F4E79"
SUBHEADER_FILL = "D9EAF7"
BORDER_COLOR = "B7B7B7"


def normalize_county_name(name: str) -> str:
    if name is None:
        return ""
    s = str(name).strip()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.upper().replace("ST.", "ST").replace("SAINT ", "ST ")
    s = s.replace("'", "")
    s = re.sub(r"\s*\(.*?\)", "", s)
    suffixes = [
        " COUNTY AND BOROUGH", " CITY AND BOROUGH", " CENSUS AREA",
        " MUNICIPALITY", " CONSOLIDATED GOVERNMENT", " UNIFIED GOVERNMENT",
        " BOROUGH", " PARISH", " COUNTY", " CITY",
    ]
    changed = True
    while changed:
        changed = False
        for suffix in suffixes:
            if s.endswith(suffix):
                s = s[: -len(suffix)]
                changed = True
                break
    s = re.sub(r"[^A-Z0-9]+", " ", s).strip()
    return re.sub(r"\s+", " ", s)


def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = list(rows[0])
    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        padded = list(row) + [None] * (len(headers) - len(row))
        data.append(dict(zip(headers, padded[: len(headers)])))
    return headers, data


def write_matrix(ws, matrix: List[List[object]]):
    ws.delete_rows(1, ws.max_row)
    for r_idx, row in enumerate(matrix, 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, value)


def recreate_sheet(wb, name: str, after_name: str | None = None):
    if name in wb.sheetnames:
        del wb[name]
    if after_name and after_name in wb.sheetnames:
        idx = wb.sheetnames.index(after_name)
        return wb.create_sheet(name, idx + 1)
    return wb.create_sheet(name)


def style_sheet(ws, max_width: int = 36):
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color=BORDER_COLOR)
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 8), max_width)


def read_fips_lookup_csv(path: str) -> Dict[Tuple[str, str], str]:
    mapping = {}
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        required = {"fips", "name", "state"}
        if not required.issubset(reader.fieldnames or []):
            raise ValueError(f"FIPS lookup must contain columns {sorted(required)}")
        for row in reader:
            if not row.get("state"):
                continue
            mapping[(normalize_county_name(row["name"]), row["state"].upper())] = str(row["fips"]).zfill(5)
    return mapping


def get_fips_for_workbook(wb, value_rows, fips_csv: str | None = None) -> Dict[Tuple[str, str], str]:
    # Preferred path: use FIPS already audited by the walkability update.
    if WALKABILITY_AUDIT_SHEET in wb.sheetnames:
        headers, rows = read_sheet(wb[WALKABILITY_AUDIT_SHEET])
        if {"County", "State", "County FIPS"}.issubset(headers):
            return {
                (r["County"], r["State"]): str(r["County FIPS"]).zfill(5)
                for r in rows
                if r.get("County") is not None and r.get("State") is not None and r.get("County FIPS") is not None
            }

    if not fips_csv:
        raise ValueError(
            "No EPA Walkability Update sheet found for FIPS mapping. "
            "Pass --fips-csv with columns fips,name,state."
        )

    lookup = read_fips_lookup_csv(fips_csv)
    overrides = {
        ("Athens-Clarke", "GA"): "13059",
        ("Dona Ana", "NM"): "35013",
        ("Baltimore city", "MD"): "24510",
        ("St. Louis city", "MO"): "29510",
        ("Richmond city", "VA"): "51760",
        ("Fairfax", "VA"): "51059",
    }
    out = {}
    missing = []
    for r in value_rows:
        key = (r["County"], r["State"])
        if key in overrides:
            out[key] = overrides[key]
            continue
        fips = lookup.get((normalize_county_name(r["County"]), str(r["State"]).upper()))
        if not fips:
            missing.append(key)
        else:
            out[key] = fips
    if missing:
        raise ValueError("Could not map these workbook rows to FIPS: " + ", ".join(map(str, missing[:20])))
    return out


def read_dbf_header_from_bytes(data: bytes):
    num_records = struct.unpack("<I", data[4:8])[0]
    header_length = struct.unpack("<H", data[8:10])[0]
    record_length = struct.unpack("<H", data[10:12])[0]
    fields = []
    pos = 32
    while pos < header_length - 1:
        desc = data[pos : pos + 32]
        if desc[0] == 0x0D:
            break
        name = desc[0:11].split(b"\x00", 1)[0].decode("ascii", "ignore")
        ftype = chr(desc[11])
        length = desc[16]
        dec = desc[17]
        fields.append((name, ftype, length, dec))
        pos += 32
    return num_records, header_length, record_length, fields


def iter_dbf_records_from_zip(zip_path: str):
    with zipfile.ZipFile(zip_path) as z:
        members = [n for n in z.namelist() if n.lower().endswith(".dbf")]
        if not members:
            raise ValueError("ZIP does not contain a .dbf file")
        member = members[0]
        with z.open(member) as f:
            header = f.read(32)
            header_length = struct.unpack("<H", header[8:10])[0]
            record_length = struct.unpack("<H", header[10:12])[0]
            rest = f.read(header_length - 32)
            num_records, _, _, fields = read_dbf_header_from_bytes(header + rest)

            offsets = []
            offset = 1  # deletion flag
            for name, ftype, length, dec in fields:
                offsets.append((name, ftype, offset, length, dec))
                offset += length

            for _ in range(num_records):
                rec = f.read(record_length)
                if len(rec) < record_length:
                    break
                if rec[0:1] == b"*":
                    continue
                out = {}
                for name, ftype, off, length, dec in offsets:
                    text = rec[off : off + length].decode("latin-1", "ignore").strip()
                    if text == "":
                        value = None
                    elif ftype in ("N", "F"):
                        try:
                            value = float(text) if ("." in text or ftype == "F") else int(text)
                        except ValueError:
                            value = None
                    else:
                        value = text
                    out[name] = value
                yield out


def build_transit_rollup(zip_path: str):
    fields = ["TrAccess_I", "Pct_Jobs_b", "Pct_Pop_by", "Pct_Wrks_b", "Jobs_byTr", "Pop_byTr", "Wrks_byTr"]
    agg = defaultdict(lambda: {k: [] for k in fields + ["CBSA_Name"]})
    total_records = 0

    for rec in iter_dbf_records_from_zip(zip_path):
        total_records += 1
        geoid = str(rec.get("GEOID10") or "").strip()
        if len(geoid) < 5:
            continue
        fips = geoid[:5]
        for field in fields:
            value = rec.get(field)
            if isinstance(value, (int, float)):
                # -99999 appears as a sentinel in TrAccess_I for a tiny number of records.
                if value < 0:
                    continue
                agg[fips][field].append(float(value))
        if rec.get("CBSA_Name"):
            agg[fips]["CBSA_Name"].append(str(rec["CBSA_Name"]))

    rollup = {}
    for fips, vals in agg.items():
        d = {"County FIPS": fips}
        for field in fields:
            x = vals[field]
            d[f"{field}_mean"] = float(statistics.mean(x)) if x else None
            d[f"{field}_median"] = float(statistics.median(x)) if x else None
            d[f"{field}_min"] = float(min(x)) if x else None
            d[f"{field}_max"] = float(max(x)) if x else None
            d[f"{field}_block_groups"] = int(len(x))
        d["EPA block groups"] = max(d[f"{field}_block_groups"] for field in fields)
        d["Primary CBSA"] = Counter(vals["CBSA_Name"]).most_common(1)[0][0] if vals["CBSA_Name"] else None

        components = [d["TrAccess_I_mean"], d["Pct_Jobs_b_mean"], d["Pct_Pop_by_mean"], d["Pct_Wrks_b_mean"]]
        if all(v is not None for v in components):
            tr, jobs, pop, wrks = [min(max(float(v), 0.0), 1.0) for v in components]
            d["EPA transit composite score"] = 100.0 * (0.50 * tr + 0.30 * jobs + 0.10 * pop + 0.10 * wrks)
        else:
            d["EPA transit composite score"] = None
        rollup[fips] = d

    return rollup, total_records


def rank_average(values: Iterable[float], higher_is_better: bool) -> List[float]:
    pairs = [(float(v), i) for i, v in enumerate(values)]
    pairs.sort(key=lambda t: t[0], reverse=higher_is_better)
    ranks = [0.0] * len(pairs)
    i = 0
    while i < len(pairs):
        j = i + 1
        while j < len(pairs) and pairs[j][0] == pairs[i][0]:
            j += 1
        avg_rank = (i + 1 + j) / 2.0
        for k in range(i, j):
            ranks[pairs[k][1]] = avg_rank
        i = j
    return ranks


def rescale_to_band(value: float, src_min: float, src_max: float, dst_min: float, dst_max: float) -> float:
    """Linearly rescale value from source range to destination range."""
    value = float(value)
    if src_max <= src_min:
        return (dst_min + dst_max) / 2.0
    x = (value - src_min) / (src_max - src_min)
    x = min(max(x, 0.0), 1.0)
    return dst_min + x * (dst_max - dst_min)


@njit
def run_randomized_runoff(ranks: np.ndarray, iterations: int, seed: int):
    np.random.seed(seed)
    n, f = ranks.shape
    wins = np.zeros(n, np.int64)
    elim_sum = np.zeros(n, np.float64)
    factor_order = np.empty(f, np.int64)
    active_idx = np.empty(n, np.int64)
    ties = np.empty(n, np.int64)

    for _ in range(iterations):
        for i in range(n):
            active_idx[i] = i
        active_count = n
        pos = f
        round_no = 0

        while active_count > 1:
            if pos >= f:
                for j in range(f):
                    factor_order[j] = j
                for j in range(f - 1, 0, -1):
                    k = np.random.randint(j + 1)
                    tmp = factor_order[j]
                    factor_order[j] = factor_order[k]
                    factor_order[k] = tmp
                pos = 0

            col = factor_order[pos]
            pos += 1

            worst = -1e308
            tie_count = 0
            for p in range(active_count):
                i = active_idx[p]
                v = ranks[i, col]
                if v > worst:
                    worst = v
                    ties[0] = p
                    tie_count = 1
                elif v == worst:
                    ties[tie_count] = p
                    tie_count += 1

            loser_pos = ties[np.random.randint(tie_count)]
            loser = active_idx[loser_pos]
            round_no += 1
            elim_sum[loser] += round_no
            active_count -= 1
            active_idx[loser_pos] = active_idx[active_count]

        winner = active_idx[0]
        wins[winner] += 1
        elim_sum[winner] += n

    return wins, elim_sum / iterations


def clean_cell(value):
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def update_workbook(
    input_workbook: str,
    epa_transit_zip: str,
    output_workbook: str,
    county_rollup_csv: str | None,
    fips_csv: str | None,
    iterations: int,
    seed: int,
    missing_policy: str,
):
    wb_values = load_workbook(input_workbook, data_only=True)
    rank_headers, rank_rows = read_sheet(wb_values[RANK_MATRIX_SHEET])
    value_headers, value_rows = read_sheet(wb_values[VALUE_MATRIX_SHEET])

    fips_map = get_fips_for_workbook(wb_values, value_rows, fips_csv)
    transit_rollup, total_bg_records = build_transit_rollup(epa_transit_zip)

    # Precompute matched/missing sets from the original pre-transit-update workbook.
    # For the guardrail policy, measured EPA counties occupy score tier >= 1.0;
    # missing counties occupy [0.001, 0.999] based on their old modeled estimate.
    workbook_keys = []
    old_score_by_key = {}
    matched_keys = set()
    missing_keys = set()
    for vrow in value_rows:
        key = (vrow["County"], vrow["State"])
        fips = fips_map[key]
        epa = transit_rollup.get(fips, {})
        epa_score = epa.get("EPA transit composite score")
        workbook_keys.append(key)
        old_score_by_key[key] = float(vrow[TRANSIT_VALUE_COL])
        if epa_score is None:
            missing_keys.add(key)
        else:
            matched_keys.add(key)

    old_missing_scores = [old_score_by_key[k] for k in missing_keys]
    missing_old_min = min(old_missing_scores) if old_missing_scores else None
    missing_old_max = max(old_missing_scores) if old_missing_scores else None

    value_by_key = {(r["County"], r["State"]): r for r in value_rows}
    original = {}
    for r in rank_rows:
        key = (r["County"], r["State"])
        original[key] = {
            "Original runoff rank": r["Runoff rank"],
            "Original avg rank": r["Avg Rank Python Check"],
            "Original public transportation rank": r[TRANSIT_RANK_COL],
            "Original public transportation score": value_by_key[key][TRANSIT_VALUE_COL],
            "Original runoff avg elimination round": r["Runoff avg elimination round"],
            "Original runoff wins": r["Runoff wins"],
            "Original runoff win rate": r["Runoff win rate"],
        }

    updated_scores = []
    audit_rows = []
    matched_count = 0

    for vrow in value_rows:
        key = (vrow["County"], vrow["State"])
        fips = fips_map[key]
        epa = transit_rollup.get(fips, {})
        epa_score = epa.get("EPA transit composite score")
        old_score = float(vrow[TRANSIT_VALUE_COL])
        if epa_score is None:
            if missing_policy == "zero":
                new_score = 0.0
                status = "No EPA DBF coverage; set to zero by --missing-policy zero"
            elif missing_policy == "error":
                raise ValueError(f"No EPA DBF coverage for {key} / {fips}")
            elif missing_policy == "retain_prior":
                new_score = old_score
                status = "No EPA DBF coverage; retained prior modeled estimate"
            elif missing_policy == "guardrail_low_band":
                new_score = rescale_to_band(old_score, missing_old_min, missing_old_max, 0.001, 0.999)
                status = "No EPA DBF coverage; prior modeled estimate rescaled into lower guardrail band below all EPA-matched counties"
            else:
                raise ValueError(f"Unsupported missing policy: {missing_policy}")
        else:
            matched_count += 1
            if missing_policy == "guardrail_low_band":
                # Shift all measured EPA scores into a higher tier. This preserves the order of
                # measured counties while guaranteeing every EPA-covered county ranks above
                # every non-covered county, even when raw EPA composite is 0.
                new_score = 1.0 + float(epa_score)
                status = "EPA DBF matched; raw EPA composite shifted into measured-data tier"
            else:
                new_score = float(epa_score)
                status = "EPA DBF matched; replaced prior estimate"

        vrow[TRANSIT_VALUE_COL] = new_score
        updated_scores.append(new_score)

        audit_rows.append({
            "County": key[0],
            "State": key[1],
            "County FIPS": fips,
            "Transit update status": status,
            "Primary CBSA": epa.get("Primary CBSA"),
            "EPA block groups": epa.get("EPA block groups", 0),
            "EPA transit composite score": epa.get("EPA transit composite score"),
            "Transit data tier": "EPA measured" if epa_score is not None else "Missing EPA/GTFS; lower fallback band",
            "Missing fallback old score min": missing_old_min,
            "Missing fallback old score max": missing_old_max,
            "EPA TrAccess index mean": epa.get("TrAccess_I_mean"),
            "EPA pct jobs by transit mean": epa.get("Pct_Jobs_b_mean"),
            "EPA pct pop by transit mean": epa.get("Pct_Pop_by_mean"),
            "EPA pct workers by transit mean": epa.get("Pct_Wrks_b_mean"),
            "Original public transportation score": old_score,
            "New public transportation score": new_score,
        })

    new_pt_ranks = rank_average(updated_scores, higher_is_better=True)
    for idx, r in enumerate(rank_rows):
        r[TRANSIT_RANK_COL] = new_pt_ranks[idx]

    factor_cols = [
        c for c in rank_headers
        if isinstance(c, str) and c.endswith(" rank") and c not in {"Avg-based rank", "Runoff rank"}
    ]
    if len(factor_cols) != 21:
        raise ValueError(f"Expected 21 direct rank factor columns, found {len(factor_cols)}")

    factor_array = np.array([[float(r[c]) for c in factor_cols] for r in rank_rows], dtype=np.float64)
    avg_ranks = factor_array.mean(axis=1)
    avg_based_ranks = rank_average(avg_ranks, higher_is_better=False)
    wins, avg_elim = run_randomized_runoff(factor_array, iterations, seed)

    sort_df = []
    for i, r in enumerate(rank_rows):
        sort_df.append((i, avg_elim[i], wins[i], avg_ranks[i], str(r["State"]), str(r["County"])))
    sort_df.sort(key=lambda x: (-x[1], -x[2], x[3], x[4], x[5]))
    runoff_rank = [0] * len(rank_rows)
    for pos, (i, *_rest) in enumerate(sort_df, 1):
        runoff_rank[i] = pos

    for i, r in enumerate(rank_rows):
        r["Avg Rank"] = round(float(avg_ranks[i]), 6)
        r["Avg Rank Python Check"] = round(float(avg_ranks[i]), 6)
        r["Avg Check Delta"] = 0.0
        r["Avg-based rank"] = float(avg_based_ranks[i])
        r["Runoff avg elimination round"] = float(avg_elim[i])
        r["Runoff wins"] = int(wins[i])
        r["Runoff win rate"] = float(wins[i] / iterations)
        r["Runoff rank"] = int(runoff_rank[i])

    rank_by_key = {(r["County"], r["State"]): r for r in rank_rows}
    for a in audit_rows:
        key = (a["County"], a["State"])
        old = original[key]
        new = rank_by_key[key]
        a.update({
            "Original public transportation rank": old["Original public transportation rank"],
            "New public transportation rank": new[TRANSIT_RANK_COL],
            "Public transportation rank delta": new[TRANSIT_RANK_COL] - old["Original public transportation rank"],
            "Original avg rank": old["Original avg rank"],
            "New avg rank": new["Avg Rank Python Check"],
            "Avg rank delta": new["Avg Rank Python Check"] - old["Original avg rank"],
            "Original runoff rank": old["Original runoff rank"],
            "New runoff rank": new["Runoff rank"],
            "Runoff rank delta": new["Runoff rank"] - old["Original runoff rank"],
            "Original runoff avg elimination round": old["Original runoff avg elimination round"],
            "New runoff avg elimination round": new["Runoff avg elimination round"],
            "Runoff avg elimination delta": new["Runoff avg elimination round"] - old["Original runoff avg elimination round"],
            "Original runoff wins": old["Original runoff wins"],
            "New runoff wins": new["Runoff wins"],
            "Original runoff win rate": old["Original runoff win rate"],
            "New runoff win rate": new["Runoff win rate"],
        })

    rank_rows_sorted = sorted(rank_rows, key=lambda r: r["Runoff rank"])
    value_rows_sorted = sorted(value_rows, key=lambda r: rank_by_key[(r["County"], r["State"])]["Runoff rank"])

    top_old20 = {k for k, v in original.items() if int(v["Original runoff rank"]) <= 20}
    top_new20 = {(r["County"], r["State"]) for r in rank_rows if int(r["Runoff rank"]) <= 20}
    top_old30 = {k for k, v in original.items() if int(v["Original runoff rank"]) <= 30}
    top_new30 = {(r["County"], r["State"]) for r in rank_rows if int(r["Runoff rank"]) <= 30}

    def mean_abs(field):
        return statistics.mean(abs(float(r[field])) for r in audit_rows)

    def median_abs(field):
        return statistics.median(abs(float(r[field])) for r in audit_rows)

    impact_rows = [
        ["Input workbook", input_workbook],
        ["EPA source ZIP", epa_transit_zip],
        ["DBF block-group records read", total_bg_records],
        ["County rows", len(rank_rows)],
        ["EPA DBF county matches", matched_count],
        ["Counties in missing fallback guardrail band", len(rank_rows) - matched_count],
        ["Missing-data policy", missing_policy],
        ["Public transportation scoring formula", "EPA raw composite = 100 × (0.50*TrAccess_I + 0.30*Pct_Jobs_byTr + 0.10*Pct_Pop_byTr + 0.10*Pct_Wrks_byTr). If --missing-policy guardrail_low_band, matched score = 1 + EPA raw composite; missing score = old modeled estimate rescaled to 0.001–0.999."],
        ["MCMC iterations", iterations],
        ["MCMC seed", seed],
        ["Top-20 runoff overlap vs prior", len(top_old20 & top_new20)],
        ["Top-30 runoff overlap vs prior", len(top_old30 & top_new30)],
        ["Mean abs public-transportation rank delta", mean_abs("Public transportation rank delta")],
        ["Median abs public-transportation rank delta", median_abs("Public transportation rank delta")],
        ["Mean abs Avg Rank delta", mean_abs("Avg rank delta")],
        ["Median abs Avg Rank delta", median_abs("Avg rank delta")],
        ["Mean abs runoff-rank delta", mean_abs("Runoff rank delta")],
        ["Median abs runoff-rank delta", median_abs("Runoff rank delta")],
        ["Largest absolute runoff-rank delta", max(abs(float(r["Runoff rank delta"])) for r in audit_rows)],
    ]

    if county_rollup_csv:
        with open(county_rollup_csv, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(audit_rows[0].keys()))
            writer.writeheader()
            writer.writerows(sorted(audit_rows, key=lambda r: r["New runoff rank"]))

    wb = load_workbook(input_workbook)
    if CalcProperties is not None:
        try:
            if wb.calculation is None:
                wb.calculation = CalcProperties()
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
        except Exception:
            pass

    write_matrix(wb[RANK_MATRIX_SHEET], [rank_headers] + [[clean_cell(r.get(h)) for h in rank_headers] for r in rank_rows_sorted])
    write_matrix(wb[VALUE_MATRIX_SHEET], [value_headers] + [[clean_cell(r.get(h)) for h in value_headers] for r in value_rows_sorted])
    write_matrix(wb[TOP30_SHEET], [rank_headers] + [[clean_cell(r.get(h)) for h in rank_headers] for r in rank_rows_sorted[:30]])

    comp_headers = ["Position", "County", "State", "Original runoff rank", "New runoff rank", "Runoff rank delta",
                    "Original runoff avg elimination round", "New runoff avg", "Original runoff wins", "New wins",
                    "Original avg rank", "New Avg Rank", "Avg Rank delta"]
    comp_rows = []
    for pos, r in enumerate(rank_rows_sorted[:30], 1):
        key = (r["County"], r["State"])
        old = original[key]
        comp_rows.append([pos, r["County"], r["State"], old["Original runoff rank"], r["Runoff rank"], r["Runoff rank"] - old["Original runoff rank"],
                          old["Original runoff avg elimination round"], r["Runoff avg elimination round"],
                          old["Original runoff wins"], r["Runoff wins"], old["Original avg rank"], r["Avg Rank Python Check"], r["Avg Rank Python Check"] - old["Original avg rank"]])
    write_matrix(wb[RESULT_COMPARISON_SHEET], [comp_headers] + comp_rows)

    ws_impact = recreate_sheet(wb, IMPACT_SHEET, after_name=RESULT_COMPARISON_SHEET)
    write_matrix(ws_impact, [["Metric", "Value"]] + impact_rows)
    ws_impact.cell(1, 4, "Top runoff movers")
    mover_headers = ["County", "State", "Original runoff rank", "New runoff rank", "Runoff rank delta",
                     "Original avg rank", "New avg rank", "Avg rank delta",
                     "Original public transportation rank", "New public transportation rank",
                     "Public transportation rank delta", "Transit update status"]
    movers = sorted(audit_rows, key=lambda r: abs(float(r["Runoff rank delta"])), reverse=True)[:15]
    for c, h in enumerate(mover_headers, 4):
        ws_impact.cell(2, c, h)
    for r_idx, row in enumerate(movers, 3):
        values = [row[h] for h in mover_headers]
        for c_idx, val in enumerate(values, 4):
            ws_impact.cell(r_idx, c_idx, val)

    ws_audit = recreate_sheet(wb, UPDATE_SHEET, after_name=IMPACT_SHEET)
    audit_headers = list(audit_rows[0].keys())
    audit_sorted = sorted(audit_rows, key=lambda r: r["New runoff rank"])
    write_matrix(ws_audit, [audit_headers] + [[clean_cell(r[h]) for h in audit_headers] for r in audit_sorted])

    run_rows = [
        ["Parameter", "Value", "Notes"],
        ["Iterations", iterations, "100k randomized runoff simulations rerun after public-transportation update."],
        ["Random seed", seed, "Same seed retained for comparability; deterministic with this script/version."],
        ["Rows/counties", len(rank_rows), "No counties added or removed."],
        ["Rank factors used", len(factor_cols), "Avg Rank excluded from random factor selection."],
        ["Column-selection rule", "Random permutation cycle", "All rank columns are picked exactly once per cycle before reshuffling."],
        ["Tie handling", "Random among tied worst ranks", "If multiple active counties tie for worst rank in selected factor."],
        ["Data import outcome", f"{matched_count} matched; {len(rank_rows)-matched_count} in missing fallback guardrail band", "EPA DBF coverage is limited to block groups in GTFS-covered regions."],
        ["Factor values changed?", "Public transportation only", "Walkability and all other factors preserved from input workbook."],
        ["Transit score formula", "Composite + guardrail", "EPA raw composite = 100 × (0.50*TrAccess_I + 0.30*Pct_Jobs_byTr + 0.10*Pct_Pop_byTr + 0.10*Pct_Wrks_byTr); with guardrail_low_band, matched score = 1 + raw composite and missing score = old modeled estimate rescaled to 0.001–0.999."],
        [None, None, None],
        ["Runoff rank", "County", "State", "Runoff avg elimination round", "Runoff wins", "Runoff win rate", "Avg Rank Python Check"],
    ]
    for r in rank_rows_sorted[:30]:
        run_rows.append([r["Runoff rank"], r["County"], r["State"], r["Runoff avg elimination round"], r["Runoff wins"], r["Runoff win rate"], r["Avg Rank Python Check"]])
    write_matrix(wb[RUNOFF_SHEET], run_rows)

    if METHODOLOGY_SHEET in wb.sheetnames:
        ws = wb[METHODOLOGY_SHEET]
        for row in ws.iter_rows(min_row=2):
            if str(row[0].value).strip().lower() == "public transportation":
                row[4].value = "Partially imported EPA transit accessibility DBF; missing EPA/GTFS counties placed in lower fallback band"
                row[5].value = ("Updated using EPA Access to Jobs and Workers via Transit DBF where county FIPS had block-group records. "
                                "County score is a simple mean of block-group composite: 100 × (0.50*TrAccess_Indexi + 0.30*Pct_Jobs_byTr + "
                                "0.10*Pct_Pop_byTr + 0.10*Pct_Wrks_byTr). For counties outside DBF coverage, prior modeled scores were rescaled to a lower fallback band below all EPA-measured counties, preserving relative order but enforcing the missing-data guardrail.")
                row[6].value = "https://www.epa.gov/smartgrowth/smart-location-mapping ; uploaded EPA_SLD_Trans45_DBF.zip ; uploaded sld_trans45_ug_0.pdf"
                break

    if LOG_SHEET in wb.sheetnames:
        ws = wb[LOG_SHEET]
        next_row = ws.max_row + 1
        log_values = [
            "EPA Access to Jobs and Workers via Transit DBF",
            "Public transportation",
            "Uploaded EPA_SLD_Trans45_DBF.zip containing SLD_Trans45.dbf plus user guide PDF",
            "Imported DBF; aggregated block-group metrics to county FIPS; computed composite score from TrAccess_Indexi, Pct_Jobs_byTr, Pct_Pop_byTr, Pct_Wrks_byTr",
            f"Updated {matched_count} EPA-measured counties; put {len(rank_rows)-matched_count} counties outside DBF coverage into lower fallback guardrail band; reran 100k MCMC",
            "https://www.epa.gov/smartgrowth/smart-location-mapping",
        ]
        for c, val in enumerate(log_values, 1):
            ws.cell(next_row, c, val)

    if "Sources Not Imported" in wb.sheetnames:
        ws = wb["Sources Not Imported"]
        rows = list(ws.iter_rows(values_only=True))
        filtered = [list(rows[0])] + [list(r) for r in rows[1:] if r and r[0] != "EPA Transit Access DBF/ZIP"]
        write_matrix(ws, filtered)

    if QA_SHEET in wb.sheetnames:
        qa_rows = [
            ["Check", "Value", "Result / Notes"],
            ["County rows", len(rank_rows), "Expected 413 based on input workbook."],
            ["Direct rank factor count", len(factor_cols), "Direct factor columns ending in ' rank', excluding Runoff rank and Avg-based rank."],
            ["EPA DBF block-group records read", total_bg_records, "Records in uploaded SLD_Trans45.dbf."],
            ["EPA DBF county matches", matched_count, "Counties with block-group transit metrics imported."],
            ["Counties in missing fallback guardrail band", len(rank_rows)-matched_count, "No EPA DBF coverage; prior modeled scores rescaled to lower band below all EPA-measured counties."],
            ["Runoff wins total", int(sum(r["Runoff wins"] for r in rank_rows)), f"Should equal iterations ({iterations})."],
            ["Runoff win rate sum", float(sum(r["Runoff win rate"] for r in rank_rows)), "Should equal 1.0, subject to floating-point precision."],
            ["Changed factor", "Public transportation only", "Other value/rank columns preserved from source workbook."],
        ]
        write_matrix(wb[QA_SHEET], qa_rows)

    for sheet_name in [RANK_MATRIX_SHEET, VALUE_MATRIX_SHEET, TOP30_SHEET, RESULT_COMPARISON_SHEET, RUNOFF_SHEET, QA_SHEET, "Sources Not Imported", IMPACT_SHEET, UPDATE_SHEET, METHODOLOGY_SHEET]:
        if sheet_name in wb.sheetnames:
            style_sheet(wb[sheet_name])

    Path(output_workbook).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_workbook)

    return {
        "matched_count": matched_count,
        "guardrail_fallback_count": len(rank_rows) - matched_count,
        "impact_rows": impact_rows,
        "top5": [(r["Runoff rank"], r["County"], r["State"], r["Avg Rank Python Check"], r["Runoff wins"]) for r in rank_rows_sorted[:5]],
    }


def parse_args():
    parser = argparse.ArgumentParser(description="Update county ranking workbook with EPA transit DBF data and rerun MCMC.")
    parser.add_argument("--input-workbook", required=True)
    parser.add_argument("--epa-transit-zip", required=True)
    parser.add_argument("--output-workbook", required=True)
    parser.add_argument("--county-rollup-csv", default=None)
    parser.add_argument("--fips-csv", default=None)
    parser.add_argument("--iterations", type=int, default=DEFAULT_ITERATIONS)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--missing-policy", choices=["guardrail_low_band", "retain_prior", "zero", "error"], default="guardrail_low_band")
    return parser.parse_args()


def main():
    args = parse_args()
    result = update_workbook(
        input_workbook=args.input_workbook,
        epa_transit_zip=args.epa_transit_zip,
        output_workbook=args.output_workbook,
        county_rollup_csv=args.county_rollup_csv,
        fips_csv=args.fips_csv,
        iterations=args.iterations,
        seed=args.seed,
        missing_policy=args.missing_policy,
    )
    print(f"Saved updated workbook: {args.output_workbook}")
    if args.county_rollup_csv:
        print(f"Saved county rollup CSV: {args.county_rollup_csv}")
    print(f"EPA DBF matches: {result['matched_count']}")
    print(f"Missing fallback guardrail band: {result['guardrail_fallback_count']}")
    print("Top 5 after update:")
    for row in result["top5"]:
        print(row)


if __name__ == "__main__":
    main()
