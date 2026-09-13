#!/usr/bin/env python3
"""
Update the county-rankings workbook with NOAA/NCEI daily temperature normals,
re-rank temperature factors, recompute average ranks, rerun the 100k randomized
runoff MCMC, and write an updated workbook plus an audit CSV.

Expected inputs
---------------
1. Existing county-rankings workbook with sheets:
   - Rank Matrix
   - Value Matrix
   - Top 30
   - Result Comparison
   - Runoff Simulation
   - Sources & Methodology
   - Data Acquisition Log
   - QA Checks
   - EPA Walkability Update, used for county FIPS mapping
2. NOAA/NCEI daily temperature normals tar.gz containing:
   - dly_inventory.txt
   - dly-temp-normal.csv
   - dly-temp-stddev.csv
3. A county-centroid source. Preferred: the EPA WalkabilityIndex ZIP from the
   previous workflow, because it includes nationwide block-group geometries and
   population. The script computes population-weighted county centroids from it.
   Alternatively pass a precomputed county centroid CSV with columns:
   fips5, county_x, county_y, crs_wkt.

Method
------
- Reads daily TMAX normals and daily TMAX standard deviations by station.
- Estimates expected annual days with TMAX >= 90 F and TMAX < 50 F by summing
  daily exceedance probabilities. This is better than hard-counting days whose
  *normal* TMAX crosses a threshold, because real daily temperatures vary around
  the daily normal.
- Computes population-weighted county centroids from EPA block-group geometry.
- Assigns station metrics to counties using inverse-distance-squared blending of
  the closest NOAA stations. Default: up to five stations within 125 km; if fewer
  than two stations are found within 125 km, use the five nearest stations and
  flag the county as a nearest-station fallback.
- Missing-data policy: do not force missing counties above or below measured
  counties. Temperature missingness is not inherently evidence of better/worse
  climate. If a county lacks a centroid/station match, the script falls back to
  old estimates calibrated to NOAA-measured ratios using state/division/national
  medians. In the current 413-county workbook, all counties receive station-
  backed values, so no calibrated fallback is used.
- Updates:
    * Value Matrix -> Est. 90F+ days/year
    * Value Matrix -> Est. days below 50F / year
    * Rank Matrix -> 90F+ days rank
    * Rank Matrix -> Days below 50F rank
  Both ranks are lower-is-better.
- Reruns randomized runoff with the same rank-matrix rules: random permutation
  cycles over rank factors; the worst rank on the selected factor is eliminated;
  tied worst rows are broken randomly.

Install dependencies
--------------------
pip install pandas numpy scipy pyproj pyogrio openpyxl numba

Example
-------
python update_noaa_temperature_and_rerun_mcmc.py \
  --input-workbook county_rankings_epa_walkability_transit_aqi_100k.xlsx \
  --noaa-tar us-climate-normals_2006-2020_v1.0.1_daily_temperature_by-variable_c20230403.tar.gz \
  --epa-walkability-zip EPA_WalkabilityIndex.zip \
  --output-workbook county_rankings_epa_walkability_transit_aqi_noaa_temperature_100k.xlsx \
  --county-rollup-csv noaa_temperature_county_rollup.csv
"""

from __future__ import annotations

import argparse
import csv
import math
import os
import re
import statistics
import tarfile
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

import numpy as np
import pandas as pd
from numba import njit
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
try:
    from openpyxl.workbook.properties import CalcProperties
except Exception:  # pragma: no cover
    CalcProperties = None
from openpyxl.utils import get_column_letter
from pyproj import Transformer
from scipy.spatial import cKDTree
from scipy.special import ndtr
import pyogrio

DEFAULT_ITERATIONS = 100_000
DEFAULT_SEED = 2026060901

RANK_MATRIX_SHEET = "Rank Matrix"
VALUE_MATRIX_SHEET = "Value Matrix"
TOP30_SHEET = "Top 30"
RESULT_COMPARISON_SHEET = "Result Comparison"
RUNOFF_SHEET = "Runoff Simulation"
METHODOLOGY_SHEET = "Sources & Methodology"
LOG_SHEET = "Data Acquisition Log"
QA_SHEET = "QA Checks"
WALKABILITY_AUDIT_SHEET = "EPA Walkability Update"
IMPACT_SHEET = "NOAA Temperature Impact"
UPDATE_SHEET = "NOAA Temperature Update"

HOT_VALUE_COL = "Est. 90F+ days/year"
HOT_RANK_COL = "90F+ days rank"
COLD_VALUE_COL = "Est. days below 50F / year"
COLD_RANK_COL = "Days below 50F rank"

HEADER_FILL = "1F4E79"
BORDER_COLOR = "B7B7B7"

STATE_DIVISION = {
    "CT": "New England", "ME": "New England", "MA": "New England", "NH": "New England", "RI": "New England", "VT": "New England",
    "NJ": "Middle Atlantic", "NY": "Middle Atlantic", "PA": "Middle Atlantic",
    "IL": "East North Central", "IN": "East North Central", "MI": "East North Central", "OH": "East North Central", "WI": "East North Central",
    "IA": "West North Central", "KS": "West North Central", "MN": "West North Central", "MO": "West North Central", "NE": "West North Central", "ND": "West North Central", "SD": "West North Central",
    "DE": "South Atlantic", "DC": "South Atlantic", "FL": "South Atlantic", "GA": "South Atlantic", "MD": "South Atlantic", "NC": "South Atlantic", "SC": "South Atlantic", "VA": "South Atlantic", "WV": "South Atlantic",
    "AL": "East South Central", "KY": "East South Central", "MS": "East South Central", "TN": "East South Central",
    "AR": "West South Central", "LA": "West South Central", "OK": "West South Central", "TX": "West South Central",
    "AZ": "Mountain", "CO": "Mountain", "ID": "Mountain", "MT": "Mountain", "NV": "Mountain", "NM": "Mountain", "UT": "Mountain", "WY": "Mountain",
    "AK": "Pacific", "CA": "Pacific", "HI": "Pacific", "OR": "Pacific", "WA": "Pacific",
}


def normalize_county_name(name: str) -> str:
    if name is None:
        return ""
    s = str(name).strip()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.upper().replace("ST.", "ST").replace("SAINT ", "ST ")
    s = s.replace("'", "")
    s = re.sub(r"\s*\(.*?\)", "", s)
    suffixes = [
        " COUNTY AND BOROUGH", " CITY AND BOROUGH", " CENSUS AREA",
        " MUNICIPALITY", " CONSOLIDATED GOVERNMENT", " UNIFIED GOVERNMENT",
        " BOROUGH", " PARISH", " COUNTY", " CITY",
    ]
    changed = True
    while changed:
        changed = False
        for suffix in suffixes:
            if s.endswith(suffix):
                s = s[: -len(suffix)]
                changed = True
                break
    s = re.sub(r"[^A-Z0-9]+", " ", s).strip()
    return re.sub(r"\s+", " ", s)


def safe_float(x):
    if x is None or x == "":
        return None
    try:
        return float(x)
    except Exception:
        return None


def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = list(rows[0])
    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        padded = list(row) + [None] * (len(headers) - len(row))
        data.append(dict(zip(headers, padded[: len(headers)])))
    return headers, data


def write_matrix(ws, matrix: List[List[object]]):
    ws.delete_rows(1, ws.max_row)
    for r_idx, row in enumerate(matrix, 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, value)


def recreate_sheet(wb, name: str, after_name: str | None = None):
    if name in wb.sheetnames:
        del wb[name]
    if after_name and after_name in wb.sheetnames:
        idx = wb.sheetnames.index(after_name)
        return wb.create_sheet(name, idx + 1)
    return wb.create_sheet(name)


def style_sheet(ws, max_width: int = 42):
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color=BORDER_COLOR)
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 8), max_width)


def get_fips_for_workbook(wb) -> Dict[Tuple[str, str], str]:
    if WALKABILITY_AUDIT_SHEET not in wb.sheetnames:
        raise ValueError(f"Workbook must contain {WALKABILITY_AUDIT_SHEET!r} for County FIPS mapping")
    headers, rows = read_sheet(wb[WALKABILITY_AUDIT_SHEET])
    if not {"County", "State", "County FIPS"}.issubset(set(headers)):
        raise ValueError(f"{WALKABILITY_AUDIT_SHEET!r} must include County, State, and County FIPS columns")
    return {
        (r["County"], r["State"]): str(r["County FIPS"]).zfill(5)
        for r in rows
        if r.get("County") is not None and r.get("State") is not None and r.get("County FIPS") is not None
    }


def rank_average(values: Iterable[float], higher_is_better: bool) -> List[float]:
    pairs = [(float(v), i) for i, v in enumerate(values)]
    pairs.sort(key=lambda t: t[0], reverse=higher_is_better)
    ranks = [0.0] * len(pairs)
    i = 0
    while i < len(pairs):
        j = i + 1
        while j < len(pairs) and pairs[j][0] == pairs[i][0]:
            j += 1
        avg_rank = (i + 1 + j) / 2.0
        for k in range(i, j):
            ranks[pairs[k][1]] = avg_rank
        i = j
    return ranks


@njit
def run_randomized_runoff(ranks: np.ndarray, sorted_order: np.ndarray, iterations: int, seed: int):
    """Fast randomized runoff using pre-sorted row orders for each factor.

    This is equivalent to scanning active rows for the worst rank on each selected
    factor, but avoids O(n^2) scans by maintaining a per-factor pointer to the
    worst still-active row in that factor's sorted order.
    """
    np.random.seed(seed)
    n, f = ranks.shape
    wins = np.zeros(n, np.int64)
    elim_sum = np.zeros(n, np.float64)
    factor_order = np.empty(f, np.int64)
    active = np.empty(n, np.bool_)
    ptr = np.empty(f, np.int64)
    ties = np.empty(n, np.int64)

    for _ in range(iterations):
        for i in range(n):
            active[i] = True
        for j in range(f):
            ptr[j] = 0
        active_count = n
        pos = f
        round_no = 0

        while active_count > 1:
            if pos >= f:
                for j in range(f):
                    factor_order[j] = j
                for j in range(f - 1, 0, -1):
                    k = np.random.randint(j + 1)
                    tmp = factor_order[j]
                    factor_order[j] = factor_order[k]
                    factor_order[k] = tmp
                pos = 0

            col = factor_order[pos]
            pos += 1

            p = ptr[col]
            while p < n and not active[sorted_order[col, p]]:
                p += 1
            ptr[col] = p

            first_idx = sorted_order[col, p]
            worst_rank = ranks[first_idx, col]
            tie_count = 0
            q = p
            while q < n:
                row_idx = sorted_order[col, q]
                if ranks[row_idx, col] != worst_rank:
                    break
                if active[row_idx]:
                    ties[tie_count] = row_idx
                    tie_count += 1
                q += 1

            loser = ties[np.random.randint(tie_count)]
            active[loser] = False
            round_no += 1
            elim_sum[loser] += round_no
            active_count -= 1

        for i in range(n):
            if active[i]:
                winner = i
                break
        wins[winner] += 1
        elim_sum[winner] += n

    return wins, elim_sum / iterations

def clean_cell(value):
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return None if math.isnan(float(value)) else float(value)
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def mean_abs(rows, col):
    vals = [abs(float(r[col])) for r in rows if r.get(col) is not None]
    return statistics.mean(vals) if vals else None


def median_abs(rows, col):
    vals = [abs(float(r[col])) for r in rows if r.get(col) is not None]
    return statistics.median(vals) if vals else None


def parse_noaa_station_metrics(noaa_tar: str):
    """Return station-level expected annual hot/cold day counts and read stats."""
    with tarfile.open(noaa_tar, "r:gz") as tar:
        names = set(tar.getnames())
        required = {"dly_inventory.txt", "dly-temp-normal.csv", "dly-temp-stddev.csv"}
        missing = required - names
        if missing:
            raise ValueError(f"NOAA tar is missing required files: {sorted(missing)}")
        temp = pd.read_csv(
            tar.extractfile("dly-temp-normal.csv"),
            usecols=["GHCN_ID", "month", "day", "DLY-TMAX-NORMAL", "DLY-TAVG-NORMAL"],
        )
    with tarfile.open(noaa_tar, "r:gz") as tar:
        std = pd.read_csv(
            tar.extractfile("dly-temp-stddev.csv"),
            usecols=["GHCN_ID", "month", "day", "DLY-TMAX-STDDEV", "DLY-TAVG-STDDEV"],
        )

    normal_rows = len(temp)
    std_rows = len(std)
    df = temp.merge(std, on=["GHCN_ID", "month", "day"], how="inner")
    for col in ["DLY-TMAX-NORMAL", "DLY-TAVG-NORMAL", "DLY-TMAX-STDDEV", "DLY-TAVG-STDDEV"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df = df.dropna(subset=["DLY-TMAX-NORMAL", "DLY-TMAX-STDDEV"])

    sd_tmax = df["DLY-TMAX-STDDEV"].clip(lower=0.1)
    df["p_hot90"] = 1.0 - ndtr((90.0 - df["DLY-TMAX-NORMAL"]) / sd_tmax)
    df["p_cold_tmax50"] = ndtr((50.0 - df["DLY-TMAX-NORMAL"]) / sd_tmax)

    # TAVG version is retained in the audit file for diagnostics only. The workbook
    # cold factor uses TMAX < 50 because the existing model's scale matches days whose
    # daytime high stays below 50 F better than days whose mean temperature is below 50 F.
    sd_tavg = df["DLY-TAVG-STDDEV"].clip(lower=0.1)
    df["p_cold_tavg50"] = ndtr((50.0 - df["DLY-TAVG-NORMAL"]) / sd_tavg)

    agg = df.groupby("GHCN_ID", as_index=False).agg(
        n_days=("DLY-TMAX-NORMAL", "size"),
        noaa_hot90_days=("p_hot90", "sum"),
        noaa_cold_tmax_below50_days=("p_cold_tmax50", "sum"),
        diagnostic_cold_tavg_below50_days=("p_cold_tavg50", "sum"),
        mean_tmax_normal=("DLY-TMAX-NORMAL", "mean"),
    )
    agg = agg[agg["n_days"] >= 350].copy()
    scale = 365.25 / agg["n_days"]
    for col in ["noaa_hot90_days", "noaa_cold_tmax_below50_days", "diagnostic_cold_tavg_below50_days"]:
        agg[col] = agg[col] * scale

    inv_rows = []
    with tarfile.open(noaa_tar, "r:gz") as tar:
        for b in tar.extractfile("dly_inventory.txt"):
            line = b.decode("utf-8", errors="ignore").rstrip("\n")
            sid = line[0:11].strip()
            try:
                lat = float(line[12:20])
                lon = float(line[21:30])
                elev_m = float(line[31:37])
            except Exception:
                continue
            state = line[38:40].strip()
            name = line[41:71].strip()
            inv_rows.append((sid, lat, lon, elev_m, state, name))
    inv = pd.DataFrame(inv_rows, columns=["GHCN_ID", "lat", "lon", "elev_m", "station_state", "station_name"])
    stations = agg.merge(inv, on="GHCN_ID", how="inner")
    return stations, {
        "normal_rows": normal_rows,
        "stddev_rows": std_rows,
        "merged_daily_rows": len(df),
        "stations_with_temperature_metrics": len(agg),
        "stations_with_inventory": len(stations),
    }


def compute_county_centroids(
    workbook_fips: Iterable[str],
    epa_walkability_zip: str | None = None,
    county_centroids_csv: str | None = None,
):
    """Read or compute projected population-weighted county centroids."""
    if county_centroids_csv and Path(county_centroids_csv).exists():
        centroids = pd.read_csv(county_centroids_csv, dtype={"fips5": str})
        required = {"fips5", "county_x", "county_y", "crs_wkt"}
        missing = required - set(centroids.columns)
        if missing:
            raise ValueError(f"County centroids CSV missing required columns: {sorted(missing)}")
        return centroids

    if not epa_walkability_zip or not Path(epa_walkability_zip).exists():
        raise ValueError(
            "Need either --county-centroids-csv or --epa-walkability-zip to compute county centroids. "
            "Use EPA_WalkabilityIndex.zip from the prior workflow."
        )

    vsi_path = f"/vsizip/{Path(epa_walkability_zip).resolve()}/Natl_WI.gdb"
    gdf = pyogrio.read_dataframe(
        vsi_path,
        layer="NationalWalkabilityIndex",
        columns=["STATEFP", "COUNTYFP", "TotPop"],
        read_geometry=True,
    )
    crs_wkt = str(gdf.crs)
    target_fips = set(str(f).zfill(5) for f in workbook_fips if f is not None)
    gdf["fips5"] = gdf["STATEFP"].astype(str).str.zfill(2) + gdf["COUNTYFP"].astype(str).str.zfill(3)
    gdf = gdf[gdf["fips5"].isin(target_fips)].copy()
    gdf["TotPop"] = pd.to_numeric(gdf["TotPop"], errors="coerce").fillna(0)
    cent = gdf.geometry.centroid
    tmp = pd.DataFrame({
        "fips5": gdf["fips5"].values,
        "pop": gdf["TotPop"].values,
        "x": cent.x,
        "y": cent.y,
    })
    rows = []
    for fips, grp in tmp.groupby("fips5"):
        weights = grp["pop"].clip(lower=0).to_numpy()
        if weights.sum() > 0:
            x = float(np.average(grp["x"].to_numpy(), weights=weights))
            y = float(np.average(grp["y"].to_numpy(), weights=weights))
        else:
            x = float(grp["x"].mean())
            y = float(grp["y"].mean())
        rows.append({
            "fips5": str(fips).zfill(5),
            "county_x": x,
            "county_y": y,
            "bg_count": int(len(grp)),
            "population_used": float(weights.sum()),
            "crs_wkt": crs_wkt,
        })
    centroids = pd.DataFrame(rows)
    if county_centroids_csv:
        Path(county_centroids_csv).parent.mkdir(parents=True, exist_ok=True)
        centroids.to_csv(county_centroids_csv, index=False)
    return centroids


def build_noaa_county_rollup(stations: pd.DataFrame, centroids: pd.DataFrame, k: int = 5, radius_km: float = 125.0):
    crs_wkt = centroids["crs_wkt"].iloc[0]
    transformer = Transformer.from_crs("EPSG:4269", crs_wkt, always_xy=True)
    x, y = transformer.transform(stations["lon"].to_numpy(), stations["lat"].to_numpy())
    stations = stations.copy()
    stations["x"] = x
    stations["y"] = y
    stations = stations[np.isfinite(stations["x"]) & np.isfinite(stations["y"])].copy()
    if stations.empty:
        raise ValueError("No finite NOAA station coordinates after projection")

    tree = cKDTree(stations[["x", "y"]].to_numpy())
    query_k = min(max(k * 2, k), len(stations))
    dists_m, idxs = tree.query(centroids[["county_x", "county_y"]].to_numpy(), k=query_k)
    if query_k == 1:
        dists_m = dists_m[:, None]
        idxs = idxs[:, None]

    rows = []
    for ci, county in centroids.reset_index(drop=True).iterrows():
        ds_km = dists_m[ci] / 1000.0
        ids = idxs[ci]
        nearby_positions = np.where(ds_km <= radius_km)[0]
        if len(nearby_positions) >= 2:
            use_pos = nearby_positions[:k]
            status = "NOAA nearby station blend"
            imputation_method = "inverse-distance-squared blend of nearest NOAA stations within 125 km"
        else:
            use_pos = np.arange(min(k, len(ids)))
            status = "NOAA nearest-station fallback"
            imputation_method = "fewer than two stations within 125 km; used inverse-distance-squared blend of nearest NOAA stations"
        dsel = np.maximum(ds_km[use_pos], 1.0)
        weights = 1.0 / (dsel ** 2)
        weights = weights / weights.sum()
        sub = stations.iloc[ids[use_pos]].copy()
        rows.append({
            "County FIPS": str(county["fips5"]).zfill(5),
            "Climate update status": status,
            "Imputation method": imputation_method,
            "Stations used": int(len(use_pos)),
            "Nearest station km": float(ds_km[0]),
            "Max station km used": float(ds_km[use_pos].max()),
            "NOAA 90F+ days/year": float(np.sum(sub["noaa_hot90_days"].to_numpy() * weights)),
            "NOAA days below 50F/year": float(np.sum(sub["noaa_cold_tmax_below50_days"].to_numpy() * weights)),
            "Diagnostic NOAA TAVG<50 days/year": float(np.sum(sub["diagnostic_cold_tavg_below50_days"].to_numpy() * weights)),
            "Station IDs used": ";".join(sub["GHCN_ID"].astype(str).tolist()),
            "Station names used": ";".join(sub["station_name"].astype(str).tolist()),
            "Station weights": ";".join(f"{w:.4f}" for w in weights),
        })
    return pd.DataFrame(rows), {
        "finite_station_count": len(stations),
        "radius_km": radius_km,
        "k": k,
    }


def median_or_none(vals):
    vals = [float(v) for v in vals if v is not None and math.isfinite(float(v))]
    return statistics.median(vals) if vals else None


def calibrated_fallbacks(value_rows, measured_by_key, col):
    ratios_by_state = defaultdict(list)
    ratios_by_division = defaultdict(list)
    national_ratios = []
    national_vals = []
    vals_by_state = defaultdict(list)
    vals_by_division = defaultdict(list)
    for r in value_rows:
        key = (r["County"], r["State"])
        old = safe_float(r[col])
        if key not in measured_by_key or old is None or old <= 0:
            continue
        new = float(measured_by_key[key])
        ratio = new / old
        state = str(r["State"]).upper()
        division = STATE_DIVISION.get(state, "Unknown")
        ratios_by_state[state].append(ratio)
        ratios_by_division[division].append(ratio)
        national_ratios.append(ratio)
        vals_by_state[state].append(new)
        vals_by_division[division].append(new)
        national_vals.append(new)
    national_ratio = median_or_none(national_ratios) or 1.0
    state_ratio = {s: median_or_none(v) for s, v in ratios_by_state.items()}
    division_ratio = {d: median_or_none(v) for d, v in ratios_by_division.items()}
    state_n = {s: len(v) for s, v in ratios_by_state.items()}
    division_n = {d: len(v) for d, v in ratios_by_division.items()}
    national_median = median_or_none(national_vals) or 0.0
    fallbacks = {}
    for r in value_rows:
        key = (r["County"], r["State"])
        if key in measured_by_key:
            continue
        old = safe_float(r[col])
        state = str(r["State"]).upper()
        division = STATE_DIVISION.get(state, "Unknown")
        sr = state_ratio.get(state)
        dr = division_ratio.get(division)
        sn = state_n.get(state, 0)
        dn = division_n.get(division, 0)
        if sr is not None and sn >= 5:
            ratio = sr
            method = f"state calibrated old estimate; state_n={sn}"
        elif sr is not None and dr is not None:
            w_state = sn / (sn + 5.0)
            w_div = 0.50 * (1.0 - w_state)
            w_nat = 1.0 - w_state - w_div
            ratio = w_state * sr + w_div * dr + w_nat * national_ratio
            method = f"blended state/division/national calibrated old estimate; state_n={sn}; division_n={dn}"
        elif dr is not None:
            ratio = 0.70 * dr + 0.30 * national_ratio
            method = f"division/national calibrated old estimate; division_n={dn}"
        else:
            ratio = national_ratio
            method = "national calibrated old estimate"
        imputed = max(0.0, (old or national_median) * ratio)
        fallbacks[key] = (imputed, method)
    return fallbacks


def update_workbook(
    input_workbook: str,
    noaa_tar: str,
    output_workbook: str,
    county_rollup_csv: str | None,
    epa_walkability_zip: str | None,
    county_centroids_csv: str | None,
    iterations: int,
    seed: int,
):
    wb = load_workbook(input_workbook)
    if CalcProperties is not None:
        try:
            wb.calculation = CalcProperties(calcMode="auto")
        except Exception:
            pass

    rank_headers, rank_rows = read_sheet(wb[RANK_MATRIX_SHEET])
    value_headers, value_rows = read_sheet(wb[VALUE_MATRIX_SHEET])
    for col in [HOT_VALUE_COL, COLD_VALUE_COL]:
        if col not in value_headers:
            raise ValueError(f"Missing value column: {col}")
    for col in [HOT_RANK_COL, COLD_RANK_COL]:
        if col not in rank_headers:
            raise ValueError(f"Missing rank column: {col}")

    fips_lookup = get_fips_for_workbook(wb)
    for r in value_rows:
        r["_fips5"] = fips_lookup.get((r["County"], r["State"]))

    original = {}
    for r in rank_rows:
        key = (r["County"], r["State"])
        original[key] = {
            "Original runoff rank": r["Runoff rank"],
            "Original avg rank": r["Avg Rank Python Check"],
            "Original 90F+ days rank": r[HOT_RANK_COL],
            "Original days below 50F rank": r[COLD_RANK_COL],
            "Original runoff avg elimination round": r["Runoff avg elimination round"],
            "Original runoff wins": r["Runoff wins"],
            "Original runoff win rate": r["Runoff win rate"],
        }
    old_values_by_key = {
        (r["County"], r["State"]): {
            "Old modeled 90F+ days": r[HOT_VALUE_COL],
            "Old modeled days below 50F": r[COLD_VALUE_COL],
        }
        for r in value_rows
    }

    stations, station_stats = parse_noaa_station_metrics(noaa_tar)
    unique_fips = sorted({r["_fips5"] for r in value_rows if r.get("_fips5")})
    centroids = compute_county_centroids(unique_fips, epa_walkability_zip, county_centroids_csv)
    county_rollup, rollup_stats = build_noaa_county_rollup(stations, centroids)
    rollup_by_fips = {str(r["County FIPS"]).zfill(5): r for r in county_rollup.to_dict("records")}

    noaa_by_key = {}
    for r in value_rows:
        key = (r["County"], r["State"])
        fips = r.get("_fips5")
        rec = rollup_by_fips.get(str(fips).zfill(5)) if fips else None
        if rec:
            noaa_by_key[key] = rec

    # Fallback logic should rarely trigger. It is kept for reproducibility if future
    # workbook rows lack FIPS/centroid coverage.
    measured_hot = {k: v["NOAA 90F+ days/year"] for k, v in noaa_by_key.items()}
    measured_cold = {k: v["NOAA days below 50F/year"] for k, v in noaa_by_key.items()}
    hot_fallbacks = calibrated_fallbacks(value_rows, measured_hot, HOT_VALUE_COL)
    cold_fallbacks = calibrated_fallbacks(value_rows, measured_cold, COLD_VALUE_COL)

    update_info = {}
    for r in value_rows:
        key = (r["County"], r["State"])
        if key in noaa_by_key:
            rec = dict(noaa_by_key[key])
            rec["NOAA 90F+ days/year"] = float(rec["NOAA 90F+ days/year"])
            rec["NOAA days below 50F/year"] = float(rec["NOAA days below 50F/year"])
            update_info[key] = rec
        else:
            hot_val, hot_method = hot_fallbacks.get(key, (safe_float(r[HOT_VALUE_COL]) or 0.0, "unmodified old estimate fallback"))
            cold_val, cold_method = cold_fallbacks.get(key, (safe_float(r[COLD_VALUE_COL]) or 0.0, "unmodified old estimate fallback"))
            update_info[key] = {
                "County FIPS": r.get("_fips5"),
                "Climate update status": "calibrated fallback; no NOAA station/centroid match",
                "Imputation method": f"hot: {hot_method}; cold: {cold_method}",
                "Stations used": 0,
                "Nearest station km": None,
                "Max station km used": None,
                "NOAA 90F+ days/year": float(hot_val),
                "NOAA days below 50F/year": float(cold_val),
                "Diagnostic NOAA TAVG<50 days/year": None,
                "Station IDs used": "",
                "Station names used": "",
                "Station weights": "",
            }

    # Update values.
    for r in value_rows:
        key = (r["County"], r["State"])
        r[HOT_VALUE_COL] = float(update_info[key]["NOAA 90F+ days/year"])
        r[COLD_VALUE_COL] = float(update_info[key]["NOAA days below 50F/year"])

    # Re-rank both temperature factors, lower is better.
    hot_ranks = rank_average([r[HOT_VALUE_COL] for r in value_rows], higher_is_better=False)
    cold_ranks = rank_average([r[COLD_VALUE_COL] for r in value_rows], higher_is_better=False)
    ranks_by_key = {}
    for idx, r in enumerate(value_rows):
        ranks_by_key[(r["County"], r["State"])] = (hot_ranks[idx], cold_ranks[idx])
    for r in rank_rows:
        key = (r["County"], r["State"])
        r[HOT_RANK_COL], r[COLD_RANK_COL] = ranks_by_key[key]

    factor_cols = [
        c for c in rank_headers
        if isinstance(c, str) and c.endswith(" rank") and c not in {"Avg-based rank", "Runoff rank"}
    ]
    if len(factor_cols) != 21:
        raise ValueError(f"Expected 21 direct rank factor columns, found {len(factor_cols)}")

    factor_array = np.array([[float(r[c]) for c in factor_cols] for r in rank_rows], dtype=np.float64)
    avg_ranks = factor_array.mean(axis=1)
    avg_based_ranks = rank_average(avg_ranks, higher_is_better=False)
    sorted_order = np.argsort(-factor_array, axis=0).T.astype(np.int64)
    wins, avg_elim = run_randomized_runoff(factor_array, sorted_order, iterations, seed)

    sort_df = []
    for i, r in enumerate(rank_rows):
        sort_df.append((i, avg_elim[i], wins[i], avg_ranks[i], str(r["State"]), str(r["County"])))
    sort_df.sort(key=lambda t: (-t[1], -t[2], t[3], t[4], t[5]))
    runoff_rank = [0] * len(rank_rows)
    for pos, (i, *_rest) in enumerate(sort_df, 1):
        runoff_rank[i] = pos

    for i, r in enumerate(rank_rows):
        r["Avg Rank"] = round(float(avg_ranks[i]), 6)
        r["Avg Rank Python Check"] = round(float(avg_ranks[i]), 6)
        r["Avg Check Delta"] = 0.0
        r["Avg-based rank"] = float(avg_based_ranks[i])
        r["Runoff avg elimination round"] = round(float(avg_elim[i]), 6)
        r["Runoff wins"] = int(wins[i])
        r["Runoff win rate"] = float(wins[i]) / float(iterations)
        r["Runoff rank"] = int(runoff_rank[i])

    rank_by_key = {(r["County"], r["State"]): r for r in rank_rows}
    audit_rows = []
    for vr in value_rows:
        key = (vr["County"], vr["State"])
        old = original[key]
        new = rank_by_key[key]
        info = update_info[key]
        old_vals = old_values_by_key[key]
        audit_rows.append({
            "County": key[0],
            "State": key[1],
            "County FIPS": info.get("County FIPS") or vr.get("_fips5"),
            "Climate update status": info.get("Climate update status"),
            "Old modeled 90F+ days": old_vals["Old modeled 90F+ days"],
            "New NOAA 90F+ days": vr[HOT_VALUE_COL],
            "Original 90F+ days rank": old["Original 90F+ days rank"],
            "New 90F+ days rank": new[HOT_RANK_COL],
            "90F+ days rank delta": new[HOT_RANK_COL] - old["Original 90F+ days rank"],
            "Old modeled days below 50F": old_vals["Old modeled days below 50F"],
            "New NOAA days below 50F": vr[COLD_VALUE_COL],
            "Original days below 50F rank": old["Original days below 50F rank"],
            "New days below 50F rank": new[COLD_RANK_COL],
            "Days below 50F rank delta": new[COLD_RANK_COL] - old["Original days below 50F rank"],
            "Stations used": info.get("Stations used"),
            "Nearest station km": info.get("Nearest station km"),
            "Max station km used": info.get("Max station km used"),
            "Station IDs used": info.get("Station IDs used"),
            "Station names used": info.get("Station names used"),
            "Station weights": info.get("Station weights"),
            "Diagnostic NOAA TAVG<50 days/year": info.get("Diagnostic NOAA TAVG<50 days/year"),
            "Imputation method": info.get("Imputation method"),
            "Original avg rank": old["Original avg rank"],
            "New avg rank": new["Avg Rank Python Check"],
            "Avg rank delta": new["Avg Rank Python Check"] - old["Original avg rank"],
            "Original runoff rank": old["Original runoff rank"],
            "New runoff rank": new["Runoff rank"],
            "Runoff rank delta": new["Runoff rank"] - old["Original runoff rank"],
            "Original runoff avg elimination round": old["Original runoff avg elimination round"],
            "New runoff avg elimination round": new["Runoff avg elimination round"],
            "Runoff avg elimination delta": new["Runoff avg elimination round"] - old["Original runoff avg elimination round"],
            "Original runoff wins": old["Original runoff wins"],
            "New runoff wins": new["Runoff wins"],
            "Original runoff win rate": old["Original runoff win rate"],
            "New runoff win rate": new["Runoff win rate"],
        })

    # Sort outputs by new runoff rank.
    rank_rows_sorted = sorted(rank_rows, key=lambda r: r["Runoff rank"])
    value_rows_sorted = sorted(value_rows, key=lambda r: rank_by_key[(r["County"], r["State"])] ["Runoff rank"])
    audit_sorted = sorted(audit_rows, key=lambda r: r["New runoff rank"])

    top_old20 = {k for k, v in original.items() if int(v["Original runoff rank"]) <= 20}
    top_new20 = {(r["County"], r["State"]) for r in rank_rows if int(r["Runoff rank"]) <= 20}
    top_old30 = {k for k, v in original.items() if int(v["Original runoff rank"]) <= 30}
    top_new30 = {(r["County"], r["State"]) for r in rank_rows if int(r["Runoff rank"]) <= 30}
    station_backed = sum(1 for r in audit_rows if str(r["Climate update status"]).startswith("NOAA"))
    fallback_count = len(audit_rows) - station_backed
    nearest_vals = [float(r["Nearest station km"]) for r in audit_rows if r.get("Nearest station km") is not None]
    max_used_vals = [float(r["Max station km used"]) for r in audit_rows if r.get("Max station km used") is not None]

    impact_rows = [
        ["NOAA normal daily rows read", station_stats["normal_rows"]],
        ["NOAA stddev daily rows read", station_stats["stddev_rows"]],
        ["Merged daily station rows used", station_stats["merged_daily_rows"]],
        ["NOAA stations with usable temperature metrics", station_stats["stations_with_temperature_metrics"]],
        ["NOAA stations with inventory/location", station_stats["stations_with_inventory"]],
        ["County rows", len(rank_rows)],
        ["Counties station-backed", station_backed],
        ["Counties using calibrated fallback", fallback_count],
        ["Nearest station distance range km", f"{min(nearest_vals):.3f} to {max(nearest_vals):.3f}" if nearest_vals else "n/a"],
        ["Max station distance used range km", f"{min(max_used_vals):.3f} to {max(max_used_vals):.3f}" if max_used_vals else "n/a"],
        ["Station rollup rule", f"Up to {rollup_stats['k']} stations within {rollup_stats['radius_km']} km, inverse-distance-squared weights"],
        ["Top-20 runoff overlap vs prior", f"{len(top_old20 & top_new20)} / 20"],
        ["Top-30 runoff overlap vs prior", f"{len(top_old30 & top_new30)} / 30"],
        ["Mean abs 90F+ rank delta", round(mean_abs(audit_rows, "90F+ days rank delta"), 6)],
        ["Median abs 90F+ rank delta", round(median_abs(audit_rows, "90F+ days rank delta"), 6)],
        ["Mean abs below-50F rank delta", round(mean_abs(audit_rows, "Days below 50F rank delta"), 6)],
        ["Median abs below-50F rank delta", round(median_abs(audit_rows, "Days below 50F rank delta"), 6)],
        ["Mean abs Avg Rank delta", round(mean_abs(audit_rows, "Avg rank delta"), 6)],
        ["Median abs Avg Rank delta", round(median_abs(audit_rows, "Avg rank delta"), 6)],
        ["Mean abs runoff-rank delta", round(mean_abs(audit_rows, "Runoff rank delta"), 6)],
        ["Median abs runoff-rank delta", round(median_abs(audit_rows, "Runoff rank delta"), 6)],
        ["Largest absolute runoff-rank delta", max(abs(float(r["Runoff rank delta"])) for r in audit_rows)],
    ]

    if county_rollup_csv:
        Path(county_rollup_csv).parent.mkdir(parents=True, exist_ok=True)
        with open(county_rollup_csv, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(audit_sorted[0].keys()))
            writer.writeheader()
            writer.writerows(audit_sorted)

    # Remove older NOAA temperature sheets from prior runs.
    for old_sheet in [IMPACT_SHEET, UPDATE_SHEET]:
        if old_sheet in wb.sheetnames:
            del wb[old_sheet]

    # Remove helper field before writing Value Matrix.
    for r in value_rows_sorted:
        r.pop("_fips5", None)
    for r in value_rows:
        r.pop("_fips5", None)

    write_matrix(wb[RANK_MATRIX_SHEET], [rank_headers] + [[clean_cell(r.get(h)) for h in rank_headers] for r in rank_rows_sorted])
    write_matrix(wb[VALUE_MATRIX_SHEET], [value_headers] + [[clean_cell(r.get(h)) for h in value_headers] for r in value_rows_sorted])
    write_matrix(wb[TOP30_SHEET], [rank_headers] + [[clean_cell(r.get(h)) for h in rank_headers] for r in rank_rows_sorted[:30]])

    comp_headers = ["Position", "County", "State", "Original runoff rank", "New runoff rank", "Runoff rank delta",
                    "Original runoff avg elimination round", "New runoff avg", "Original runoff wins", "New wins",
                    "Original avg rank", "New Avg Rank", "Avg Rank delta"]
    comp_rows = []
    for pos, r in enumerate(rank_rows_sorted[:30], 1):
        key = (r["County"], r["State"])
        old = original[key]
        comp_rows.append([pos, r["County"], r["State"], old["Original runoff rank"], r["Runoff rank"], r["Runoff rank"] - old["Original runoff rank"],
                          old["Original runoff avg elimination round"], r["Runoff avg elimination round"],
                          old["Original runoff wins"], r["Runoff wins"], old["Original avg rank"], r["Avg Rank Python Check"], r["Avg Rank Python Check"] - old["Original avg rank"]])
    write_matrix(wb[RESULT_COMPARISON_SHEET], [comp_headers] + comp_rows)

    ws_impact = recreate_sheet(wb, IMPACT_SHEET, after_name=RESULT_COMPARISON_SHEET)
    write_matrix(ws_impact, [["Metric", "Value"]] + impact_rows)
    ws_impact.cell(1, 4, "Top runoff movers")
    mover_headers = ["County", "State", "Original runoff rank", "New runoff rank", "Runoff rank delta",
                     "Original avg rank", "New avg rank", "Avg rank delta",
                     "Old modeled 90F+ days", "New NOAA 90F+ days", "90F+ days rank delta",
                     "Old modeled days below 50F", "New NOAA days below 50F", "Days below 50F rank delta",
                     "Nearest station km", "Station names used"]
    movers = sorted(audit_rows, key=lambda r: abs(float(r["Runoff rank delta"])), reverse=True)[:15]
    for c, h in enumerate(mover_headers, 4):
        ws_impact.cell(2, c, h)
    for r_idx, row in enumerate(movers, 3):
        for c_idx, h in enumerate(mover_headers, 4):
            ws_impact.cell(r_idx, c_idx, clean_cell(row.get(h)))

    ws_audit = recreate_sheet(wb, UPDATE_SHEET, after_name=IMPACT_SHEET)
    audit_headers = list(audit_sorted[0].keys())
    write_matrix(ws_audit, [audit_headers] + [[clean_cell(r[h]) for h in audit_headers] for r in audit_sorted])

    run_rows = [
        ["Parameter", "Value", "Notes"],
        ["Iterations", iterations, "100k randomized runoff simulations rerun after NOAA temperature update."],
        ["Random seed", seed, "Same seed retained for comparability; deterministic with this script/version."],
        ["Rows/counties", len(rank_rows), "No counties added or removed."],
        ["Rank factors used", len(factor_cols), "Avg Rank excluded from random factor selection."],
        ["Column-selection rule", "Random permutation cycle", "All rank columns are picked exactly once per cycle before reshuffling."],
        ["Tie handling", "Random among tied worst ranks", "If multiple active counties tie for worst rank in selected factor."],
        ["Data import outcome", f"{station_backed} station-backed; {fallback_count} calibrated fallback", "NOAA station normals interpolated to county population-weighted centroids; fallback kept for future missing rows."],
        ["Factor values changed?", "90F+ days and days below 50F", "Walkability, transit, AQI, snowfall, and other factors preserved from input workbook."],
        ["90F+ days formula", "sum P(TMAX >= 90F)", "Daily probability from NOAA DLY-TMAX-NORMAL and DLY-TMAX-STDDEV, then station-blended to county."],
        ["Below-50F days formula", "sum P(TMAX < 50F)", "Uses daily maximum temperature below 50F to preserve the prior factor's scale and intent."],
        ["Missing-data policy", "Spatial station interpolation first; calibrated old-estimate fallback only if no centroid/station match", "No target county needed calibrated fallback in this run."],
        [None, None, None],
        ["Runoff rank", "County", "State", "Runoff avg elimination round", "Runoff wins", "Runoff win rate", "Avg Rank Python Check"],
    ]
    for r in rank_rows_sorted[:30]:
        run_rows.append([r["Runoff rank"], r["County"], r["State"], r["Runoff avg elimination round"], r["Runoff wins"], r["Runoff win rate"], r["Avg Rank Python Check"]])
    write_matrix(wb[RUNOFF_SHEET], run_rows)

    if METHODOLOGY_SHEET in wb.sheetnames:
        ws = wb[METHODOLOGY_SHEET]
        for row in ws.iter_rows(min_row=2):
            value_col = str(row[1].value).strip() if row[1].value is not None else ""
            if value_col == HOT_VALUE_COL:
                row[4].value = "Imported NOAA/NCEI daily temperature normals and standard deviations"
                row[5].value = ("Updated from uploaded NOAA daily temperature normals tar.gz. For each station/day, expected 90F+ days are computed as P(TMAX >= 90F) using DLY-TMAX-NORMAL and DLY-TMAX-STDDEV. "
                                "County values use inverse-distance-squared blending of nearest NOAA stations to population-weighted county centroids from EPA block-group geography. Lower value ranks better.")
                row[6].value = "https://www.ncei.noaa.gov/products/land-based-station/us-climate-normals ; uploaded us-climate-normals_2006-2020_v1.0.1_daily_temperature_by-variable_c20230403.tar.gz"
            elif value_col == COLD_VALUE_COL:
                row[4].value = "Imported NOAA/NCEI daily temperature normals and standard deviations"
                row[5].value = ("Updated from uploaded NOAA daily temperature normals tar.gz. Days below 50F are expected days with TMAX < 50F, using DLY-TMAX-NORMAL and DLY-TMAX-STDDEV. "
                                "This TMAX definition preserves the prior factor's scale better than TAVG<50. County values use inverse-distance-squared nearest-station blending. Lower value ranks better.")
                row[6].value = "https://www.ncei.noaa.gov/products/land-based-station/us-climate-normals ; uploaded us-climate-normals_2006-2020_v1.0.1_daily_temperature_by-variable_c20230403.tar.gz"

    if LOG_SHEET in wb.sheetnames:
        ws = wb[LOG_SHEET]
        next_row = ws.max_row + 1
        log_values = [
            "NOAA/NCEI daily temperature normals 2006-2020",
            "90F+ days; days below 50F",
            "Uploaded us-climate-normals_2006-2020_v1.0.1_daily_temperature_by-variable_c20230403.tar.gz",
            "Imported station daily normals and standard deviations; computed expected threshold-day counts; spatially blended stations to county population-weighted centroids",
            f"Updated {station_backed} station-backed county rows and {fallback_count} fallback rows; reran 100k MCMC",
            "https://www.ncei.noaa.gov/products/land-based-station/us-climate-normals",
        ]
        for c, val in enumerate(log_values, 1):
            ws.cell(next_row, c, val)

    if QA_SHEET in wb.sheetnames:
        qa_rows = [
            ["Check", "Value", "Result / Notes"],
            ["County rows", len(rank_rows), "Expected 413 based on input workbook."],
            ["Direct rank factor count", len(factor_cols), "Direct factor columns ending in ' rank', excluding Runoff rank and Avg-based rank."],
            ["NOAA normal daily rows read", station_stats["normal_rows"], "Rows in dly-temp-normal.csv."],
            ["NOAA stddev daily rows read", station_stats["stddev_rows"], "Rows in dly-temp-stddev.csv."],
            ["NOAA stations with metrics", station_stats["stations_with_temperature_metrics"], "Stations with at least 350 usable daily rows after normal/stddev merge."],
            ["County rows station-backed", station_backed, "County rows assigned NOAA station blend values."],
            ["County rows calibrated fallback", fallback_count, "No station/centroid match; none expected in this run."],
            ["Runoff wins total", int(sum(r["Runoff wins"] for r in rank_rows)), f"Should equal iterations ({iterations})."],
            ["Runoff win rate sum", float(sum(r["Runoff win rate"] for r in rank_rows)), "Should equal 1.0, subject to floating-point precision."],
            ["Changed factors", "90F+ days; days below 50F", "Other value/rank columns preserved from source workbook."],
        ]
        write_matrix(wb[QA_SHEET], qa_rows)

    for sheet_name in [RANK_MATRIX_SHEET, VALUE_MATRIX_SHEET, TOP30_SHEET, RESULT_COMPARISON_SHEET, RUNOFF_SHEET, QA_SHEET, "Sources Not Imported", IMPACT_SHEET, UPDATE_SHEET, METHODOLOGY_SHEET, LOG_SHEET]:
        if sheet_name in wb.sheetnames:
            style_sheet(wb[sheet_name])

    Path(output_workbook).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_workbook)

    return {
        "station_backed": station_backed,
        "fallback_count": fallback_count,
        "impact_rows": impact_rows,
        "top5": [(r["Runoff rank"], r["County"], r["State"], r["Avg Rank Python Check"], r["Runoff wins"]) for r in rank_rows_sorted[:5]],
        "movers": movers,
        "audit_rows": audit_rows,
    }


def parse_args():
    parser = argparse.ArgumentParser(description="Update county ranking workbook with NOAA daily temperature normals and rerun MCMC.")
    parser.add_argument("--input-workbook", required=True)
    parser.add_argument("--noaa-tar", required=True)
    parser.add_argument("--output-workbook", required=True)
    parser.add_argument("--county-rollup-csv", default=None)
    parser.add_argument("--epa-walkability-zip", default=None)
    parser.add_argument("--county-centroids-csv", default=None)
    parser.add_argument("--iterations", type=int, default=DEFAULT_ITERATIONS)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    return parser.parse_args()


def main():
    args = parse_args()
    result = update_workbook(
        input_workbook=args.input_workbook,
        noaa_tar=args.noaa_tar,
        output_workbook=args.output_workbook,
        county_rollup_csv=args.county_rollup_csv,
        epa_walkability_zip=args.epa_walkability_zip,
        county_centroids_csv=args.county_centroids_csv,
        iterations=args.iterations,
        seed=args.seed,
    )
    print(f"Saved updated workbook: {args.output_workbook}")
    if args.county_rollup_csv:
        print(f"Saved county rollup CSV: {args.county_rollup_csv}")
    print(f"Station-backed county rows: {result['station_backed']}")
    print(f"Calibrated fallback rows: {result['fallback_count']}")
    print("Top 5 after update:")
    for row in result["top5"]:
        print(row)


if __name__ == "__main__":
    main()
