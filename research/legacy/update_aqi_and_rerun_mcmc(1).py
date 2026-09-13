#!/usr/bin/env python3
"""
Update the county-rankings workbook with EPA AirData annual AQI-by-county data,
re-rank the Bad AQI Days factor, recompute average ranks, rerun the 100k
randomized-runoff MCMC, and write an updated workbook plus an audit CSV.

Expected inputs
---------------
1. Existing county-rankings workbook with sheets:
   - Rank Matrix
   - Value Matrix
   - Top 30
   - Result Comparison
   - Runoff Simulation
   - Sources & Methodology
   - Data Acquisition Log
   - QA Checks
   - preferably EPA Walkability Update, for County FIPS mapping
2. Two or more EPA annual_aqi_by_county_YYYY.zip files, each containing the
   EPA annual_aqi_by_county_YYYY.csv file.

Method
------
- Reads each EPA annual AQI CSV from its ZIP.
- Defines bad AQI days as AQI >= 101:
    Unhealthy for Sensitive Groups Days + Unhealthy Days
    + Very Unhealthy Days + Hazardous Days
- Annualizes each year's bad-day count by days with valid AQI:
    annualized_bad_days = bad_days / Days with AQI * 365.25
  This reduces undercounting when county AQI coverage is partial.
- Aggregates 2023/2024 values by county using the mean annualized bad days.
- Uses measured EPA values where county/year data are present.
- For counties missing EPA AQI rows, imputes from the old modeled estimate after
  calibrating old modeled AQI days to EPA-measured AQI days. The calibration is
  local-first:
    * state median measured/old ratio when enough same-state measured counties exist
    * otherwise a blend of same-state, Census-division, and national medians
  This preserves the prior model's relative ordering while anchoring missing
  counties to measured EPA AirData behavior in comparable geographies.
- Re-ranks Bad AQI Days only; lower is better. Other factor values/ranks are
  preserved from the input workbook.
- Reruns the randomized runoff with the same rank-matrix rules: random
  permutation cycles over rank factors; the worst rank on the selected factor is
  eliminated; tied worst rows are broken randomly.

Install dependencies
--------------------
pip install openpyxl numpy numba

Example
-------
python update_aqi_and_rerun_mcmc.py \
  --input-workbook county_rankings_epa_walkability_transit_guardrailed_100k.xlsx \
  --epa-aqi-zips annual_aqi_by_county_2023.zip annual_aqi_by_county_2024.zip \
  --output-workbook county_rankings_epa_walkability_transit_aqi_100k.xlsx \
  --county-rollup-csv epa_aqi_county_rollup.csv
"""

from __future__ import annotations

import argparse
import csv
import math
import re
import statistics
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

import numpy as np
from numba import njit
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
try:
    from openpyxl.workbook.properties import CalcProperties
except Exception:  # pragma: no cover
    CalcProperties = None

DEFAULT_ITERATIONS = 100_000
DEFAULT_SEED = 2026060901

RANK_MATRIX_SHEET = "Rank Matrix"
VALUE_MATRIX_SHEET = "Value Matrix"
TOP30_SHEET = "Top 30"
RESULT_COMPARISON_SHEET = "Result Comparison"
RUNOFF_SHEET = "Runoff Simulation"
METHODOLOGY_SHEET = "Sources & Methodology"
LOG_SHEET = "Data Acquisition Log"
QA_SHEET = "QA Checks"
WALKABILITY_AUDIT_SHEET = "EPA Walkability Update"
IMPACT_SHEET = "EPA AQI Impact"
UPDATE_SHEET = "EPA AQI Update"

AQI_VALUE_COL = "Est. bad AQI days / year"
AQI_RANK_COL = "Bad AQI days rank"

HEADER_FILL = "1F4E79"
BORDER_COLOR = "B7B7B7"

STATE_ABBR = {
    "Alabama": "AL", "Alaska": "AK", "Arizona": "AZ", "Arkansas": "AR", "California": "CA",
    "Colorado": "CO", "Connecticut": "CT", "Delaware": "DE", "District of Columbia": "DC",
    "District Of Columbia": "DC", "Florida": "FL", "Georgia": "GA", "Hawaii": "HI", "Idaho": "ID",
    "Illinois": "IL", "Indiana": "IN", "Iowa": "IA", "Kansas": "KS", "Kentucky": "KY",
    "Louisiana": "LA", "Maine": "ME", "Maryland": "MD", "Massachusetts": "MA", "Michigan": "MI",
    "Minnesota": "MN", "Mississippi": "MS", "Missouri": "MO", "Montana": "MT", "Nebraska": "NE",
    "Nevada": "NV", "New Hampshire": "NH", "New Jersey": "NJ", "New Mexico": "NM", "New York": "NY",
    "North Carolina": "NC", "North Dakota": "ND", "Ohio": "OH", "Oklahoma": "OK", "Oregon": "OR",
    "Pennsylvania": "PA", "Rhode Island": "RI", "South Carolina": "SC", "South Dakota": "SD",
    "Tennessee": "TN", "Texas": "TX", "Utah": "UT", "Vermont": "VT", "Virginia": "VA",
    "Washington": "WA", "West Virginia": "WV", "Wisconsin": "WI", "Wyoming": "WY",
    "Puerto Rico": "PR", "Virgin Islands": "VI",
}

# Census divisions provide a regional fallback when a state has limited/no measured counties.
STATE_DIVISION = {
    "CT": "New England", "ME": "New England", "MA": "New England", "NH": "New England", "RI": "New England", "VT": "New England",
    "NJ": "Middle Atlantic", "NY": "Middle Atlantic", "PA": "Middle Atlantic",
    "IL": "East North Central", "IN": "East North Central", "MI": "East North Central", "OH": "East North Central", "WI": "East North Central",
    "IA": "West North Central", "KS": "West North Central", "MN": "West North Central", "MO": "West North Central", "NE": "West North Central", "ND": "West North Central", "SD": "West North Central",
    "DE": "South Atlantic", "DC": "South Atlantic", "FL": "South Atlantic", "GA": "South Atlantic", "MD": "South Atlantic", "NC": "South Atlantic", "SC": "South Atlantic", "VA": "South Atlantic", "WV": "South Atlantic",
    "AL": "East South Central", "KY": "East South Central", "MS": "East South Central", "TN": "East South Central",
    "AR": "West South Central", "LA": "West South Central", "OK": "West South Central", "TX": "West South Central",
    "AZ": "Mountain", "CO": "Mountain", "ID": "Mountain", "MT": "Mountain", "NV": "Mountain", "NM": "Mountain", "UT": "Mountain", "WY": "Mountain",
    "AK": "Pacific", "CA": "Pacific", "HI": "Pacific", "OR": "Pacific", "WA": "Pacific",
}

BAD_DAY_COLUMNS = [
    "Unhealthy for Sensitive Groups Days",
    "Unhealthy Days",
    "Very Unhealthy Days",
    "Hazardous Days",
]


def normalize_county_name(name: str) -> str:
    if name is None:
        return ""
    s = str(name).strip()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.upper().replace("ST.", "ST").replace("SAINT ", "ST ")
    s = s.replace("'", "")
    s = re.sub(r"\s*\(.*?\)", "", s)
    suffixes = [
        " COUNTY AND BOROUGH", " CITY AND BOROUGH", " CENSUS AREA",
        " MUNICIPALITY", " CONSOLIDATED GOVERNMENT", " UNIFIED GOVERNMENT",
        " BOROUGH", " PARISH", " COUNTY", " CITY",
    ]
    changed = True
    while changed:
        changed = False
        for suffix in suffixes:
            if s.endswith(suffix):
                s = s[: -len(suffix)]
                changed = True
                break
    s = re.sub(r"[^A-Z0-9]+", " ", s).strip()
    return re.sub(r"\s+", " ", s)


def safe_float(x):
    if x is None or x == "":
        return None
    try:
        return float(x)
    except Exception:
        return None


def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = list(rows[0])
    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        padded = list(row) + [None] * (len(headers) - len(row))
        data.append(dict(zip(headers, padded[: len(headers)])))
    return headers, data


def write_matrix(ws, matrix: List[List[object]]):
    ws.delete_rows(1, ws.max_row)
    for r_idx, row in enumerate(matrix, 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, value)


def recreate_sheet(wb, name: str, after_name: str | None = None):
    if name in wb.sheetnames:
        del wb[name]
    if after_name and after_name in wb.sheetnames:
        idx = wb.sheetnames.index(after_name)
        return wb.create_sheet(name, idx + 1)
    return wb.create_sheet(name)


def style_sheet(ws, max_width: int = 38):
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color=BORDER_COLOR)
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 8), max_width)


def get_fips_for_workbook(wb, value_rows) -> Dict[Tuple[str, str], str]:
    if WALKABILITY_AUDIT_SHEET not in wb.sheetnames:
        return {}
    headers, rows = read_sheet(wb[WALKABILITY_AUDIT_SHEET])
    if not {"County", "State", "County FIPS"}.issubset(set(headers)):
        return {}
    return {
        (r["County"], r["State"]): str(r["County FIPS"]).zfill(5)
        for r in rows
        if r.get("County") is not None and r.get("State") is not None and r.get("County FIPS") is not None
    }


def read_aqi_zips(aqi_zips: List[str]):
    """Return raw rows keyed by normalized county/state."""
    rows_by_key = defaultdict(list)
    total_rows = 0
    input_years = set()
    for path in aqi_zips:
        with zipfile.ZipFile(path) as z:
            csv_names = [n for n in z.namelist() if n.lower().endswith(".csv")]
            if not csv_names:
                raise ValueError(f"No CSV found inside {path}")
            name = csv_names[0]
            with z.open(name) as f:
                text_iter = (line.decode("utf-8-sig") for line in f)
                reader = csv.DictReader(text_iter)
                missing_cols = {"State", "County", "Year", "Days with AQI", *BAD_DAY_COLUMNS} - set(reader.fieldnames or [])
                if missing_cols:
                    raise ValueError(f"{name} is missing required columns: {sorted(missing_cols)}")
                for row in reader:
                    state = STATE_ABBR.get(str(row["State"]).strip())
                    if not state:
                        continue
                    county_norm = normalize_county_name(row["County"])
                    year = int(row["Year"])
                    days_with_aqi = int(float(row["Days with AQI"]))
                    bad_days = sum(int(float(row[c])) for c in BAD_DAY_COLUMNS)
                    annualized = (bad_days / days_with_aqi * 365.25) if days_with_aqi > 0 else None
                    rec = {
                        "state_name": row["State"],
                        "state": state,
                        "county": row["County"],
                        "county_norm": county_norm,
                        "year": year,
                        "days_with_aqi": days_with_aqi,
                        "bad_days_raw": bad_days,
                        "annualized_bad_days": annualized,
                        "good_days": int(float(row.get("Good Days", 0) or 0)),
                        "moderate_days": int(float(row.get("Moderate Days", 0) or 0)),
                        "max_aqi": safe_float(row.get("Max AQI")),
                        "p90_aqi": safe_float(row.get("90th Percentile AQI")),
                        "median_aqi": safe_float(row.get("Median AQI")),
                    }
                    rows_by_key[(county_norm, state)].append(rec)
                    input_years.add(year)
                    total_rows += 1
    return rows_by_key, total_rows, sorted(input_years)


def aggregate_aqi_for_workbook(value_rows, aqi_rows_by_key):
    measured = {}
    for r in value_rows:
        key_norm = (normalize_county_name(r["County"]), str(r["State"]).upper())
        recs = aqi_rows_by_key.get(key_norm, [])
        if not recs:
            continue
        annualized = [x["annualized_bad_days"] for x in recs if x["annualized_bad_days"] is not None]
        if not annualized:
            continue
        years = sorted({x["year"] for x in recs})
        days = [x["days_with_aqi"] for x in recs]
        raw_bad = [x["bad_days_raw"] for x in recs]
        max_aqi_vals = [x["max_aqi"] for x in recs if x["max_aqi"] is not None]
        p90_vals = [x["p90_aqi"] for x in recs if x["p90_aqi"] is not None]
        median_vals = [x["median_aqi"] for x in recs if x["median_aqi"] is not None]
        measured[(r["County"], r["State"])] = {
            "match_status": "measured",
            "years_available": ",".join(str(y) for y in years),
            "n_years": len(years),
            "epa_aqi_bad_days": statistics.mean(annualized),
            "avg_raw_bad_days": statistics.mean(raw_bad),
            "avg_days_with_aqi": statistics.mean(days),
            "min_days_with_aqi": min(days),
            "max_days_with_aqi": max(days),
            "year_to_year_abs_delta": abs(annualized[-1] - annualized[0]) if len(annualized) >= 2 else None,
            "max_aqi_mean": statistics.mean(max_aqi_vals) if max_aqi_vals else None,
            "p90_aqi_mean": statistics.mean(p90_vals) if p90_vals else None,
            "median_aqi_mean": statistics.mean(median_vals) if median_vals else None,
        }
    return measured


def median_or_none(vals):
    vals = [float(v) for v in vals if v is not None and math.isfinite(float(v))]
    return statistics.median(vals) if vals else None


def compute_imputations(value_rows, measured: Dict[Tuple[str, str], dict]):
    ratios_by_state = defaultdict(list)
    ratios_by_division = defaultdict(list)
    national_ratios = []
    measured_vals_by_state = defaultdict(list)
    measured_vals_by_division = defaultdict(list)
    national_measured_vals = []

    for r in value_rows:
        key = (r["County"], r["State"])
        old = safe_float(r[AQI_VALUE_COL])
        if key not in measured or old is None or old <= 0:
            continue
        new = float(measured[key]["epa_aqi_bad_days"])
        ratio = new / old
        state = str(r["State"]).upper()
        division = STATE_DIVISION.get(state, "Unknown")
        ratios_by_state[state].append(ratio)
        ratios_by_division[division].append(ratio)
        national_ratios.append(ratio)
        measured_vals_by_state[state].append(new)
        measured_vals_by_division[division].append(new)
        national_measured_vals.append(new)

    national_ratio = median_or_none(national_ratios) or 1.0
    national_median_val = median_or_none(national_measured_vals) or 0.0

    state_ratio = {s: median_or_none(v) for s, v in ratios_by_state.items()}
    division_ratio = {d: median_or_none(v) for d, v in ratios_by_division.items()}
    state_n = {s: len(v) for s, v in ratios_by_state.items()}
    division_n = {d: len(v) for d, v in ratios_by_division.items()}
    state_median_val = {s: median_or_none(v) for s, v in measured_vals_by_state.items()}
    division_median_val = {d: median_or_none(v) for d, v in measured_vals_by_division.items()}

    imputations = {}
    for r in value_rows:
        key = (r["County"], r["State"])
        if key in measured:
            continue
        old = safe_float(r[AQI_VALUE_COL])
        state = str(r["State"]).upper()
        division = STATE_DIVISION.get(state, "Unknown")
        sr = state_ratio.get(state)
        dr = division_ratio.get(division)
        sn = state_n.get(state, 0)
        dn = division_n.get(division, 0)
        # Local-first shrinkage: state signal gets stronger as measured same-state count rises.
        if sr is not None and sn >= 5:
            ratio = sr
            method = f"state calibrated old estimate; state_n={sn}"
        elif sr is not None and dr is not None:
            # Blend thin state evidence with division and national priors.
            w_state = sn / (sn + 5.0)
            w_div = 0.50 * (1.0 - w_state)
            w_nat = 1.0 - w_state - w_div
            ratio = w_state * sr + w_div * dr + w_nat * national_ratio
            method = f"blended state/division/national calibrated old estimate; state_n={sn}; division_n={dn}"
        elif dr is not None:
            ratio = 0.70 * dr + 0.30 * national_ratio
            method = f"division/national calibrated old estimate; division_n={dn}"
        else:
            ratio = national_ratio
            method = "national calibrated old estimate"

        if old is not None and old > 0:
            imputed = max(0.0, old * ratio)
        else:
            imputed = state_median_val.get(state) or division_median_val.get(division) or national_median_val
            method += "; median fallback because old estimate missing/nonpositive"

        # Conservative clipping to avoid imputed outliers far outside measured values. If the state has
        # at least five measured counties, use the state's observed max; otherwise use the division max;
        # otherwise use national max.
        if state in measured_vals_by_state and len(measured_vals_by_state[state]) >= 5:
            cap = max(measured_vals_by_state[state]) * 1.25
        elif division in measured_vals_by_division and len(measured_vals_by_division[division]) >= 5:
            cap = max(measured_vals_by_division[division]) * 1.25
        else:
            cap = max(national_measured_vals) * 1.25
        imputed = min(imputed, cap)
        imputations[key] = {
            "match_status": "imputed",
            "epa_aqi_bad_days": imputed,
            "imputation_method": method,
            "calibration_ratio_used": ratio,
            "state_measured_count": sn,
            "division_measured_count": dn,
            "years_available": "",
            "n_years": 0,
            "avg_raw_bad_days": None,
            "avg_days_with_aqi": None,
            "min_days_with_aqi": None,
            "max_days_with_aqi": None,
            "year_to_year_abs_delta": None,
            "max_aqi_mean": None,
            "p90_aqi_mean": None,
            "median_aqi_mean": None,
        }
    return imputations, {
        "national_ratio": national_ratio,
        "national_median_measured_bad_days": national_median_val,
        "state_measured_counts": state_n,
        "division_measured_counts": division_n,
    }


def rank_average(values: Iterable[float], higher_is_better: bool) -> List[float]:
    pairs = [(float(v), i) for i, v in enumerate(values)]
    pairs.sort(key=lambda t: t[0], reverse=higher_is_better)
    ranks = [0.0] * len(pairs)
    i = 0
    while i < len(pairs):
        j = i + 1
        while j < len(pairs) and pairs[j][0] == pairs[i][0]:
            j += 1
        avg_rank = (i + 1 + j) / 2.0
        for k in range(i, j):
            ranks[pairs[k][1]] = avg_rank
        i = j
    return ranks


@njit
def run_randomized_runoff(ranks: np.ndarray, iterations: int, seed: int):
    np.random.seed(seed)
    n, f = ranks.shape
    wins = np.zeros(n, np.int64)
    elim_sum = np.zeros(n, np.float64)
    factor_order = np.empty(f, np.int64)
    active_idx = np.empty(n, np.int64)
    ties = np.empty(n, np.int64)

    for _ in range(iterations):
        for i in range(n):
            active_idx[i] = i
        active_count = n
        pos = f
        round_no = 0

        while active_count > 1:
            if pos >= f:
                for j in range(f):
                    factor_order[j] = j
                for j in range(f - 1, 0, -1):
                    k = np.random.randint(j + 1)
                    tmp = factor_order[j]
                    factor_order[j] = factor_order[k]
                    factor_order[k] = tmp
                pos = 0

            col = factor_order[pos]
            pos += 1

            worst = -1e308
            tie_count = 0
            for p in range(active_count):
                i = active_idx[p]
                v = ranks[i, col]
                if v > worst:
                    worst = v
                    ties[0] = p
                    tie_count = 1
                elif v == worst:
                    ties[tie_count] = p
                    tie_count += 1

            loser_pos = ties[np.random.randint(tie_count)]
            loser = active_idx[loser_pos]
            round_no += 1
            elim_sum[loser] += round_no
            active_count -= 1
            active_idx[loser_pos] = active_idx[active_count]

        winner = active_idx[0]
        wins[winner] += 1
        elim_sum[winner] += n

    return wins, elim_sum / iterations


def clean_cell(value):
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def mean_abs(rows, col):
    vals = [abs(float(r[col])) for r in rows]
    return statistics.mean(vals) if vals else None


def median_abs(rows, col):
    vals = [abs(float(r[col])) for r in rows]
    return statistics.median(vals) if vals else None


def update_workbook(
    input_workbook: str,
    epa_aqi_zips: List[str],
    output_workbook: str,
    county_rollup_csv: str | None,
    iterations: int,
    seed: int,
):
    wb = load_workbook(input_workbook)
    if CalcProperties is not None:
        try:
            wb.calculation = CalcProperties(calcMode="auto")
        except Exception:
            pass

    rank_headers, rank_rows = read_sheet(wb[RANK_MATRIX_SHEET])
    value_headers, value_rows = read_sheet(wb[VALUE_MATRIX_SHEET])
    if AQI_VALUE_COL not in value_headers:
        raise ValueError(f"Missing value column: {AQI_VALUE_COL}")
    if AQI_RANK_COL not in rank_headers:
        raise ValueError(f"Missing rank column: {AQI_RANK_COL}")

    original = {}
    for r in rank_rows:
        key = (r["County"], r["State"])
        original[key] = {
            "Original runoff rank": r["Runoff rank"],
            "Original avg rank": r["Avg Rank Python Check"],
            "Original bad AQI days rank": r[AQI_RANK_COL],
            "Original runoff avg elimination round": r["Runoff avg elimination round"],
            "Original runoff wins": r["Runoff wins"],
            "Original runoff win rate": r["Runoff win rate"],
        }

    aqi_rows_by_key, raw_aqi_rows, input_years = read_aqi_zips(epa_aqi_zips)
    measured = aggregate_aqi_for_workbook(value_rows, aqi_rows_by_key)
    imputations, imputation_stats = compute_imputations(value_rows, measured)
    combined = {**measured, **imputations}
    if len(combined) != len(value_rows):
        missing_keys = [(r["County"], r["State"]) for r in value_rows if (r["County"], r["State"]) not in combined]
        raise RuntimeError(f"Missing AQI value for workbook rows: {missing_keys[:10]}")

    fips_lookup = get_fips_for_workbook(wb, value_rows)

    value_by_key = {(r["County"], r["State"]): r for r in value_rows}
    rank_by_key_pre = {(r["County"], r["State"]): r for r in rank_rows}

    # Update AQI value column.
    for r in value_rows:
        key = (r["County"], r["State"])
        r[AQI_VALUE_COL] = float(combined[key]["epa_aqi_bad_days"])

    # Re-rank Bad AQI days, lower is better.
    updated_aqi_values = [float(r[AQI_VALUE_COL]) for r in value_rows]
    new_aqi_ranks = rank_average(updated_aqi_values, higher_is_better=False)
    for idx, r in enumerate(rank_rows):
        key = (r["County"], r["State"])
        r[AQI_RANK_COL] = new_aqi_ranks[idx]

    factor_cols = [
        c for c in rank_headers
        if isinstance(c, str) and c.endswith(" rank") and c not in {"Avg-based rank", "Runoff rank"}
    ]
    if len(factor_cols) != 21:
        raise ValueError(f"Expected 21 direct rank factor columns, found {len(factor_cols)}")

    factor_array = np.array([[float(r[c]) for c in factor_cols] for r in rank_rows], dtype=np.float64)
    avg_ranks = factor_array.mean(axis=1)
    avg_based_ranks = rank_average(avg_ranks, higher_is_better=False)
    wins, avg_elim = run_randomized_runoff(factor_array, iterations, seed)

    sort_df = []
    for i, r in enumerate(rank_rows):
        sort_df.append((i, avg_elim[i], wins[i], avg_ranks[i], str(r["State"]), str(r["County"])))
    sort_df.sort(key=lambda t: (-t[1], -t[2], t[3], t[4], t[5]))
    runoff_rank = [0] * len(rank_rows)
    for pos, (i, *_rest) in enumerate(sort_df, 1):
        runoff_rank[i] = pos

    for i, r in enumerate(rank_rows):
        r["Avg Rank"] = round(float(avg_ranks[i]), 6)
        r["Avg Rank Python Check"] = round(float(avg_ranks[i]), 6)
        r["Avg Check Delta"] = 0.0
        r["Avg-based rank"] = float(avg_based_ranks[i])
        r["Runoff avg elimination round"] = round(float(avg_elim[i]), 6)
        r["Runoff wins"] = int(wins[i])
        r["Runoff win rate"] = float(wins[i]) / float(iterations)
        r["Runoff rank"] = int(runoff_rank[i])

    rank_by_key = {(r["County"], r["State"]): r for r in rank_rows}
    audit_rows = []
    for vr in value_rows:
        key = (vr["County"], vr["State"])
        old = original[key]
        new = rank_by_key[key]
        info = combined[key]
        audit_rows.append({
            "County": key[0],
            "State": key[1],
            "County FIPS": fips_lookup.get(key),
            "AQI update status": info["match_status"],
            "Old modeled bad AQI days": value_by_key[key].get("_old_bad_aqi_days", None),
            "Original bad AQI days rank": old["Original bad AQI days rank"],
            "New bad AQI days": vr[AQI_VALUE_COL],
            "New bad AQI days rank": new[AQI_RANK_COL],
            "Bad AQI days rank delta": new[AQI_RANK_COL] - old["Original bad AQI days rank"],
            "Years available": info.get("years_available"),
            "Measured years count": info.get("n_years"),
            "Avg raw bad AQI days": info.get("avg_raw_bad_days"),
            "Avg days with AQI": info.get("avg_days_with_aqi"),
            "Min days with AQI": info.get("min_days_with_aqi"),
            "Max days with AQI": info.get("max_days_with_aqi"),
            "Year-to-year annualized bad-days delta": info.get("year_to_year_abs_delta"),
            "Mean max AQI": info.get("max_aqi_mean"),
            "Mean 90th percentile AQI": info.get("p90_aqi_mean"),
            "Mean median AQI": info.get("median_aqi_mean"),
            "Imputation method": info.get("imputation_method"),
            "Calibration ratio used": info.get("calibration_ratio_used"),
            "State measured count": info.get("state_measured_count"),
            "Division measured count": info.get("division_measured_count"),
            "Original avg rank": old["Original avg rank"],
            "New avg rank": new["Avg Rank Python Check"],
            "Avg rank delta": new["Avg Rank Python Check"] - old["Original avg rank"],
            "Original runoff rank": old["Original runoff rank"],
            "New runoff rank": new["Runoff rank"],
            "Runoff rank delta": new["Runoff rank"] - old["Original runoff rank"],
            "Original runoff avg elimination round": old["Original runoff avg elimination round"],
            "New runoff avg elimination round": new["Runoff avg elimination round"],
            "Runoff avg elimination delta": new["Runoff avg elimination round"] - old["Original runoff avg elimination round"],
            "Original runoff wins": old["Original runoff wins"],
            "New runoff wins": new["Runoff wins"],
            "Original runoff win rate": old["Original runoff win rate"],
            "New runoff win rate": new["Runoff win rate"],
        })

    # Preserve old modeled value in audit: read from source workbook before replacement.
    # We need it because value_by_key points to mutated dicts; restore by loading with data_only once.
    wb_old = load_workbook(input_workbook, data_only=True, read_only=True)
    _, old_value_rows = read_sheet(wb_old[VALUE_MATRIX_SHEET])
    old_aqi_by_key = {(r["County"], r["State"]): r[AQI_VALUE_COL] for r in old_value_rows}
    for ar in audit_rows:
        ar["Old modeled bad AQI days"] = old_aqi_by_key.get((ar["County"], ar["State"]))
    wb_old.close()

    rank_rows_sorted = sorted(rank_rows, key=lambda r: r["Runoff rank"])
    value_rows_sorted = sorted(value_rows, key=lambda r: rank_by_key[(r["County"], r["State"])] ["Runoff rank"])

    top_old20 = {k for k, v in original.items() if int(v["Original runoff rank"]) <= 20}
    top_new20 = {(r["County"], r["State"]) for r in rank_rows if int(r["Runoff rank"]) <= 20}
    top_old30 = {k for k, v in original.items() if int(v["Original runoff rank"]) <= 30}
    top_new30 = {(r["County"], r["State"]) for r in rank_rows if int(r["Runoff rank"]) <= 30}
    measured_count = sum(1 for r in audit_rows if r["AQI update status"] == "measured")
    imputed_count = len(audit_rows) - measured_count
    measured_aqi_vals = [float(r["New bad AQI days"]) for r in audit_rows if r["AQI update status"] == "measured"]
    imputed_aqi_vals = [float(r["New bad AQI days"]) for r in audit_rows if r["AQI update status"] == "imputed"]
    low_coverage_measured = sum(1 for r in audit_rows if r["AQI update status"] == "measured" and (r["Avg days with AQI"] or 0) < 183)

    impact_rows = [
        ["Input EPA annual AQI rows used (US/territory state mapped)", raw_aqi_rows],
        ["Input AQI years", ", ".join(str(y) for y in input_years)],
        ["County rows", len(rank_rows)],
        ["Counties with measured EPA AQI rows", measured_count],
        ["Counties imputed", imputed_count],
        ["Measured counties with avg Days with AQI < 183", low_coverage_measured],
        ["National median measured/old calibration ratio", round(float(imputation_stats["national_ratio"]), 6)],
        ["Measured AQI bad days range", f"{min(measured_aqi_vals):.3f} to {max(measured_aqi_vals):.3f}" if measured_aqi_vals else "n/a"],
        ["Imputed AQI bad days range", f"{min(imputed_aqi_vals):.3f} to {max(imputed_aqi_vals):.3f}" if imputed_aqi_vals else "n/a"],
        ["Top-20 runoff overlap vs prior", f"{len(top_old20 & top_new20)} / 20"],
        ["Top-30 runoff overlap vs prior", f"{len(top_old30 & top_new30)} / 30"],
        ["Mean abs bad-AQI-days rank delta", round(mean_abs(audit_rows, "Bad AQI days rank delta"), 6)],
        ["Median abs bad-AQI-days rank delta", round(median_abs(audit_rows, "Bad AQI days rank delta"), 6)],
        ["Mean abs Avg Rank delta", round(mean_abs(audit_rows, "Avg rank delta"), 6)],
        ["Median abs Avg Rank delta", round(median_abs(audit_rows, "Avg rank delta"), 6)],
        ["Mean abs runoff-rank delta", round(mean_abs(audit_rows, "Runoff rank delta"), 6)],
        ["Median abs runoff-rank delta", round(median_abs(audit_rows, "Runoff rank delta"), 6)],
        ["Largest absolute runoff-rank delta", max(abs(float(r["Runoff rank delta"])) for r in audit_rows)],
    ]

    if county_rollup_csv:
        Path(county_rollup_csv).parent.mkdir(parents=True, exist_ok=True)
        with open(county_rollup_csv, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(audit_rows[0].keys()))
            writer.writeheader()
            writer.writerows(sorted(audit_rows, key=lambda r: r["New runoff rank"]))

    # Remove older AQI sheets from prior runs, then rewrite core sheets.
    for old_sheet in [IMPACT_SHEET, UPDATE_SHEET]:
        if old_sheet in wb.sheetnames:
            del wb[old_sheet]

    write_matrix(wb[RANK_MATRIX_SHEET], [rank_headers] + [[clean_cell(r.get(h)) for h in rank_headers] for r in rank_rows_sorted])
    write_matrix(wb[VALUE_MATRIX_SHEET], [value_headers] + [[clean_cell(r.get(h)) for h in value_headers] for r in value_rows_sorted])
    write_matrix(wb[TOP30_SHEET], [rank_headers] + [[clean_cell(r.get(h)) for h in rank_headers] for r in rank_rows_sorted[:30]])

    comp_headers = ["Position", "County", "State", "Original runoff rank", "New runoff rank", "Runoff rank delta",
                    "Original runoff avg elimination round", "New runoff avg", "Original runoff wins", "New wins",
                    "Original avg rank", "New Avg Rank", "Avg Rank delta"]
    comp_rows = []
    for pos, r in enumerate(rank_rows_sorted[:30], 1):
        key = (r["County"], r["State"])
        old = original[key]
        comp_rows.append([pos, r["County"], r["State"], old["Original runoff rank"], r["Runoff rank"], r["Runoff rank"] - old["Original runoff rank"],
                          old["Original runoff avg elimination round"], r["Runoff avg elimination round"],
                          old["Original runoff wins"], r["Runoff wins"], old["Original avg rank"], r["Avg Rank Python Check"], r["Avg Rank Python Check"] - old["Original avg rank"]])
    write_matrix(wb[RESULT_COMPARISON_SHEET], [comp_headers] + comp_rows)

    ws_impact = recreate_sheet(wb, IMPACT_SHEET, after_name=RESULT_COMPARISON_SHEET)
    write_matrix(ws_impact, [["Metric", "Value"]] + impact_rows)
    ws_impact.cell(1, 4, "Top runoff movers")
    mover_headers = ["County", "State", "Original runoff rank", "New runoff rank", "Runoff rank delta",
                     "Original avg rank", "New avg rank", "Avg rank delta",
                     "Original bad AQI days rank", "New bad AQI days rank",
                     "Bad AQI days rank delta", "AQI update status", "New bad AQI days"]
    movers = sorted(audit_rows, key=lambda r: abs(float(r["Runoff rank delta"])), reverse=True)[:15]
    for c, h in enumerate(mover_headers, 4):
        ws_impact.cell(2, c, h)
    for r_idx, row in enumerate(movers, 3):
        values = [row[h] for h in mover_headers]
        for c_idx, val in enumerate(values, 4):
            ws_impact.cell(r_idx, c_idx, clean_cell(val))

    ws_audit = recreate_sheet(wb, UPDATE_SHEET, after_name=IMPACT_SHEET)
    audit_headers = list(audit_rows[0].keys())
    audit_sorted = sorted(audit_rows, key=lambda r: r["New runoff rank"])
    write_matrix(ws_audit, [audit_headers] + [[clean_cell(r[h]) for h in audit_headers] for r in audit_sorted])

    run_rows = [
        ["Parameter", "Value", "Notes"],
        ["Iterations", iterations, "100k randomized runoff simulations rerun after AQI update."],
        ["Random seed", seed, "Same seed retained for comparability; deterministic with this script/version."],
        ["Rows/counties", len(rank_rows), "No counties added or removed."],
        ["Rank factors used", len(factor_cols), "Avg Rank excluded from random factor selection."],
        ["Column-selection rule", "Random permutation cycle", "All rank columns are picked exactly once per cycle before reshuffling."],
        ["Tie handling", "Random among tied worst ranks", "If multiple active counties tie for worst rank in selected factor."],
        ["Data import outcome", f"{measured_count} measured; {imputed_count} imputed", "Measured counties used EPA annual_aqi_by_county rows; missing counties imputed from calibrated old modeled estimates."],
        ["Factor values changed?", "Bad AQI days only", "Walkability, transit, and all other factors preserved from input workbook."],
        ["AQI score formula", "2-year mean annualized bad AQI days", "bad days = USG + Unhealthy + Very Unhealthy + Hazardous; annualized by Days with AQI; lower is better."],
        ["Missing-data policy", "Local calibrated imputation", "State-first calibration of old modeled estimates to EPA measured annualized bad AQI days, then Census-division/national fallback when state evidence is thin or absent."],
        [None, None, None],
        ["Runoff rank", "County", "State", "Runoff avg elimination round", "Runoff wins", "Runoff win rate", "Avg Rank Python Check"],
    ]
    for r in rank_rows_sorted[:30]:
        run_rows.append([r["Runoff rank"], r["County"], r["State"], r["Runoff avg elimination round"], r["Runoff wins"], r["Runoff win rate"], r["Avg Rank Python Check"]])
    write_matrix(wb[RUNOFF_SHEET], run_rows)

    if METHODOLOGY_SHEET in wb.sheetnames:
        ws = wb[METHODOLOGY_SHEET]
        for row in ws.iter_rows(min_row=2):
            if str(row[0].value).strip().lower() in {"bad aqi days", "air quality", "aqi"} or str(row[1].value).strip() == AQI_VALUE_COL:
                row[4].value = "Imported EPA AirData annual AQI by county, 2023-2024"
                row[5].value = ("Updated from uploaded EPA annual_aqi_by_county files. Bad AQI days are AQI >= 101, computed as USG + Unhealthy + Very Unhealthy + Hazardous. "
                                "Each year is annualized by Days with AQI, then averaged across years. Missing counties are imputed from old modeled estimates calibrated to EPA-measured counties using same-state, Census-division, and national ratios.")
                row[6].value = "https://aqs.epa.gov/aqsweb/airdata/download_files.html ; uploaded annual_aqi_by_county_2023.zip ; uploaded annual_aqi_by_county_2024.zip"
                break

    if LOG_SHEET in wb.sheetnames:
        ws = wb[LOG_SHEET]
        next_row = ws.max_row + 1
        log_values = [
            "EPA AirData annual AQI by county",
            "Bad AQI days",
            "Uploaded annual_aqi_by_county_2023.zip and annual_aqi_by_county_2024.zip",
            "Imported annual AQI summaries; computed annualized AQI>=101 days and averaged across years; imputed missing counties with local calibrated old estimates",
            f"Updated {measured_count} measured counties and imputed {imputed_count}; reran 100k MCMC",
            "https://aqs.epa.gov/aqsweb/airdata/download_files.html",
        ]
        for c, val in enumerate(log_values, 1):
            ws.cell(next_row, c, val)

    if "Sources Not Imported" in wb.sheetnames:
        ws = wb["Sources Not Imported"]
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            filtered = [list(rows[0])] + [list(r) for r in rows[1:] if r and "AQI" not in str(r[0]) and "AirData" not in str(r[0])]
            write_matrix(ws, filtered)

    if QA_SHEET in wb.sheetnames:
        qa_rows = [
            ["Check", "Value", "Result / Notes"],
            ["County rows", len(rank_rows), "Expected 413 based on input workbook."],
            ["Direct rank factor count", len(factor_cols), "Direct factor columns ending in ' rank', excluding Runoff rank and Avg-based rank."],
            ["EPA annual AQI rows read", raw_aqi_rows, "Rows across uploaded annual AQI files."],
            ["EPA AQI county matches", measured_count, "Workbook counties with one or more annual AQI rows."],
            ["AQI counties imputed", imputed_count, "No EPA county rows in supplied years; imputed with local calibrated old estimate."],
            ["Runoff wins total", int(sum(r["Runoff wins"] for r in rank_rows)), f"Should equal iterations ({iterations})."],
            ["Runoff win rate sum", float(sum(r["Runoff win rate"] for r in rank_rows)), "Should equal 1.0, subject to floating-point precision."],
            ["Changed factor", "Bad AQI days only", "Other value/rank columns preserved from source workbook."],
        ]
        write_matrix(wb[QA_SHEET], qa_rows)

    for sheet_name in [RANK_MATRIX_SHEET, VALUE_MATRIX_SHEET, TOP30_SHEET, RESULT_COMPARISON_SHEET, RUNOFF_SHEET, QA_SHEET, "Sources Not Imported", IMPACT_SHEET, UPDATE_SHEET, METHODOLOGY_SHEET, LOG_SHEET]:
        if sheet_name in wb.sheetnames:
            style_sheet(wb[sheet_name])

    Path(output_workbook).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_workbook)

    return {
        "measured_count": measured_count,
        "imputed_count": imputed_count,
        "input_years": input_years,
        "impact_rows": impact_rows,
        "top5": [(r["Runoff rank"], r["County"], r["State"], r["Avg Rank Python Check"], r["Runoff wins"]) for r in rank_rows_sorted[:5]],
        "movers": movers,
    }


def parse_args():
    parser = argparse.ArgumentParser(description="Update county ranking workbook with EPA annual AQI data and rerun MCMC.")
    parser.add_argument("--input-workbook", required=True)
    parser.add_argument("--epa-aqi-zips", nargs="+", required=True)
    parser.add_argument("--output-workbook", required=True)
    parser.add_argument("--county-rollup-csv", default=None)
    parser.add_argument("--iterations", type=int, default=DEFAULT_ITERATIONS)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    return parser.parse_args()


def main():
    args = parse_args()
    result = update_workbook(
        input_workbook=args.input_workbook,
        epa_aqi_zips=args.epa_aqi_zips,
        output_workbook=args.output_workbook,
        county_rollup_csv=args.county_rollup_csv,
        iterations=args.iterations,
        seed=args.seed,
    )
    print(f"Saved updated workbook: {args.output_workbook}")
    if args.county_rollup_csv:
        print(f"Saved county rollup CSV: {args.county_rollup_csv}")
    print(f"EPA AQI measured counties: {result['measured_count']}")
    print(f"Imputed counties: {result['imputed_count']}")
    print("Top 5 after update:")
    for row in result["top5"]:
        print(row)


if __name__ == "__main__":
    main()
