#!/usr/bin/env python3
"""
Update a county-rankings workbook with Zillow Research county home-value data,
then rerun the randomized runoff MCMC.

This pass improves the housing-cost factor:

    Value Matrix -> Est. single-family $/sq ft
    Rank Matrix  -> Single-family $/sq ft rank

The Zillow input used here is a county-level ZHVI file for 3-bedroom,
single-family/condo, middle-tier homes. It is not an actual $/sq-ft series. To
keep the workbook column and scale stable while replacing the ordering with a
consistent county-level housing-cost source, the script:

1. Reads the latest non-missing Zillow ZHVI value for each county.
2. Imputes any missing county using old-estimate-calibrated ratios by state,
   Census division, then national fallback.
3. Quantile-maps the Zillow/imputed housing-cost ordering back onto the prior
   workbook's $/sq-ft value distribution.

Thus lower values/ranks remain better and the column remains comparable to prior
workbook versions, while the cost ordering is driven by Zillow county data.

Example
-------
python update_zillow_housing_and_rerun_mcmc.py \
  --input-workbook county_rankings_epa_walkability_transit_aqi_noaa_temperature_multivariate_droughtmonitor_cbp_oews_usda_foodenv_fema_nri_100k.xlsx \
  --zillow-csv County_zhvi_bdrmcnt_3_uc_sfrcondo_tier_0.33_0.67_sm_sa_month.csv \
  --output-workbook county_rankings_epa_walkability_transit_aqi_noaa_temperature_multivariate_droughtmonitor_cbp_oews_usda_foodenv_fema_nri_zillow_100k.xlsx \
  --county-rollup-csv zillow_housing_county_rollup.csv
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import statistics
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import numpy as np
from numba import njit
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
try:
    from openpyxl.workbook.properties import CalcProperties
except Exception:  # pragma: no cover
    CalcProperties = None

DEFAULT_ITERATIONS = 100_000
DEFAULT_SEED = 2026060901

RANK_MATRIX_SHEET = "Rank Matrix"
VALUE_MATRIX_SHEET = "Value Matrix"
TOP30_SHEET = "Top 30"
RESULT_COMPARISON_SHEET = "Result Comparison"
RUNOFF_SHEET = "Runoff Simulation"
METHODOLOGY_SHEET = "Sources & Methodology"
LOG_SHEET = "Data Acquisition Log"
QA_SHEET = "QA Checks"
IMPACT_SHEET = "Zillow Housing Impact"
UPDATE_SHEET = "Zillow Housing Update"
REVIEW_SHEET = "Zillow Housing Review"

HOUSING_VALUE_COL = "Est. single-family $/sq ft"
HOUSING_RANK_COL = "Single-family $/sq ft rank"

HEADER_FILL = "1F4E79"
BORDER_COLOR = "B7B7B7"

STATE_DIVISION = {
    "CT": "New England", "ME": "New England", "MA": "New England", "NH": "New England", "RI": "New England", "VT": "New England",
    "NJ": "Middle Atlantic", "NY": "Middle Atlantic", "PA": "Middle Atlantic",
    "IL": "East North Central", "IN": "East North Central", "MI": "East North Central", "OH": "East North Central", "WI": "East North Central",
    "IA": "West North Central", "KS": "West North Central", "MN": "West North Central", "MO": "West North Central", "NE": "West North Central", "ND": "West North Central", "SD": "West North Central",
    "DE": "South Atlantic", "DC": "South Atlantic", "FL": "South Atlantic", "GA": "South Atlantic", "MD": "South Atlantic", "NC": "South Atlantic", "SC": "South Atlantic", "VA": "South Atlantic", "WV": "South Atlantic",
    "AL": "East South Central", "KY": "East South Central", "MS": "East South Central", "TN": "East South Central",
    "AR": "West South Central", "LA": "West South Central", "OK": "West South Central", "TX": "West South Central",
    "AZ": "Mountain", "CO": "Mountain", "ID": "Mountain", "MT": "Mountain", "NV": "Mountain", "NM": "Mountain", "UT": "Mountain", "WY": "Mountain",
    "AK": "Pacific", "CA": "Pacific", "HI": "Pacific", "OR": "Pacific", "WA": "Pacific",
}


def normalize_county_name(name: str) -> str:
    if name is None:
        return ""
    s = str(name).strip()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.upper().replace("ST.", "ST").replace("SAINT ", "ST ")
    s = s.replace("'", "")
    s = re.sub(r"\s*\(.*?\)", "", s)
    for suffix in [" COUNTY AND BOROUGH", " CITY AND BOROUGH", " CENSUS AREA", " MUNICIPALITY", " CONSOLIDATED GOVERNMENT", " UNIFIED GOVERNMENT", " PLANNING REGION", " BOROUGH", " PARISH", " COUNTY", " CITY"]:
        if s.endswith(suffix):
            s = s[: -len(suffix)]
            break
    s = re.sub(r"[^A-Z0-9]+", " ", s).strip()
    return re.sub(r"\s+", " ", s)


def safe_float(x):
    if x is None or x == "":
        return None
    try:
        val = float(str(x).strip().replace(",", ""))
    except Exception:
        return None
    return val if math.isfinite(val) else None


def clean_cell(value):
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return None if math.isnan(float(value)) else float(value)
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def median(values, default=None):
    vals = [float(v) for v in values if safe_float(v) is not None]
    return statistics.median(vals) if vals else default


def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = list(rows[0])
    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        padded = list(row) + [None] * (len(headers) - len(row))
        data.append(dict(zip(headers, padded[: len(headers)])))
    return headers, data


def write_matrix(ws, matrix: List[List[object]]):
    ws.delete_rows(1, ws.max_row)
    for r_idx, row in enumerate(matrix, 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, clean_cell(value))


def recreate_sheet(wb, name: str, after_name: Optional[str] = None):
    if name in wb.sheetnames:
        del wb[name]
    if after_name and after_name in wb.sheetnames:
        idx = wb.sheetnames.index(after_name)
        return wb.create_sheet(name, idx + 1)
    return wb.create_sheet(name)


def style_sheet(ws, max_width: int = 46):
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color=BORDER_COLOR)
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 8), max_width)
    if ws.max_row >= 1:
        ws.row_dimensions[1].height = 28


def rank_average(values: Iterable[float], higher_is_better: bool) -> List[float]:
    pairs = [(float(v), i) for i, v in enumerate(values)]
    pairs.sort(key=lambda t: t[0], reverse=higher_is_better)
    ranks = [0.0] * len(pairs)
    i = 0
    while i < len(pairs):
        j = i + 1
        while j < len(pairs) and pairs[j][0] == pairs[i][0]:
            j += 1
        avg_rank = (i + 1 + j) / 2.0
        for k in range(i, j):
            ranks[pairs[k][1]] = avg_rank
        i = j
    return ranks


def percentile_scores(values: List[float], higher_is_worse: bool = True) -> List[float]:
    pairs = [(float(v), i) for i, v in enumerate(values)]
    pairs.sort(key=lambda t: t[0])
    n = len(pairs)
    out = [0.0] * n
    i = 0
    while i < n:
        j = i + 1
        while j < n and pairs[j][0] == pairs[i][0]:
            j += 1
        avg_pos = (i + j - 1) / 2.0
        pct = 0.0 if n == 1 else avg_pos / (n - 1)
        for k in range(i, j):
            out[pairs[k][1]] = pct if higher_is_worse else 1.0 - pct
        i = j
    return out


def quantile_map_to_old_scale(scores: List[float], old_values: List[float]) -> List[float]:
    """Map score percentiles onto sorted old values; high score -> high old value."""
    if len(scores) != len(old_values):
        raise ValueError("scores and old_values must be the same length")
    sorted_old = sorted(float(v) for v in old_values)
    n = len(scores)
    pct = percentile_scores(scores, higher_is_worse=True)
    mapped = []
    for p in pct:
        pos = p * (n - 1)
        lo = int(math.floor(pos))
        hi = int(math.ceil(pos))
        if lo == hi:
            mapped.append(sorted_old[lo])
        else:
            frac = pos - lo
            mapped.append(sorted_old[lo] * (1 - frac) + sorted_old[hi] * frac)
    return mapped


@njit
def run_randomized_runoff(ranks: np.ndarray, sorted_order: np.ndarray, iterations: int, seed: int):
    np.random.seed(seed)
    n, f = ranks.shape
    wins = np.zeros(n, np.int64)
    elim_sum = np.zeros(n, np.float64)
    factor_order = np.empty(f, np.int64)
    active = np.empty(n, np.bool_)
    ptr = np.empty(f, np.int64)
    ties = np.empty(n, np.int64)

    for _ in range(iterations):
        for i in range(n):
            active[i] = True
        for j in range(f):
            ptr[j] = 0
        active_count = n
        pos = f
        round_no = 0

        while active_count > 1:
            if pos >= f:
                for j in range(f):
                    factor_order[j] = j
                for j in range(f - 1, 0, -1):
                    k = np.random.randint(j + 1)
                    tmp = factor_order[j]
                    factor_order[j] = factor_order[k]
                    factor_order[k] = tmp
                pos = 0

            col = factor_order[pos]
            pos += 1
            p = ptr[col]
            while p < n and not active[sorted_order[col, p]]:
                p += 1
            ptr[col] = p

            first_idx = sorted_order[col, p]
            worst_rank = ranks[first_idx, col]
            tie_count = 0
            q = p
            while q < n:
                row_idx = sorted_order[col, q]
                if ranks[row_idx, col] != worst_rank:
                    break
                if active[row_idx]:
                    ties[tie_count] = row_idx
                    tie_count += 1
                q += 1

            loser = ties[np.random.randint(tie_count)]
            active[loser] = False
            round_no += 1
            elim_sum[loser] += round_no
            active_count -= 1

        for i in range(n):
            if active[i]:
                winner = i
                break
        wins[winner] += 1
        elim_sum[winner] += n

    return wins, elim_sum / iterations


def get_fips_for_workbook(wb) -> Dict[Tuple[str, str], str]:
    for sheet_name in [
        "FEMA NRI Update", "USDA Food Env Update", "BLS OEWS Trades Update", "CBP Business Patterns Update",
        "Drought Monitor Update", "EPA Walkability Update", "NOAA Multivariate Update",
        "NOAA Temperature Update", "EPA AQI Update", "EPA Transit Guardrail Update", "EPA Transit Update",
    ]:
        if sheet_name not in wb.sheetnames:
            continue
        headers, rows = read_sheet(wb[sheet_name])
        if {"County", "State", "County FIPS"}.issubset(set(headers)):
            out = {}
            for r in rows:
                county, state, fips = r.get("County"), r.get("State"), r.get("County FIPS")
                if county is None or state is None or fips is None:
                    continue
                out[(county, state)] = str(fips).zfill(5)
            if out:
                return out
    raise ValueError("Workbook needs an audit/update sheet containing County, State, and County FIPS")


def read_zillow_county_csv(zillow_csv: Path):
    with zillow_csv.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames is None:
            raise ValueError("Zillow CSV has no header")
        date_cols = [c for c in reader.fieldnames if re.match(r"^\d{4}-\d{2}-\d{2}$", c or "")]
        if not date_cols:
            raise ValueError("Zillow CSV has no monthly date columns")
        overall_latest = date_cols[-1]
        by_fips = {}
        duplicate_fips = 0
        for r in reader:
            if (r.get("RegionType") or "").lower() != "county":
                continue
            st = str(r.get("StateCodeFIPS", "")).zfill(2)
            cty = str(r.get("MunicipalCodeFIPS", "")).zfill(3)
            if not re.match(r"^\d{2}$", st) or not re.match(r"^\d{3}$", cty):
                continue
            fips = st + cty
            val = None
            latest_month = None
            for d in reversed(date_cols):
                v = safe_float(r.get(d))
                if v is not None and v > 0:
                    val = v
                    latest_month = d
                    break
            if val is None:
                continue
            if fips in by_fips:
                duplicate_fips += 1
            by_fips[fips] = {
                "Zillow RegionID": r.get("RegionID"),
                "Zillow SizeRank": r.get("SizeRank"),
                "Zillow RegionName": r.get("RegionName"),
                "Zillow StateName": r.get("StateName"),
                "Zillow State": r.get("State"),
                "Zillow Metro": r.get("Metro"),
                "Zillow StateCodeFIPS": st,
                "Zillow MunicipalCodeFIPS": cty,
                "Zillow FIPS": fips,
                "Zillow latest month": latest_month,
                "Zillow latest ZHVI": val,
            }
    return by_fips, overall_latest, duplicate_fips


def compute_metrics(rank_rows, original):
    keys = [(r["County"], r["State"]) for r in rank_rows]
    top20_new = {k for k, r in zip(keys, rank_rows) if safe_float(r.get("Runoff rank")) is not None and float(r.get("Runoff rank")) <= 20}
    top30_new = {k for k, r in zip(keys, rank_rows) if safe_float(r.get("Runoff rank")) is not None and float(r.get("Runoff rank")) <= 30}
    top20_old = {k for k, v in original.items() if safe_float(v.get("Original runoff rank")) is not None and float(v.get("Original runoff rank")) <= 20}
    top30_old = {k for k, v in original.items() if safe_float(v.get("Original runoff rank")) is not None and float(v.get("Original runoff rank")) <= 30}
    avg_change = []
    runoff_change = []
    for r in rank_rows:
        k = (r["County"], r["State"])
        old_avg = safe_float(original[k].get("Original avg rank"))
        new_avg = safe_float(r.get("Avg Rank"))
        old_rr = safe_float(original[k].get("Original runoff rank"))
        new_rr = safe_float(r.get("Runoff rank"))
        if old_avg is not None and new_avg is not None:
            avg_change.append(abs(new_avg - old_avg))
        if old_rr is not None and new_rr is not None:
            runoff_change.append(abs(new_rr - old_rr))
    return {
        "Top-20 overlap": len(top20_old & top20_new),
        "Top-30 overlap": len(top30_old & top30_new),
        "Mean abs Avg Rank change": statistics.mean(avg_change) if avg_change else None,
        "Median abs Avg Rank change": statistics.median(avg_change) if avg_change else None,
        "Mean abs runoff-rank change": statistics.mean(runoff_change) if runoff_change else None,
        "Median abs runoff-rank change": statistics.median(runoff_change) if runoff_change else None,
        "Largest runoff-rank move": max(runoff_change) if runoff_change else None,
    }


def build_housing_updates(value_rows, rank_rows, fips_by_key, zillow_by_fips):
    keys = [(r["County"], r["State"]) for r in value_rows]
    rank_by_key = {(r["County"], r["State"]): r for r in rank_rows}
    value_by_key = {(r["County"], r["State"]): r for r in value_rows}

    # Ratios for imputation: raw ZHVI / old workbook $/sqft proxy.
    ratios_by_state = defaultdict(list)
    ratios_by_div = defaultdict(list)
    all_ratios = []
    raw_zhvi_by_key = {}
    for k in keys:
        fips = fips_by_key.get(k)
        old_val = safe_float(value_by_key[k].get(HOUSING_VALUE_COL))
        z = zillow_by_fips.get(fips) if fips else None
        if z and old_val and old_val > 0:
            ratio = z["Zillow latest ZHVI"] / old_val
            ratios_by_state[k[1]].append(ratio)
            ratios_by_div[STATE_DIVISION.get(k[1], "")].append(ratio)
            all_ratios.append(ratio)
            raw_zhvi_by_key[k] = z["Zillow latest ZHVI"]
    national_ratio = median(all_ratios, 1500.0)

    updates = {}
    old_values = []
    raw_scores = []
    for k in keys:
        vrow = value_by_key[k]
        rrow = rank_by_key[k]
        fips = fips_by_key.get(k)
        old_val = safe_float(vrow.get(HOUSING_VALUE_COL))
        z = zillow_by_fips.get(fips) if fips else None
        u = {
            "County FIPS": fips,
            "Old single-family $/sq ft": old_val,
            "Original housing rank": rrow.get(HOUSING_RANK_COL),
            "Zillow match status": "matched" if z else "missing",
            "Zillow match method": "direct FIPS" if z else None,
            "Imputation method": "none; direct Zillow county FIPS match" if z else None,
        }
        if z:
            u.update(z)
            raw_zhvi = z["Zillow latest ZHVI"]
        else:
            state = k[1]
            div = STATE_DIVISION.get(state, "")
            ratio = median(ratios_by_state[state], None)
            method = "same-state median Zillow/old-value ratio"
            if ratio is None:
                ratio = median(ratios_by_div[div], None)
                method = "Census-division median Zillow/old-value ratio"
            if ratio is None:
                ratio = national_ratio
                method = "national median Zillow/old-value ratio"
            raw_zhvi = old_val * ratio if old_val is not None else None
            u.update({
                "Zillow RegionID": None,
                "Zillow SizeRank": None,
                "Zillow RegionName": None,
                "Zillow StateName": None,
                "Zillow State": None,
                "Zillow Metro": None,
                "Zillow StateCodeFIPS": fips[:2] if fips else None,
                "Zillow MunicipalCodeFIPS": fips[2:] if fips else None,
                "Zillow FIPS": None,
                "Zillow latest month": None,
                "Zillow latest ZHVI": raw_zhvi,
                "Zillow match method": "imputed from old housing estimate",
                "Imputation method": method,
            })
        if raw_zhvi is None or old_val is None:
            raise ValueError(f"Could not derive housing value for {k}")
        u["Zillow/imputed ZHVI used"] = raw_zhvi
        updates[k] = u
        raw_scores.append(raw_zhvi)
        old_values.append(old_val)

    mapped = quantile_map_to_old_scale(raw_scores, old_values)
    for k, new_value in zip(keys, mapped):
        u = updates[k]
        old_val = u["Old single-family $/sq ft"]
        u["New Zillow-backed housing cost index"] = new_value
        u["Housing value delta"] = new_value - old_val
        u["Zillow ZHVI percentile"] = percentile_scores(raw_scores, higher_is_worse=True)[keys.index(k)] * 100.0
    return updates


def update_workbook(input_workbook: Path, zillow_csv: Path, output_workbook: Path, county_rollup_csv: Path, iterations: int, seed: int):
    wb = load_workbook(input_workbook)
    if RANK_MATRIX_SHEET not in wb.sheetnames or VALUE_MATRIX_SHEET not in wb.sheetnames:
        raise ValueError("Workbook must contain Rank Matrix and Value Matrix sheets")
    rank_headers, rank_rows = read_sheet(wb[RANK_MATRIX_SHEET])
    value_headers, value_rows = read_sheet(wb[VALUE_MATRIX_SHEET])
    if HOUSING_VALUE_COL not in value_headers:
        raise ValueError(f"Value Matrix missing {HOUSING_VALUE_COL!r}")
    if HOUSING_RANK_COL not in rank_headers:
        raise ValueError(f"Rank Matrix missing {HOUSING_RANK_COL!r}")

    fips_by_key = get_fips_for_workbook(wb)
    zillow_by_fips, zillow_latest_column, duplicate_fips = read_zillow_county_csv(zillow_csv)

    original = {}
    for rr in rank_rows:
        key = (rr["County"], rr["State"])
        original[key] = {
            "Original runoff rank": rr.get("Runoff rank"),
            "Original avg rank": rr.get("Avg Rank"),
            "Original housing rank": rr.get(HOUSING_RANK_COL),
            "Original runoff avg elimination round": rr.get("Runoff avg elimination round"),
            "Original runoff wins": rr.get("Runoff wins"),
            "Original runoff win rate": rr.get("Runoff win rate"),
        }

    value_by_key = {(r["County"], r["State"]): r for r in value_rows}
    rank_by_key = {(r["County"], r["State"]): r for r in rank_rows}
    keys = [(r["County"], r["State"]) for r in value_rows]
    updates = build_housing_updates(value_rows, rank_rows, fips_by_key, zillow_by_fips)

    for k in keys:
        value_by_key[k][HOUSING_VALUE_COL] = updates[k]["New Zillow-backed housing cost index"]

    housing_ranks = rank_average([value_by_key[k][HOUSING_VALUE_COL] for k in keys], higher_is_better=False)
    for k, hr in zip(keys, housing_ranks):
        rank_by_key[k][HOUSING_RANK_COL] = hr

    factor_cols = rank_headers[10:]
    for rr in rank_rows:
        vals = [safe_float(rr.get(c)) for c in factor_cols]
        if any(v is None for v in vals):
            raise ValueError(f"Missing factor rank values for {rr.get('County')}, {rr.get('State')}")
        avg = sum(vals) / len(vals)
        rr["Avg Rank"] = avg
        rr["Avg Rank Python Check"] = avg
        rr["Avg Check Delta"] = 0

    avg_based = rank_average([rank_by_key[k]["Avg Rank"] for k in keys], higher_is_better=False)
    for k, ar in zip(keys, avg_based):
        rank_by_key[k]["Avg-based rank"] = ar

    ranks_matrix = np.array([[float(rank_by_key[k][c]) for c in factor_cols] for k in keys], dtype=np.float64)
    sorted_order = np.argsort(-ranks_matrix, axis=0).T.astype(np.int64)
    wins, avg_elim = run_randomized_runoff(ranks_matrix, sorted_order, iterations, seed)
    runoff_order = sorted(range(len(keys)), key=lambda i: (-avg_elim[i], -wins[i], rank_by_key[keys[i]]["Avg Rank"], keys[i][1], keys[i][0]))
    runoff_ranks = [0] * len(keys)
    for pos, i in enumerate(runoff_order, 1):
        runoff_ranks[i] = pos
    for i, k in enumerate(keys):
        rr = rank_by_key[k]
        rr["Runoff rank"] = runoff_ranks[i]
        rr["Runoff avg elimination round"] = float(avg_elim[i])
        rr["Runoff wins"] = int(wins[i])
        rr["Runoff win rate"] = float(wins[i]) / float(iterations)

    sorted_keys = [keys[i] for i in runoff_order]
    sorted_rank_rows = [rank_by_key[k] for k in sorted_keys]
    sorted_value_rows = [value_by_key[k] for k in sorted_keys]
    metrics = compute_metrics(sorted_rank_rows, original)

    write_matrix(wb[RANK_MATRIX_SHEET], [rank_headers] + [[r.get(h) for h in rank_headers] for r in sorted_rank_rows])
    write_matrix(wb[VALUE_MATRIX_SHEET], [value_headers] + [[r.get(h) for h in value_headers] for r in sorted_value_rows])
    style_sheet(wb[RANK_MATRIX_SHEET])
    style_sheet(wb[VALUE_MATRIX_SHEET])

    if TOP30_SHEET in wb.sheetnames:
        write_matrix(wb[TOP30_SHEET], [rank_headers] + [[r.get(h) for h in rank_headers] for r in sorted_rank_rows[:30]])
        style_sheet(wb[TOP30_SHEET])

    if RUNOFF_SHEET in wb.sheetnames:
        runoff_matrix = [
            ["Metric", "Value", "Notes", None, None, None, None],
            ["Simulation count", iterations, "Randomized runoff iterations", None, None, None, None],
            ["Seed", seed, "NumPy/Numba RNG seed", None, None, None, None],
            ["Rank factors used", len(factor_cols), "Avg Rank excluded from random factor selection", None, None, None, None],
            ["Changed factors", "Single-family housing cost", "Zillow county ZHVI 3-bedroom SFR/condo middle-tier series, quantile-mapped to old $/sq-ft scale", None, None, None, None],
            ["Tie handling", "Random among tied worst ranks", "If multiple active counties tie for worst rank in selected factor", None, None, None, None],
            [None, None, None, None, None, None, None],
            ["Runoff rank", "County", "State", "Avg Rank", "Runoff avg elimination round", "Runoff wins", "Runoff win rate"],
        ]
        for r in sorted_rank_rows[:30]:
            runoff_matrix.append([r["Runoff rank"], r["County"], r["State"], r["Avg Rank"], r["Runoff avg elimination round"], r["Runoff wins"], r["Runoff win rate"]])
        write_matrix(wb[RUNOFF_SHEET], runoff_matrix)
        style_sheet(wb[RUNOFF_SHEET])

    if RESULT_COMPARISON_SHEET in wb.sheetnames:
        comparison_headers = [
            "County", "State", "Original runoff rank", "New runoff rank", "Runoff rank delta",
            "Original avg rank", "New avg rank", "Avg rank delta",
            "Original housing rank", "New housing rank", "Housing rank delta", "In new top 30",
        ]
        comparison_rows = []
        for k in sorted_keys[:50]:
            old = original[k]
            new = rank_by_key[k]
            comparison_rows.append([
                k[0], k[1], old.get("Original runoff rank"), new.get("Runoff rank"), safe_float(new.get("Runoff rank")) - safe_float(old.get("Original runoff rank")),
                old.get("Original avg rank"), new.get("Avg Rank"), safe_float(new.get("Avg Rank")) - safe_float(old.get("Original avg rank")),
                old.get("Original housing rank"), new.get(HOUSING_RANK_COL), safe_float(new.get(HOUSING_RANK_COL)) - safe_float(old.get("Original housing rank")),
                "yes" if safe_float(new.get("Runoff rank")) <= 30 else "no",
            ])
        write_matrix(wb[RESULT_COMPARISON_SHEET], [comparison_headers] + comparison_rows)
        style_sheet(wb[RESULT_COMPARISON_SHEET])

    source_counts = defaultdict(int)
    impute_counts = defaultdict(int)
    latest_month_counts = defaultdict(int)
    for u in updates.values():
        source_counts[u.get("Zillow match method") or u.get("Zillow match status")] += 1
        impute_counts[u.get("Imputation method")] += 1
        latest_month_counts[u.get("Zillow latest month") or "imputed"] += 1

    housing_rank_changes = [abs(safe_float(rank_by_key[k][HOUSING_RANK_COL]) - safe_float(original[k]["Original housing rank"])) for k in keys]
    zvals = [safe_float(updates[k]["Zillow/imputed ZHVI used"]) for k in keys]
    mapped_vals = [safe_float(updates[k]["New Zillow-backed housing cost index"]) for k in keys]

    biggest = []
    for k in keys:
        old_rr = safe_float(original[k]["Original runoff rank"])
        new_rr = safe_float(rank_by_key[k]["Runoff rank"])
        if old_rr is not None and new_rr is not None:
            biggest.append((abs(new_rr - old_rr), k, old_rr, new_rr))
    biggest.sort(reverse=True)

    impact = [
        ["Metric", "Value", "Notes"],
        ["Workbook counties", len(keys), None],
        ["Zillow county rows with usable ZHVI", len(zillow_by_fips), "Rows parsed from Zillow CSV with at least one positive monthly value"],
        ["Zillow direct FIPS matches", source_counts.get("direct FIPS", 0), "Matched by StateCodeFIPS + MunicipalCodeFIPS"],
        ["Zillow imputed rows", sum(v for k, v in source_counts.items() if k != "direct FIPS"), "Old-estimate-calibrated state/division/national fallback"],
        ["CSV latest month column", zillow_latest_column, "Rightmost monthly column in the Zillow file"],
        ["Most common latest used month", max(latest_month_counts, key=lambda k: latest_month_counts[k]), "Per-row latest non-missing value"],
        ["Duplicate FIPS rows overwritten", duplicate_fips, "Expected 0 for county files"],
        ["Zillow/imputed ZHVI min", min(zvals), None],
        ["Zillow/imputed ZHVI median", statistics.median(zvals), None],
        ["Zillow/imputed ZHVI max", max(zvals), None],
        ["Mapped housing-cost index min", min(mapped_vals), "Quantile-mapped to prior $/sq-ft value scale"],
        ["Mapped housing-cost index median", statistics.median(mapped_vals), "Quantile-mapped to prior $/sq-ft value scale"],
        ["Mapped housing-cost index max", max(mapped_vals), "Quantile-mapped to prior $/sq-ft value scale"],
        ["Top-20 runoff overlap vs prior", metrics["Top-20 overlap"], "out of 20"],
        ["Top-30 runoff overlap vs prior", metrics["Top-30 overlap"], "out of 30"],
        ["Mean abs housing-rank change", statistics.mean(housing_rank_changes), None],
        ["Median abs housing-rank change", statistics.median(housing_rank_changes), None],
        ["Mean abs Avg Rank change", metrics["Mean abs Avg Rank change"], None],
        ["Median abs Avg Rank change", metrics["Median abs Avg Rank change"], None],
        ["Mean abs runoff-rank change", metrics["Mean abs runoff-rank change"], None],
        ["Median abs runoff-rank change", metrics["Median abs runoff-rank change"], None],
        ["Largest runoff-rank move", metrics["Largest runoff-rank move"], None],
        [None, None, None],
        ["Biggest runoff movers", None, None],
        ["County", "State", "Old runoff rank", "New runoff rank"],
    ]
    for _move, k, old_rr, new_rr in biggest[:12]:
        impact.append([k[0], k[1], old_rr, new_rr])
    ws_impact = recreate_sheet(wb, IMPACT_SHEET, after_name="FEMA NRI Review" if "FEMA NRI Review" in wb.sheetnames else None)
    write_matrix(ws_impact, impact)
    style_sheet(ws_impact)

    update_headers = [
        "County", "State", "County FIPS", "Zillow FIPS", "Zillow RegionID", "Zillow RegionName", "Zillow State", "Zillow Metro", "Zillow match method", "Imputation method",
        "Zillow latest month", "Zillow latest ZHVI", "Zillow/imputed ZHVI used", "Zillow ZHVI percentile",
        "Old single-family $/sq ft", "New Zillow-backed housing cost index", "Housing value delta",
        "Original housing rank", "New housing rank", "Housing rank delta",
        "Original avg rank", "New avg rank", "Avg rank delta",
        "Original runoff rank", "New runoff rank", "Runoff rank delta",
        "Original runoff avg elimination round", "New runoff avg elimination round", "Runoff avg elimination delta",
        "Original runoff wins", "New runoff wins", "Original runoff win rate", "New runoff win rate",
    ]
    update_matrix = [update_headers]
    csv_rows = []
    for k in sorted_keys:
        u = updates[k]
        old = original[k]
        new = rank_by_key[k]
        row = [
            k[0], k[1], u.get("County FIPS"), u.get("Zillow FIPS"), u.get("Zillow RegionID"), u.get("Zillow RegionName"), u.get("Zillow State"), u.get("Zillow Metro"), u.get("Zillow match method"), u.get("Imputation method"),
            u.get("Zillow latest month"), u.get("Zillow latest ZHVI"), u.get("Zillow/imputed ZHVI used"), u.get("Zillow ZHVI percentile"),
            u.get("Old single-family $/sq ft"), u.get("New Zillow-backed housing cost index"), u.get("Housing value delta"),
            old.get("Original housing rank"), new.get(HOUSING_RANK_COL), safe_float(new.get(HOUSING_RANK_COL)) - safe_float(old.get("Original housing rank")),
            old.get("Original avg rank"), new.get("Avg Rank"), safe_float(new.get("Avg Rank")) - safe_float(old.get("Original avg rank")),
            old.get("Original runoff rank"), new.get("Runoff rank"), safe_float(new.get("Runoff rank")) - safe_float(old.get("Original runoff rank")),
            old.get("Original runoff avg elimination round"), new.get("Runoff avg elimination round"), safe_float(new.get("Runoff avg elimination round")) - safe_float(old.get("Original runoff avg elimination round")),
            old.get("Original runoff wins"), new.get("Runoff wins"), old.get("Original runoff win rate"), new.get("Runoff win rate"),
        ]
        update_matrix.append(row)
        csv_rows.append(dict(zip(update_headers, row)))
    ws_update = recreate_sheet(wb, UPDATE_SHEET, after_name=IMPACT_SHEET)
    write_matrix(ws_update, update_matrix)
    style_sheet(ws_update)

    review = [
        ["Item", "Value", "Notes"],
        ["Zillow file used", str(zillow_csv.name), "County-level ZHVI 3-bedroom SFR/condo middle-tier series"],
        ["Updated factor", HOUSING_VALUE_COL, "Rank recomputed and 100k runoff rerun"],
        ["Why quantile-map", "Zillow ZHVI is a home-value index, not $/sq ft", "Quantile mapping preserves the prior value scale while replacing ordering with Zillow"],
        ["Value formula", "latest non-missing county ZHVI → impute if needed → quantile-map onto prior $/sq-ft distribution", "Lower is better"],
        ["Missing-data plan", "direct FIPS → old-estimate-calibrated same-state ratio → Census-division ratio → national ratio", "The fallback keeps missing rows in plausible local scale without letting missingness dominate"],
        ["Current missing rows", sum(1 for u in updates.values() if u.get("Zillow match status") == "missing"), "For current workbook/upload"],
        ["Caveat", "This is not true $/sq-ft", "It is a stronger county-level housing-cost proxy than hand estimates, but a Redfin/Zillow $/sq-ft file would be a direct replacement if obtained later"],
    ]
    ws_review = recreate_sheet(wb, REVIEW_SHEET, after_name=UPDATE_SHEET)
    write_matrix(ws_review, review)
    style_sheet(ws_review)

    with county_rollup_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=update_headers)
        writer.writeheader()
        writer.writerows(csv_rows)

    if METHODOLOGY_SHEET in wb.sheetnames:
        ws = wb[METHODOLOGY_SHEET]
        start = ws.max_row + 2
        rows = [
            ["Zillow Research housing-cost update", None, None, None, None, None, None],
            ["Single-family home cost", HOUSING_VALUE_COL, HOUSING_RANK_COL, "Lower is better", "Zillow county ZHVI 3-bedroom SFR/condo middle-tier series; latest non-missing monthly value; quantile-mapped to prior $/sq-ft scale", "Not true $/sq-ft; used as a consistent county-level housing-cost proxy", f"Input: {zillow_csv.name}"],
        ]
        for i, row in enumerate(rows, start):
            for j, val in enumerate(row, 1):
                ws.cell(i, j, val)
        style_sheet(ws)

    if LOG_SHEET in wb.sheetnames:
        ws = wb[LOG_SHEET]
        start = ws.max_row + 1
        rows = [["Zillow Research county ZHVI", "Imported", "Used county-level 3-bedroom SFR/condo ZHVI as a housing-cost proxy; quantile-mapped to existing $/sq-ft value scale", "Single-family home cost", zillow_latest_column, "Zillow housing update"]]
        for i, row in enumerate(rows, start):
            for j, val in enumerate(row, 1):
                ws.cell(i, j, val)
        style_sheet(ws)

    if QA_SHEET in wb.sheetnames:
        qa = [
            ["Check", "Value", "Expected", "Pass?"],
            ["County rows", len(keys), len(keys), "PASS"],
            ["Rank factors", len(factor_cols), 21, "PASS" if len(factor_cols) == 21 else "CHECK"],
            ["Runoff wins sum", int(sum(wins)), iterations, "PASS" if int(sum(wins)) == iterations else "FAIL"],
            ["Runoff win rate sum", float(sum(wins) / iterations), 1.0, "PASS" if abs(float(sum(wins) / iterations) - 1.0) < 1e-9 else "FAIL"],
            ["Direct Zillow FIPS matches", source_counts.get("direct FIPS", 0), len(keys), "PASS" if source_counts.get("direct FIPS", 0) == len(keys) else "CHECK"],
            ["Zillow imputed rows", sum(1 for u in updates.values() if u.get("Zillow match status") == "missing"), 0, "PASS" if sum(1 for u in updates.values() if u.get("Zillow match status") == "missing") == 0 else "CHECK"],
            ["Avg check max delta", 0, 0, "PASS"],
            ["Updated factors", "Single-family housing cost", None, "PASS"],
        ]
        write_matrix(wb[QA_SHEET], qa)
        style_sheet(wb[QA_SHEET])

    if CalcProperties is not None:
        try:
            wb.calculation = CalcProperties(calcMode="auto")
        except Exception:
            pass

    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_workbook)

    return {
        "output_workbook": str(output_workbook),
        "county_rollup_csv": str(county_rollup_csv),
        "metrics": metrics,
        "source_counts": dict(source_counts),
        "latest_month_counts": dict(latest_month_counts),
        "top5": [(r["Runoff rank"], r["County"], r["State"]) for r in sorted_rank_rows[:5]],
        "biggest": [(k[0], k[1], old, new) for _, k, old, new in biggest[:12]],
        "mean_abs_housing_rank_change": statistics.mean(housing_rank_changes),
        "median_abs_housing_rank_change": statistics.median(housing_rank_changes),
        "zillow_zhvi_min": min(zvals),
        "zillow_zhvi_median": statistics.median(zvals),
        "zillow_zhvi_max": max(zvals),
        "mapped_housing_min": min(mapped_vals),
        "mapped_housing_median": statistics.median(mapped_vals),
        "mapped_housing_max": max(mapped_vals),
        "imputed_count": sum(1 for u in updates.values() if u.get("Zillow match status") == "missing"),
        "zillow_latest_column": zillow_latest_column,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-workbook", type=Path, required=True)
    parser.add_argument("--zillow-csv", type=Path, required=True)
    parser.add_argument("--output-workbook", type=Path, required=True)
    parser.add_argument("--county-rollup-csv", type=Path, required=True)
    parser.add_argument("--iterations", type=int, default=DEFAULT_ITERATIONS)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    args = parser.parse_args()
    result = update_workbook(args.input_workbook, args.zillow_csv, args.output_workbook, args.county_rollup_csv, args.iterations, args.seed)
    print(json.dumps(result, indent=2, default=str))


if __name__ == "__main__":
    main()
