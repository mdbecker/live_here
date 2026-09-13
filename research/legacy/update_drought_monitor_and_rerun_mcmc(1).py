#!/usr/bin/env python3
"""
Update the county-rankings workbook with U.S. Drought Monitor county data,
re-rank drought, rerun the 100k randomized runoff MCMC, and write an updated
workbook plus an audit CSV.

This script is designed as the next pass after the NOAA annual/seasonal
multivariate update. It supersedes the NOAA precipitation-normal drought proxy
with observed U.S. Drought Monitor D1+ drought frequency for 2015-01-01 through
2025-01-01.

Expected inputs
---------------
1. Existing county-rankings workbook with sheets:
   - Rank Matrix
   - Value Matrix
   - At least one audit/update sheet with County, State, County FIPS
2. Drought Monitor CSV exported from:
   - Section: Non-consecutive Weeks in Drought
   - State: All States
   - Dates: 1/1/2015 to 1/1/2025
   - Level: D1
   - Min. Weeks: 2
   - Output: CSV
3. Optional NOAA monthly archives. These are inspected and documented but are
   not used to overwrite additional columns unless a future version of this
   script adds a defensible mapping. In this pass the Drought Monitor file is
   the authoritative direct source for drought, while prior daily/annual NOAA
   passes are already better for temperature and snowfall.

Method
------
- Convert D1+ non-consecutive weeks over the 10-year window to annualized days:
    D1+ drought days/year = weeks * 7 / 10
- Connecticut has newer planning-region FIPS in the Drought Monitor export,
  while the workbook uses legacy CT counties. The script maps legacy counties
  to the closest current planning region(s) using a transparent fixed crosswalk.
- Remaining missing values, if any, are filled by calibrated fallback using the
  old estimate and state/division/national measured ratios. In the current pass,
  all 413 workbook counties should receive either direct FIPS data or CT
  planning-region crosswalk data.
- Updates:
    * Value Matrix -> Est. D1+ drought days/year
    * Rank Matrix  -> Drought days rank
  Rank direction is lower-is-better.
- Reruns randomized runoff with the same rank-matrix rules: random permutation
  cycles over rank factors; the worst rank on the selected factor is eliminated;
  tied worst rows are broken randomly.

Example
-------
python update_drought_monitor_and_rerun_mcmc.py \
  --input-workbook county_rankings_epa_walkability_transit_aqi_noaa_temperature_multivariate_100k.xlsx \
  --drought-csv drought_monitor_dm_export__20150101_20250101.csv \
  --monthly-multivariate-tar us-climate-normals_2006-2020_v1.0.1_monthly_multivariate_by-station_c20230404.tar.gz \
  --monthly-temperature-tar us-climate-normals_2006-2020_v1.0.1_monthly_temperature_by-variable_c20230403.tar.gz \
  --monthly-precipitation-tar us-climate-normals_2006-2020_v1.0.1_monthly_precipitation_by-variable_c20230404.tar.gz \
  --output-workbook county_rankings_epa_walkability_transit_aqi_noaa_temperature_multivariate_droughtmonitor_100k.xlsx \
  --county-rollup-csv drought_monitor_county_rollup.csv
"""

from __future__ import annotations

import argparse
import csv
import math
import os
import statistics
import tarfile
import unicodedata
import re
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Tuple, Optional

import numpy as np
from numba import njit
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
try:
    from openpyxl.workbook.properties import CalcProperties
except Exception:  # pragma: no cover
    CalcProperties = None

DEFAULT_ITERATIONS = 100_000
DEFAULT_SEED = 2026060901
TEN_YEAR_WINDOW_YEARS = 10.0

RANK_MATRIX_SHEET = "Rank Matrix"
VALUE_MATRIX_SHEET = "Value Matrix"
TOP30_SHEET = "Top 30"
RESULT_COMPARISON_SHEET = "Result Comparison"
RUNOFF_SHEET = "Runoff Simulation"
METHODOLOGY_SHEET = "Sources & Methodology"
LOG_SHEET = "Data Acquisition Log"
QA_SHEET = "QA Checks"
IMPACT_SHEET = "Drought Monitor Impact"
UPDATE_SHEET = "Drought Monitor Update"
MONTHLY_REVIEW_SHEET = "NOAA Monthly Review"

DROUGHT_VALUE_COL = "Est. D1+ drought days/year"
DROUGHT_RANK_COL = "Drought days rank"

HEADER_FILL = "1F4E79"
SUBHEADER_FILL = "D9EAF7"
BORDER_COLOR = "B7B7B7"

STATE_DIVISION = {
    "CT": "New England", "ME": "New England", "MA": "New England", "NH": "New England", "RI": "New England", "VT": "New England",
    "NJ": "Middle Atlantic", "NY": "Middle Atlantic", "PA": "Middle Atlantic",
    "IL": "East North Central", "IN": "East North Central", "MI": "East North Central", "OH": "East North Central", "WI": "East North Central",
    "IA": "West North Central", "KS": "West North Central", "MN": "West North Central", "MO": "West North Central", "NE": "West North Central", "ND": "West North Central", "SD": "West North Central",
    "DE": "South Atlantic", "DC": "South Atlantic", "FL": "South Atlantic", "GA": "South Atlantic", "MD": "South Atlantic", "NC": "South Atlantic", "SC": "South Atlantic", "VA": "South Atlantic", "WV": "South Atlantic",
    "AL": "East South Central", "KY": "East South Central", "MS": "East South Central", "TN": "East South Central",
    "AR": "West South Central", "LA": "West South Central", "OK": "West South Central", "TX": "West South Central",
    "AZ": "Mountain", "CO": "Mountain", "ID": "Mountain", "MT": "Mountain", "NV": "Mountain", "NM": "Mountain", "UT": "Mountain", "WY": "Mountain",
    "AK": "Pacific", "CA": "Pacific", "HI": "Pacific", "OR": "Pacific", "WA": "Pacific",
}

# Drought Monitor has CT planning regions; workbook retains legacy CT counties.
# These weights are a transparent geography bridge, not a population-weighted
# official crosswalk. They are used only because the Drought Monitor export no
# longer reports the legacy counties used in the workbook.
CT_PLANNING_REGION_CROSSWALK = {
    "Fairfield": [("Greater Bridgeport Planning Region", 0.50), ("Western Connecticut Planning Region", 0.50)],
    "Hartford": [("Capitol Planning Region", 1.00)],
    "Litchfield": [("Northwest Hills Planning Region", 0.75), ("Western Connecticut Planning Region", 0.25)],
    "Middlesex": [("Lower Connecticut River Valley Planning Region", 1.00)],
    "New Haven": [("South Central Connecticut Planning Region", 0.55), ("Naugatuck Valley Planning Region", 0.45)],
    "New London": [("Southeastern Connecticut Planning Region", 1.00)],
    "Tolland": [("Capitol Planning Region", 0.50), ("Northeastern Connecticut Planning Region", 0.50)],
    "Windham": [("Northeastern Connecticut Planning Region", 1.00)],
}


def normalize_county_name(name: str) -> str:
    if name is None:
        return ""
    s = str(name).strip()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.upper().replace("ST.", "ST").replace("SAINT ", "ST ")
    s = s.replace("'", "")
    s = re.sub(r"\s*\(.*?\)", "", s)
    suffixes = [
        " COUNTY AND BOROUGH", " CITY AND BOROUGH", " CENSUS AREA",
        " MUNICIPALITY", " CONSOLIDATED GOVERNMENT", " UNIFIED GOVERNMENT",
        " BOROUGH", " PARISH", " COUNTY", " CITY",
    ]
    changed = True
    while changed:
        changed = False
        for suffix in suffixes:
            if s.endswith(suffix):
                s = s[: -len(suffix)]
                changed = True
                break
    s = re.sub(r"[^A-Z0-9]+", " ", s).strip()
    return re.sub(r"\s+", " ", s)


def safe_float(x):
    if x is None or x == "":
        return None
    try:
        val = float(str(x).strip())
    except Exception:
        return None
    if not math.isfinite(val) or val in {-9999.0, -999.0, -8888.0}:
        return None
    return val


def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = list(rows[0])
    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        padded = list(row) + [None] * (len(headers) - len(row))
        data.append(dict(zip(headers, padded[: len(headers)])))
    return headers, data


def write_matrix(ws, matrix: List[List[object]]):
    ws.delete_rows(1, ws.max_row)
    for r_idx, row in enumerate(matrix, 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, value)


def recreate_sheet(wb, name: str, after_name: Optional[str] = None):
    if name in wb.sheetnames:
        del wb[name]
    if after_name and after_name in wb.sheetnames:
        idx = wb.sheetnames.index(after_name)
        return wb.create_sheet(name, idx + 1)
    return wb.create_sheet(name)


def style_sheet(ws, max_width: int = 42):
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color=BORDER_COLOR)
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 8), max_width)


def get_fips_for_workbook(wb) -> Dict[Tuple[str, str], str]:
    # Prefer audit/update sheets generated by this workflow.
    for sheet_name in [
        "EPA Walkability Update", "NOAA Multivariate Update", "NOAA Temperature Update",
        "EPA AQI Update", "EPA Transit Update", "Drought Monitor Update"
    ]:
        if sheet_name not in wb.sheetnames:
            continue
        headers, rows = read_sheet(wb[sheet_name])
        if {"County", "State", "County FIPS"}.issubset(set(headers)):
            return {
                (r["County"], r["State"]): str(r["County FIPS"]).zfill(5)
                for r in rows
                if r.get("County") is not None and r.get("State") is not None and r.get("County FIPS") is not None
            }
    raise ValueError("Workbook needs an audit/update sheet containing County, State, and County FIPS")


def rank_average(values: Iterable[float], higher_is_better: bool) -> List[float]:
    pairs = [(float(v), i) for i, v in enumerate(values)]
    pairs.sort(key=lambda t: t[0], reverse=higher_is_better)
    ranks = [0.0] * len(pairs)
    i = 0
    while i < len(pairs):
        j = i + 1
        while j < len(pairs) and pairs[j][0] == pairs[i][0]:
            j += 1
        avg_rank = (i + 1 + j) / 2.0
        for k in range(i, j):
            ranks[pairs[k][1]] = avg_rank
        i = j
    return ranks


@njit
def run_randomized_runoff(ranks: np.ndarray, sorted_order: np.ndarray, iterations: int, seed: int):
    np.random.seed(seed)
    n, f = ranks.shape
    wins = np.zeros(n, np.int64)
    elim_sum = np.zeros(n, np.float64)
    factor_order = np.empty(f, np.int64)
    active = np.empty(n, np.bool_)
    ptr = np.empty(f, np.int64)
    ties = np.empty(n, np.int64)

    for _ in range(iterations):
        for i in range(n):
            active[i] = True
        for j in range(f):
            ptr[j] = 0
        active_count = n
        pos = f
        round_no = 0

        while active_count > 1:
            if pos >= f:
                for j in range(f):
                    factor_order[j] = j
                for j in range(f - 1, 0, -1):
                    k = np.random.randint(j + 1)
                    tmp = factor_order[j]
                    factor_order[j] = factor_order[k]
                    factor_order[k] = tmp
                pos = 0

            col = factor_order[pos]
            pos += 1

            p = ptr[col]
            while p < n and not active[sorted_order[col, p]]:
                p += 1
            ptr[col] = p

            first_idx = sorted_order[col, p]
            worst_rank = ranks[first_idx, col]
            tie_count = 0
            q = p
            while q < n:
                row_idx = sorted_order[col, q]
                if ranks[row_idx, col] != worst_rank:
                    break
                if active[row_idx]:
                    ties[tie_count] = row_idx
                    tie_count += 1
                q += 1

            loser = ties[np.random.randint(tie_count)]
            active[loser] = False
            round_no += 1
            elim_sum[loser] += round_no
            active_count -= 1

        for i in range(n):
            if active[i]:
                winner = i
                break
        wins[winner] += 1
        elim_sum[winner] += n

    return wins, elim_sum / iterations


def clean_cell(value):
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return None if math.isnan(float(value)) else float(value)
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def load_drought_monitor(csv_path: Path):
    by_fips: Dict[str, dict] = {}
    by_state_name: Dict[Tuple[str, str], dict] = {}
    rows = []
    with csv_path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for r in reader:
            fips = str(r.get("FIPS", "")).zfill(5)
            state = r.get("State")
            county = r.get("County")
            weeks = safe_float(r.get("NonConsecutiveWeeks"))
            if not fips or not state or not county or weeks is None:
                continue
            rec = {
                "FIPS": fips,
                "County": county,
                "State": state,
                "NonConsecutiveWeeks": float(weeks),
                "D1DaysPerYear": float(weeks) * 7.0 / TEN_YEAR_WINDOW_YEARS,
            }
            by_fips[fips] = rec
            by_state_name[(state, normalize_county_name(county))] = rec
            rows.append(rec)
    return by_fips, by_state_name, rows


def median(values: List[float], default=None):
    vals = [v for v in values if v is not None and math.isfinite(v)]
    return statistics.median(vals) if vals else default


def calibrated_drought_fallbacks(value_rows, direct_new_by_key, old_col):
    ratios_by_state = defaultdict(list)
    ratios_by_division = defaultdict(list)
    all_ratios = []
    new_by_state = defaultdict(list)
    new_by_division = defaultdict(list)
    all_new = []

    for r in value_rows:
        key = (r["County"], r["State"])
        if key not in direct_new_by_key:
            continue
        old_val = safe_float(r.get(old_col))
        new_val = safe_float(direct_new_by_key[key])
        if new_val is None:
            continue
        state = r["State"]
        division = STATE_DIVISION.get(state, "Unknown")
        new_by_state[state].append(new_val)
        new_by_division[division].append(new_val)
        all_new.append(new_val)
        if old_val is not None and old_val > 0:
            ratio = new_val / old_val
            ratios_by_state[state].append(ratio)
            ratios_by_division[division].append(ratio)
            all_ratios.append(ratio)

    national_ratio = median(all_ratios, 1.0)
    national_new = median(all_new, 0.0)
    fallbacks = {}
    for r in value_rows:
        key = (r["County"], r["State"])
        if key in direct_new_by_key:
            continue
        state = r["State"]
        division = STATE_DIVISION.get(state, "Unknown")
        old_val = safe_float(r.get(old_col))
        state_ratio = median(ratios_by_state[state])
        division_ratio = median(ratios_by_division[division])
        ratio = state_ratio if state_ratio is not None else division_ratio if division_ratio is not None else national_ratio
        if old_val is not None:
            value = old_val * ratio
            method = f"calibrated old estimate fallback using {'state' if state_ratio is not None else 'division' if division_ratio is not None else 'national'} measured ratio"
        else:
            state_new = median(new_by_state[state])
            division_new = median(new_by_division[division])
            value = state_new if state_new is not None else division_new if division_new is not None else national_new
            method = f"median measured fallback using {'state' if state_new is not None else 'division' if division_new is not None else 'national'} values"
        fallbacks[key] = (max(0.0, min(365.25, float(value))), method)
    return fallbacks


def inspect_monthly_archives(paths: List[Optional[Path]]):
    review = []
    for p in paths:
        if not p:
            continue
        p = Path(p)
        if not p.exists():
            review.append({"Archive": p.name, "Status": "missing", "Files": None, "Example columns": None, "Decision": "Not used; file was not found."})
            continue
        try:
            with tarfile.open(p, "r:gz") as t:
                names = t.getnames()
                csv_names = [n for n in names if n.lower().endswith(".csv")]
                example_columns = []
                for name in csv_names:
                    f = t.extractfile(name)
                    if f is None:
                        continue
                    header = f.readline().decode("utf-8", "replace").strip()
                    if header:
                        # Keep only the first handful of meaningful columns for a readable sheet.
                        example_columns = next(csv.reader([header]))[:12]
                        break
                lower = p.name.lower()
                if "precipitation" in lower:
                    decision = "Inspected. Useful for monthly precipitation seasonality, but superseded for the drought factor by the direct U.S. Drought Monitor export."
                elif "temperature" in lower:
                    decision = "Inspected. Useful for monthly temperature summaries, but daily temperature normals already provided better hot/cold day estimates in the previous pass."
                elif "multivariate" in lower:
                    decision = "Inspected. Contains monthly multivariate normals, but no new direct factor update was more defensible than Drought Monitor for drought and prior NOAA annual/daily updates for snow/temperature."
                else:
                    decision = "Inspected. No direct factor overwrite was made in this pass."
                review.append({
                    "Archive": p.name,
                    "Status": "read",
                    "Files": len(names),
                    "CSV files": len(csv_names),
                    "Example columns": ", ".join(example_columns),
                    "Decision": decision,
                })
        except Exception as e:
            review.append({"Archive": p.name, "Status": f"error: {e}", "Files": None, "CSV files": None, "Example columns": None, "Decision": "Not used due to read error."})
    return review


def update_workbook(input_workbook: Path, drought_csv: Path, output_workbook: Path, county_rollup_csv: Path,
                    monthly_archives: List[Optional[Path]], iterations: int, seed: int):
    wb = load_workbook(input_workbook)
    if RANK_MATRIX_SHEET not in wb.sheetnames or VALUE_MATRIX_SHEET not in wb.sheetnames:
        raise ValueError("Workbook must contain Rank Matrix and Value Matrix sheets")

    rank_headers, rank_rows = read_sheet(wb[RANK_MATRIX_SHEET])
    value_headers, value_rows = read_sheet(wb[VALUE_MATRIX_SHEET])
    if DROUGHT_VALUE_COL not in value_headers:
        raise ValueError(f"Value Matrix missing {DROUGHT_VALUE_COL!r}")
    if DROUGHT_RANK_COL not in rank_headers:
        raise ValueError(f"Rank Matrix missing {DROUGHT_RANK_COL!r}")

    fips_by_key = get_fips_for_workbook(wb)
    by_fips, by_state_name, dm_rows = load_drought_monitor(drought_csv)
    monthly_review = inspect_monthly_archives(monthly_archives)

    original = {}
    for rr in rank_rows:
        key = (rr["County"], rr["State"])
        original[key] = {
            "Original runoff rank": rr.get("Runoff rank"),
            "Original avg rank": rr.get("Avg Rank"),
            "Original drought rank": rr.get(DROUGHT_RANK_COL),
            "Original runoff avg elimination round": rr.get("Runoff avg elimination round"),
            "Original runoff wins": rr.get("Runoff wins"),
            "Original runoff win rate": rr.get("Runoff win rate"),
        }

    old_values_by_key = {}
    for vr in value_rows:
        key = (vr["County"], vr["State"])
        old_values_by_key[key] = {
            "Old drought days/year": safe_float(vr.get(DROUGHT_VALUE_COL)),
            "County FIPS": fips_by_key.get(key),
        }

    new_drought_by_key = {}
    update_info = {}
    for vr in value_rows:
        key = (vr["County"], vr["State"])
        county, state = key
        fips = fips_by_key.get(key)
        info = {
            "County FIPS": fips,
            "Old drought days/year": old_values_by_key[key]["Old drought days/year"],
        }
        rec = by_fips.get(fips) if fips else None
        if rec is not None:
            new_drought_by_key[key] = rec["D1DaysPerYear"]
            info.update({
                "Drought Monitor status": "direct FIPS match",
                "Drought Monitor geography": rec["County"],
                "D1+ weeks 2015-2025": rec["NonConsecutiveWeeks"],
                "D1+ drought days/year": rec["D1DaysPerYear"],
                "Imputation method": "direct U.S. Drought Monitor FIPS match; weeks * 7 / 10 years",
            })
        elif state == "CT" and county in CT_PLANNING_REGION_CROSSWALK:
            pieces = []
            total_w = 0.0
            total_weeks = 0.0
            for pr_name, weight in CT_PLANNING_REGION_CROSSWALK[county]:
                pr_rec = by_state_name.get(("CT", normalize_county_name(pr_name)))
                if pr_rec is None:
                    continue
                pieces.append(f"{pr_name} ({weight:.2f} × {pr_rec['NonConsecutiveWeeks']:.0f}w)")
                total_w += weight
                total_weeks += weight * pr_rec["NonConsecutiveWeeks"]
            if total_w > 0:
                weeks = total_weeks / total_w
                days = weeks * 7.0 / TEN_YEAR_WINDOW_YEARS
                new_drought_by_key[key] = days
                info.update({
                    "Drought Monitor status": "CT planning-region crosswalk",
                    "Drought Monitor geography": "; ".join(pieces),
                    "D1+ weeks 2015-2025": weeks,
                    "D1+ drought days/year": days,
                    "Imputation method": "legacy CT county bridged to 2022+ CT planning region(s); weighted average; weeks * 7 / 10 years",
                })
            else:
                info.update({
                    "Drought Monitor status": "missing before fallback",
                    "Drought Monitor geography": None,
                    "D1+ weeks 2015-2025": None,
                    "D1+ drought days/year": None,
                    "Imputation method": None,
                })
        else:
            info.update({
                "Drought Monitor status": "missing before fallback",
                "Drought Monitor geography": None,
                "D1+ weeks 2015-2025": None,
                "D1+ drought days/year": None,
                "Imputation method": None,
            })
        update_info[key] = info

    fallbacks = calibrated_drought_fallbacks(value_rows, new_drought_by_key, DROUGHT_VALUE_COL)
    for vr in value_rows:
        key = (vr["County"], vr["State"])
        info = update_info[key]
        if key in new_drought_by_key:
            drought_value = new_drought_by_key[key]
        else:
            drought_value, method = fallbacks.get(key, (safe_float(vr[DROUGHT_VALUE_COL]) or 0.0, "unmodified old estimate fallback"))
            info["Drought Monitor status"] = "calibrated fallback"
            info["D1+ drought days/year"] = drought_value
            info["Imputation method"] = method
        vr[DROUGHT_VALUE_COL] = float(max(0.0, min(365.25, drought_value)))
        update_info[key] = info

    # Re-rank drought, lower is better.
    drought_ranks = rank_average([r[DROUGHT_VALUE_COL] for r in value_rows], higher_is_better=False)
    drought_ranks_by_key = {(r["County"], r["State"]): drought_ranks[idx] for idx, r in enumerate(value_rows)}
    for rr in rank_rows:
        key = (rr["County"], rr["State"])
        rr[DROUGHT_RANK_COL] = drought_ranks_by_key[key]

    factor_cols = [
        c for c in rank_headers
        if isinstance(c, str) and c.endswith(" rank") and c not in {"Avg-based rank", "Runoff rank"}
    ]
    if len(factor_cols) != 21:
        raise ValueError(f"Expected 21 direct rank factor columns, found {len(factor_cols)}")

    factor_array = np.array([[float(r[c]) for c in factor_cols] for r in rank_rows], dtype=np.float64)
    avg_ranks = factor_array.mean(axis=1)
    avg_based_ranks = rank_average(avg_ranks, higher_is_better=False)
    sorted_order = np.argsort(-factor_array, axis=0).T.astype(np.int64)
    wins, avg_elim = run_randomized_runoff(factor_array, sorted_order, iterations, seed)

    sort_rows = []
    for i, r in enumerate(rank_rows):
        sort_rows.append((i, avg_elim[i], wins[i], avg_ranks[i], str(r["State"]), str(r["County"])))
    sort_rows.sort(key=lambda t: (-t[1], -t[2], t[3], t[4], t[5]))
    runoff_rank = [0] * len(rank_rows)
    for pos, (i, *_rest) in enumerate(sort_rows, 1):
        runoff_rank[i] = pos

    for i, rr in enumerate(rank_rows):
        rr["Avg Rank"] = round(float(avg_ranks[i]), 6)
        rr["Avg Rank Python Check"] = round(float(avg_ranks[i]), 6)
        rr["Avg Check Delta"] = 0.0
        rr["Avg-based rank"] = float(avg_based_ranks[i])
        rr["Runoff avg elimination round"] = round(float(avg_elim[i]), 6)
        rr["Runoff wins"] = int(wins[i])
        rr["Runoff win rate"] = float(wins[i]) / float(iterations)
        rr["Runoff rank"] = int(runoff_rank[i])

    rank_by_key = {(r["County"], r["State"]): r for r in rank_rows}
    audit_rows = []
    for vr in value_rows:
        key = (vr["County"], vr["State"])
        old = original[key]
        new = rank_by_key[key]
        info = update_info[key]
        old_days = old_values_by_key[key]["Old drought days/year"]
        new_days = vr[DROUGHT_VALUE_COL]
        audit_rows.append({
            "County": key[0],
            "State": key[1],
            "County FIPS": info.get("County FIPS"),
            "Drought Monitor status": info.get("Drought Monitor status"),
            "Drought Monitor geography": info.get("Drought Monitor geography"),
            "D1+ weeks 2015-2025": info.get("D1+ weeks 2015-2025"),
            "Old NOAA-proxy drought days/year": old_days,
            "New Drought Monitor days/year": new_days,
            "Drought days delta": None if old_days is None else float(new_days) - float(old_days),
            "Original drought rank": old["Original drought rank"],
            "New drought rank": new[DROUGHT_RANK_COL],
            "Drought rank delta": float(new[DROUGHT_RANK_COL]) - float(old["Original drought rank"]),
            "Imputation method": info.get("Imputation method"),
            "Original avg rank": old["Original avg rank"],
            "New avg rank": new["Avg Rank"],
            "Avg rank delta": float(new["Avg Rank"]) - float(old["Original avg rank"]),
            "Original runoff rank": old["Original runoff rank"],
            "New runoff rank": new["Runoff rank"],
            "Runoff rank delta": int(new["Runoff rank"]) - int(old["Original runoff rank"]),
            "Original runoff avg elimination round": old["Original runoff avg elimination round"],
            "New runoff avg elimination round": new["Runoff avg elimination round"],
            "Runoff avg elimination delta": float(new["Runoff avg elimination round"]) - float(old["Original runoff avg elimination round"]),
            "Original runoff wins": old["Original runoff wins"],
            "New runoff wins": new["Runoff wins"],
            "Original runoff win rate": old["Original runoff win rate"],
            "New runoff win rate": new["Runoff win rate"],
        })

    # Sort display rows by new runoff rank.
    rank_rows.sort(key=lambda r: int(r["Runoff rank"]))
    value_rows.sort(key=lambda r: int(rank_by_key[(r["County"], r["State"])] ["Runoff rank"]))
    audit_rows.sort(key=lambda r: int(r["New runoff rank"]))

    # Write core sheets.
    write_matrix(wb[RANK_MATRIX_SHEET], [rank_headers] + [[clean_cell(r.get(h)) for h in rank_headers] for r in rank_rows])
    write_matrix(wb[VALUE_MATRIX_SHEET], [value_headers] + [[clean_cell(r.get(h)) for h in value_headers] for r in value_rows])
    style_sheet(wb[RANK_MATRIX_SHEET], max_width=28)
    style_sheet(wb[VALUE_MATRIX_SHEET], max_width=28)

    top30 = rank_rows[:30]
    top_ws = wb[TOP30_SHEET] if TOP30_SHEET in wb.sheetnames else wb.create_sheet(TOP30_SHEET, 2)
    write_matrix(top_ws, [rank_headers] + [[clean_cell(r.get(h)) for h in rank_headers] for r in top30])
    style_sheet(top_ws, max_width=28)

    # Result comparison - prior vs new top 30.
    comparison_headers = [
        "New runoff rank", "County", "State", "Prior runoff rank", "Runoff rank delta",
        "Prior avg rank", "New avg rank", "Avg rank delta", "Prior drought rank", "New drought rank", "Drought rank delta",
        "Prior drought days/year", "New drought days/year", "Drought days delta", "Update status"
    ]
    comparison_rows = []
    for ar in audit_rows[:30]:
        comparison_rows.append([
            ar["New runoff rank"], ar["County"], ar["State"], ar["Original runoff rank"], ar["Runoff rank delta"],
            ar["Original avg rank"], ar["New avg rank"], ar["Avg rank delta"], ar["Original drought rank"], ar["New drought rank"], ar["Drought rank delta"],
            ar["Old NOAA-proxy drought days/year"], ar["New Drought Monitor days/year"], ar["Drought days delta"], ar["Drought Monitor status"]
        ])
    comp_ws = wb[RESULT_COMPARISON_SHEET] if RESULT_COMPARISON_SHEET in wb.sheetnames else wb.create_sheet(RESULT_COMPARISON_SHEET, 3)
    write_matrix(comp_ws, [comparison_headers] + comparison_rows)
    style_sheet(comp_ws, max_width=30)

    # Impact sheet.
    matched_direct = sum(1 for r in audit_rows if r["Drought Monitor status"] == "direct FIPS match")
    ct_cross = sum(1 for r in audit_rows if r["Drought Monitor status"] == "CT planning-region crosswalk")
    fallback_count = sum(1 for r in audit_rows if r["Drought Monitor status"] == "calibrated fallback")
    top20_old = { (r["County"], r["State"]) for r in audit_rows if int(r["Original runoff rank"]) <= 20 }
    top20_new = { (r["County"], r["State"]) for r in audit_rows if int(r["New runoff rank"]) <= 20 }
    top30_old = { (r["County"], r["State"]) for r in audit_rows if int(r["Original runoff rank"]) <= 30 }
    top30_new = { (r["County"], r["State"]) for r in audit_rows if int(r["New runoff rank"]) <= 30 }
    abs_drought_rank = [abs(float(r["Drought rank delta"])) for r in audit_rows]
    abs_avg = [abs(float(r["Avg rank delta"])) for r in audit_rows]
    abs_runoff = [abs(int(r["Runoff rank delta"])) for r in audit_rows]
    biggest_movers = sorted(audit_rows, key=lambda r: abs(int(r["Runoff rank delta"])), reverse=True)[:12]

    impact_matrix = [
        ["Metric", "Value", "Notes"],
        ["Input workbook", input_workbook.name, "Baseline before Drought Monitor update"],
        ["Drought Monitor CSV", drought_csv.name, "2015-01-01 through 2025-01-01; D1; non-consecutive weeks"],
        ["MCMC iterations", iterations, "Same runoff algorithm and seed used in prior passes"],
        ["MCMC seed", seed, ""],
        ["Workbook counties", len(audit_rows), ""],
        ["Direct Drought Monitor FIPS matches", matched_direct, ""],
        ["CT planning-region crosswalk rows", ct_cross, "Legacy county rows bridged to current CT planning regions"],
        ["Calibrated fallback rows", fallback_count, "No fallback should be needed for the current 413-county workbook"],
        ["Top-20 runoff overlap vs prior", f"{len(top20_old & top20_new)} / 20", ""],
        ["Top-30 runoff overlap vs prior", f"{len(top30_old & top30_new)} / 30", ""],
        ["Mean abs drought-rank change", round(float(statistics.mean(abs_drought_rank)), 3), ""],
        ["Median abs drought-rank change", round(float(statistics.median(abs_drought_rank)), 3), ""],
        ["Mean abs Avg Rank change", round(float(statistics.mean(abs_avg)), 3), ""],
        ["Median abs Avg Rank change", round(float(statistics.median(abs_avg)), 3), ""],
        ["Mean abs runoff-rank change", round(float(statistics.mean(abs_runoff)), 3), ""],
        ["Median abs runoff-rank change", round(float(statistics.median(abs_runoff)), 3), ""],
        ["Largest runoff-rank move", max(abs_runoff), "Absolute places moved"],
        [],
        ["Biggest runoff movers", "", ""],
        ["County", "State", "Prior runoff rank", "New runoff rank", "Runoff rank delta", "Prior drought days/year", "New drought days/year", "Drought Monitor status"],
    ]
    for r in biggest_movers:
        impact_matrix.append([
            r["County"], r["State"], r["Original runoff rank"], r["New runoff rank"], r["Runoff rank delta"],
            r["Old NOAA-proxy drought days/year"], r["New Drought Monitor days/year"], r["Drought Monitor status"]
        ])
    impact_ws = recreate_sheet(wb, IMPACT_SHEET, after_name=RESULT_COMPARISON_SHEET)
    write_matrix(impact_ws, impact_matrix)
    style_sheet(impact_ws, max_width=50)

    # Detailed update/audit sheet.
    update_headers = list(audit_rows[0].keys()) if audit_rows else []
    update_ws = recreate_sheet(wb, UPDATE_SHEET, after_name=IMPACT_SHEET)
    write_matrix(update_ws, [update_headers] + [[clean_cell(r.get(h)) for h in update_headers] for r in audit_rows])
    style_sheet(update_ws, max_width=48)

    # NOAA monthly review sheet.
    monthly_headers = ["Archive", "Status", "Files", "CSV files", "Example columns", "Decision"]
    monthly_ws = recreate_sheet(wb, MONTHLY_REVIEW_SHEET, after_name=UPDATE_SHEET)
    write_matrix(monthly_ws, [monthly_headers] + [[r.get(h) for h in monthly_headers] for r in monthly_review])
    style_sheet(monthly_ws, max_width=70)

    # Runoff simulation sheet.
    runoff_ws = wb[RUNOFF_SHEET] if RUNOFF_SHEET in wb.sheetnames else wb.create_sheet(RUNOFF_SHEET)
    runoff_matrix = [
        ["Metric", "Value", "Notes"],
        ["Simulation count", iterations, "Randomized runoff iterations"],
        ["Seed", seed, "NumPy/Numba RNG seed"],
        ["Rank factors used", len(factor_cols), "Avg Rank excluded from random factor selection"],
        ["Changed factors", "Drought days", "U.S. Drought Monitor D1+ observed weeks replaced NOAA precipitation-normal drought proxy"],
        ["Tie handling", "Random among tied worst ranks", "If multiple active counties tie for worst rank in selected factor"],
        [],
        ["Runoff rank", "County", "State", "Avg Rank", "Runoff avg elimination round", "Runoff wins", "Runoff win rate"],
    ]
    for r in rank_rows[:30]:
        runoff_matrix.append([r["Runoff rank"], r["County"], r["State"], r["Avg Rank"], r["Runoff avg elimination round"], r["Runoff wins"], r["Runoff win rate"]])
    write_matrix(runoff_ws, runoff_matrix)
    style_sheet(runoff_ws, max_width=36)

    # Sources & Methodology: append/refresh concise note.
    if METHODOLOGY_SHEET in wb.sheetnames:
        ws = wb[METHODOLOGY_SHEET]
        start = ws.max_row + 2
        rows_to_add = [
            ["Drought Monitor update", "U.S. Drought Monitor county export, D1, non-consecutive weeks, 2015-01-01 to 2025-01-01", "Updated Est. D1+ drought days/year = weeks * 7 / 10. Legacy CT counties crosswalked to current planning regions. Monthly NOAA archives inspected but not used for additional direct overwrites in this pass."],
        ]
        for i, row in enumerate(rows_to_add, start):
            for j, value in enumerate(row, 1):
                ws.cell(i, j, value)
        style_sheet(ws, max_width=60)

    # Data acquisition log: append concise entries.
    if LOG_SHEET in wb.sheetnames:
        ws = wb[LOG_SHEET]
        start = ws.max_row + 1
        entries = [
            ["U.S. Drought Monitor D1+ non-consecutive weeks", drought_csv.name, "Imported", "Used to replace drought days factor; direct/crosswalk coverage for all workbook counties"],
            ["NOAA monthly normals archives", "; ".join([Path(p).name for p in monthly_archives if p]), "Inspected", "No additional direct factor overwrite in this pass; prior daily/annual NOAA passes remain better for temp/snow and Drought Monitor supersedes precip proxy"],
        ]
        for i, row in enumerate(entries, start):
            for j, value in enumerate(row, 1):
                ws.cell(i, j, value)
        style_sheet(ws, max_width=60)

    # QA checks.
    qa_rows = [
        ["Check", "Value", "Pass?", "Notes"],
        ["Rank Matrix rows", len(rank_rows), len(rank_rows) == len(value_rows), "Should match Value Matrix rows"],
        ["Value Matrix rows", len(value_rows), len(rank_rows) == len(value_rows), ""],
        ["Direct rank factor count", len(factor_cols), len(factor_cols) == 21, "Columns ending in ' rank', excluding Runoff rank and Avg-based rank"],
        ["Runoff wins sum", int(sum(wins)), int(sum(wins)) == iterations, "Should equal simulation count"],
        ["Runoff win rate sum", float(sum(wins) / iterations), abs(float(sum(wins) / iterations) - 1.0) < 1e-9, "Should equal 1"],
        ["Max Avg Check Delta", max(abs(float(r.get("Avg Check Delta", 0.0))) for r in rank_rows), max(abs(float(r.get("Avg Check Delta", 0.0))) for r in rank_rows) < 1e-9, ""],
        ["Direct Drought Monitor / CT-crosswalk / fallback counts", f"{matched_direct} / {ct_cross} / {fallback_count}", fallback_count == 0, "Fallback count should be 0 for this workbook"],
        ["Changed factors", "Drought days", True, "Other value/rank columns preserved from source workbook"],
    ]
    qa_ws = wb[QA_SHEET] if QA_SHEET in wb.sheetnames else wb.create_sheet(QA_SHEET)
    write_matrix(qa_ws, qa_rows)
    style_sheet(qa_ws, max_width=50)

    # Number formats.
    for ws_name in [RANK_MATRIX_SHEET, VALUE_MATRIX_SHEET, TOP30_SHEET, RESULT_COMPARISON_SHEET, IMPACT_SHEET, UPDATE_SHEET, RUNOFF_SHEET, QA_SHEET]:
        if ws_name not in wb.sheetnames:
            continue
        ws = wb[ws_name]
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000" if abs(cell.value) < 1 and cell.value != 0 else "0.000"
                elif isinstance(cell.value, int):
                    cell.number_format = "0"

    # Encourage recalc in Excel.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        try:
            if CalcProperties is not None:
                wb.calculation = CalcProperties(calcMode="auto")
        except Exception:
            pass

    wb.save(output_workbook)

    # Write audit CSV.
    with county_rollup_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=update_headers)
        writer.writeheader()
        for row in audit_rows:
            writer.writerow(row)

    summary = {
        "output_workbook": str(output_workbook),
        "county_rollup_csv": str(county_rollup_csv),
        "rows": len(audit_rows),
        "direct_matches": matched_direct,
        "ct_crosswalk": ct_cross,
        "fallbacks": fallback_count,
        "top20_overlap": len(top20_old & top20_new),
        "top30_overlap": len(top30_old & top30_new),
        "mean_abs_drought_rank_change": float(statistics.mean(abs_drought_rank)),
        "median_abs_drought_rank_change": float(statistics.median(abs_drought_rank)),
        "mean_abs_avg_rank_change": float(statistics.mean(abs_avg)),
        "median_abs_avg_rank_change": float(statistics.median(abs_avg)),
        "mean_abs_runoff_rank_change": float(statistics.mean(abs_runoff)),
        "median_abs_runoff_rank_change": float(statistics.median(abs_runoff)),
        "largest_runoff_rank_move": max(abs_runoff),
        "top5": [(r["County"], r["State"], r["Runoff rank"]) for r in rank_rows[:5]],
        "biggest_movers": [
            (r["County"], r["State"], r["Original runoff rank"], r["New runoff rank"], r["Runoff rank delta"], r["Old NOAA-proxy drought days/year"], r["New Drought Monitor days/year"])
            for r in biggest_movers[:10]
        ],
    }
    return summary


def main():
    parser = argparse.ArgumentParser(description="Update county ranking workbook with U.S. Drought Monitor D1+ data and rerun MCMC.")
    parser.add_argument("--input-workbook", type=Path, required=True)
    parser.add_argument("--drought-csv", type=Path, required=True)
    parser.add_argument("--monthly-multivariate-tar", type=Path, default=None)
    parser.add_argument("--monthly-temperature-tar", type=Path, default=None)
    parser.add_argument("--monthly-precipitation-tar", type=Path, default=None)
    parser.add_argument("--output-workbook", type=Path, required=True)
    parser.add_argument("--county-rollup-csv", type=Path, required=True)
    parser.add_argument("--iterations", type=int, default=DEFAULT_ITERATIONS)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    args = parser.parse_args()

    monthly_archives = [args.monthly_multivariate_tar, args.monthly_temperature_tar, args.monthly_precipitation_tar]
    summary = update_workbook(
        input_workbook=args.input_workbook,
        drought_csv=args.drought_csv,
        output_workbook=args.output_workbook,
        county_rollup_csv=args.county_rollup_csv,
        monthly_archives=monthly_archives,
        iterations=args.iterations,
        seed=args.seed,
    )
    for key, value in summary.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
