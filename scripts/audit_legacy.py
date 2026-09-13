"""Read-only audit of recovered workbooks. Requires the optional research extra.

Run from the repository root: python scripts/audit_legacy.py
Writes research evidence only; never supplies data to the factor pipeline.
"""

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
LEGACY = ROOT / "research" / "legacy"
DOCS = ROOT / "docs" / "research"


def read_matrix(workbook, name):
    rows = iter(workbook[name].values)
    headers = next(rows)
    return [dict(zip(headers, r)) for r in rows if any(v is not None for v in r)]


def main():
    reports, values, inventory = [], {}, []
    for path in sorted(LEGACY.iterdir()):
        if path.suffix not in {".csv", ".py", ".xlsx"}:
            continue
        inventory.append({"name": path.name, "path": str(path.relative_to(ROOT)),
                          "bytes": path.stat().st_size,
                          "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                          "role": "historical_reference_only"})
        if path.suffix != ".xlsx":
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            ranks = read_matrix(workbook, "Rank Matrix")
            raw = read_matrix(workbook, "Value Matrix")
            values[path.name] = {(r["County"], r["State"]): r for r in raw}
            columns = [c for c in ranks[0] if c.endswith(" rank") and c not in {"Runoff rank", "Avg-based rank"}]
            top = sorted(ranks, key=lambda r: r["Runoff rank"])[:5]
            reports.append({"file": path.name, "sheets": workbook.sheetnames,
                            "county_rows": len(raw), "ranking_rows": len(ranks),
                            "ranking_factor_count": len(columns),
                            "wins": sum(r["Runoff wins"] for r in ranks),
                            "blank_factor_values": sum(r[c] is None for r in raw for c in r if c not in {"County", "State"}),
                            "top5": [{k: r[k] for k in ("County", "State", "Runoff rank", "Avg-based rank")} for r in top]})
        finally:
            workbook.close()
    corrected = values["county_rankings_added_candidates_sourcefix_100k(1).xlsx"]
    earlier = values["county_rankings_full_data_reuse_100k(1).xlsx"]
    woodland = values["county_rankings_with_preserved_woodland_estimate.xlsx"]
    shared = sorted(earlier.keys() & corrected.keys())
    changed = [{"county": key, "column": c, "earlier": earlier[key][c], "corrected": corrected[key][c]}
               for key in shared for c in earlier[key]
               if c not in {"County", "State"} and earlier[key][c] != corrected[key][c]]
    woodland_col = "Est. preserved woodland share (proxy)"
    deviations = [abs(r[woodland_col] - min(.9, r["Est. trees / acre"] / 250)
                      * (.1 + .4 * r["Outdoor recreation score"] / 100)) for r in woodland.values()]
    result = {"audit_type": "saved_values_only_no_rerun", "workbooks": reports,
              "shared_earlier_corrected_counties": len(shared), "changed_shared_factor_cells": changed,
              "woodland_formula_max_absolute_error": max(deviations),
              "woodland_rows_differing_from_corrected": sum(
                  any(woodland[key][c] != corrected[key][c] for c in corrected[key])
                  for key in woodland.keys() & corrected.keys())}
    DOCS.mkdir(parents=True, exist_ok=True)
    (DOCS / "attachment-inventory.json").write_text(json.dumps(inventory, indent=2) + "\n")
    (DOCS / "workbook-audit.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps({"attachments": len(inventory), "workbooks": len(reports),
                      "shared_counties": len(shared), "changed_shared_factor_cells": len(changed),
                      "woodland_max_formula_error": max(deviations)}, indent=2))


if __name__ == "__main__":
    main()
